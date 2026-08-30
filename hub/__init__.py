"""Smart 1 Hub — the shell application.

Owns: login/logout, dashboard, Client 360, Tools landing, Activity, Status,
plus serving the prebuilt Knack "Clients" app (which expects /static and
/data at the site root, so the hub serves those paths for it).
"""
import re
import json
import os
import shutil
import subprocess

import requests as _rq
from flask import (
    Flask, jsonify, make_response, redirect, render_template, request,
    send_from_directory,
)

from . import audit, auth, errors, knack_data
from hub.webargs import clamp_int

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CLIENTS_APP = os.path.join(ROOT, "clients_app")

MODULES = [
    {"key": "clients", "label": "Clients", "href": "/clients", "tag": "Knack"},
    {"key": "google", "label": "Google", "href": "/google/", "tag": "GA4 · GTM"},
    {"key": "sites", "label": "Sites", "href": "/sites/", "tag": "Simvoly"},
    {"key": "suite", "label": "Suite", "href": "/suite/", "tag": "GHL"},
    {"key": "scans", "label": "Site Scans", "href": "/scans/", "tag": "Insites"},
    {"key": "tools", "label": "Tools", "href": "/tools", "tag": ""},
]


def current_user():
    """The signed-in name, for hub views.

    Mirrors auth.user_from_environ: the ordinary cookie first, then the embed
    companion for a page being framed inside Smart 1 Suite. Both readers have
    to agree, or a page would render for a guard that a fetch on the same page
    then 401s — which reads as data that will not load rather than as a login
    problem.
    """
    user = auth.verify_cookie_value(request.cookies.get(auth.COOKIE_NAME))
    if user:
        return user
    try:
        from . import suite_embed as embed
        return embed.user_from_environ(request.environ)
    except Exception:  # noqa: BLE001
        return None


_MOUNT_ACTIVE_HUB = {
    "/tools": "tools", "/qa": "qa", "/activity": "activity",
    "/diagnostics": "diagnostics", "/client360": "client360", "/seo": "seo",
    "/clients": "clients", "/status": "status",
    # Two segments, matched before the one above it. A blueprint tool that has
    # its own sidebar entry has to be able to say so: /tools/website-audit is
    # under Sales in the nav, and resolving it to "tools" highlighted Client
    # Tools instead -- nav pointing at the wrong entry is a small lie the
    # reader corrects by ignoring the highlight.
    "/tools/website-audit": "website_audit",
}


def _hub_active(path: str) -> str:
    """Which sidebar entry a hub path belongs to. Longest prefix wins."""
    parts = (path or "/").strip("/").split("/")
    for depth in (2, 1):
        key = "/" + "/".join(parts[:depth])
        if key in _MOUNT_ACTIVE_HUB:
            return _MOUNT_ACTIVE_HUB[key]
    return ""


def _read_document(raw: bytes, filename: str) -> str:
    """Text out of a PDF or DOCX. Returns "" rather than raising.

    A scanned PDF has no text layer, so this legitimately returns nothing —
    the caller says so plainly rather than reporting a failure.
    """
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader  # type: ignore
            import io as _io
            reader = PdfReader(_io.BytesIO(raw))
            return "\n".join((pg.extract_text() or "") for pg in reader.pages[:40])
        if name.endswith((".docx", ".doc")):
            import io as _io
            from docx import Document
            return "\n".join(p.text for p in Document(_io.BytesIO(raw)).paragraphs)
        return raw.decode("utf-8", errors="ignore")
    except Exception:  # noqa: BLE001
        return ""


def _proposal_text_for(client: str, filename: str) -> str:
    """Read a proposal already uploaded against a client.

    Two storages, and only one of them is a URL. With Cloudinary configured
    the record holds an https link and is fetched. Without it the file went to
    the data disk and the record holds a RELATIVE path -- "/api/client/…" --
    which requests cannot fetch and which used to make every locally stored
    proposal silently unreadable: the fetch raised, the except swallowed it,
    and the caller got "" as though the document were empty. Local files are
    read off the disk instead, which is also the faster answer.

    The record's own stored location is used, never one a caller supplies:
    this reads a URL, so accepting one would be an SSRF hole.
    """
    try:
        from . import proposals
        for rec in proposals.list_proposals(client):
            if not (rec.get("filename") == filename or rec.get("id") == filename):
                continue
            url = str(rec.get("url") or "")
            name = rec.get("filename", "")
            if url.startswith("http://") or url.startswith("https://"):
                import requests as _r
                resp = _r.get(url, timeout=30)
                if not resp.ok:
                    return ""
                return _read_document(resp.content, name)
            if url:
                path = os.path.join(proposals._local_dir(),
                                    os.path.basename(url))
                # basename() above, and this check, because a filename is the
                # one part of the record a person typed.
                if os.path.commonpath([os.path.realpath(path),
                                       os.path.realpath(proposals._local_dir())]
                                      ) != os.path.realpath(proposals._local_dir()):
                    return ""
                with open(path, "rb") as fh:
                    return _read_document(fh.read(), name)
            return ""
    except Exception:  # noqa: BLE001
        pass
    return ""


def create_hub_app() -> Flask:
    app = Flask(
        __name__,
        template_folder="templates",
        static_folder="static",
        static_url_path="/assets",
    )
    app.config["MAX_CONTENT_LENGTH"] = 512 * 1024 * 1024

    # every unhandled exception in the hub app lands in the error log
    from flask import got_request_exception

    def _log_exc(sender, exception, **extra):  # noqa: ARG001
        try:
            errors.log_exception("hub", exception, path=request.path,
                                 actor=current_user() or "")
        except Exception:  # noqa: BLE001 — logging must never break a response
            pass
    got_request_exception.connect(_log_exc, app)

    # ---------------- auth ----------------
    @app.context_processor
    def _inject_demo_module():
        """Which walkthrough belongs to the page being rendered.

        The demo launcher reads <body data-module>. Hub pages had none, so no
        walkthrough button ever appeared on the Dashboard, Client 360, SEO or
        QA — which is most of where somebody would look for one.
        """
        path = (request.path or "/").rstrip("/") or "/"
        mapping = {"/": "hub", "/client360": "hub", "/seo": "seo",
                   "/qa": "qa", "/qa/stale-creative": "qa",
                   "/qa/unattached-images": "qa",
                   "/tools": "hub", "/diagnostics": "hub"}
        return {"hub_demo_module": mapping.get(path, "")}

    @app.context_processor
    def _inject_sidebar():
        """Expose the one shared nav to hub templates."""
        from .sidebar import render_sidebar, collapses_by_default

        def hub_sidebar(active=""):
            try:
                # A page inside a frame sits in somebody else's navigation
                # already. Suppressing the injector in after_request is not
                # enough: base.html calls this global directly, so Client 360
                # framed in Suite would still render its own full sidebar and
                # the injector -- which skips a body that already has one --
                # would see nothing to do and agree that all was well.
                from . import suite_embed as embed
                if embed.is_embedded(request.environ):
                    return ""
            except Exception:  # noqa: BLE001
                pass
            try:
                # The same one decision the injector makes. base.html calls
                # this global directly, so a hub page rendering its own nav
                # has to reach the same answer or a tool would behave one way
                # through its template and another through the injector.
                try:
                    collapsed = collapses_by_default(request.path)
                except Exception:  # noqa: BLE001
                    collapsed = False
                return render_sidebar(active or "", is_admin=viewer_is_admin(),
                                      collapsed_default=collapsed).decode()
            except Exception:  # noqa: BLE001 — nav must never break a page
                return ""
        return {"hub_sidebar": hub_sidebar}

    @app.context_processor
    def _inject_version():
        """Every page footer shows the running build, so you can open the
        deployed site and confirm it matches what you pushed."""
        from . import version as _v
        return {"hub_version": _v.label(), "hub_version_info": _v.info()}

    @app.route("/api/version")
    def api_version():
        from . import version as _v
        return jsonify(_v.info())

    # ---------------- v7.5: diagnostics, quotas, cost ----------------
    @app.route("/diagnostics")
    def page_diagnostics():
        gate = _require_page()
        if gate:
            return gate
        return render_template("diagnostics.html", user=current_user(),
                               active="diagnostics")

    @app.route("/api/diagnostics")
    def api_diagnostics():
        """Live reachability of every external API.

        Deliberately never spends a credit: Insites has no free endpoint, so it
        reports `unverified` rather than starting a throwaway audit.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import diagnostics
        return jsonify(diagnostics.run_all())

    @app.route("/api/quotas")
    def api_quotas():
        """Monthly usage against allowances, and every provider cost estimate.

        `?live=1` also asks ElevenLabs and Cloudinary for their own counters,
        which are the authority on what those two actually bill. Off by
        default and cached for five minutes: everything else here is a local
        read of the activity log, and a page that always makes two outbound
        calls is a page that hangs when someone else's API does.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import quotas
        live = (request.args.get("live") or "").lower() in ("1", "true", "yes", "on")
        return jsonify(quotas.summary(request.args.get("month"), live=live))

    # ---------------- Google account index ----------------
    # Deliberately hub routes, not routes under /google: DispatcherMiddleware
    # routes purely by prefix, so anything under /google belongs to Google
    # Finder and a hub route there would never be reached. That trap has bitten
    # three times and /api/integrity has a high-severity check for it.
    @app.route("/api/google/index")
    def api_google_index():
        """What the stored Google index holds, and how old it is."""
        gate = _require_api()
        if gate:
            return gate
        from . import google_index
        return jsonify(google_index.status())

    @app.route("/api/google/for-client")
    def api_google_for_client():
        """Every Google resource joined to one client.

        This is what Client 360 and the tool auto-fill read. It is a scan of a
        stored dictionary — no Google call — which is the entire reason the
        360 page stopped waiting on a four-API sweep.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import google_index
        return jsonify(google_index.for_client(
            request.args.get("name", ""), request.args.get("url", "")))

    # Also a hub route, not one under /google, and under /tools rather than
    # /tools/<something-mounted> — the same trap the note above describes.
    @app.route("/tools/google-match")
    def page_google_match():
        gate = _require_page()
        if gate:
            return gate
        return render_template("google_match.html", user=current_user(),
                               active="google")

    @app.route("/api/google/read-labels", methods=["POST"])
    def api_google_read_labels():
        """Read the orphan resource labels for the business inside them.

        A POST, and a button, because the call is billed. The model is shown
        the label only — never the client book — and answers with a run of
        words out of it; that name then goes through the same
        `client_key.resolve()` `suggest_for()` already runs on the raw label,
        so the rules that decide are unchanged and a reading can only change
        which string gets asked about. It never outranks a recorded id or a
        domain, and every row still ends at a human pressing Attach.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import google_links, google_names_ai
        book = google_links._orphan_book()          # noqa: SLF001
        labels = google_names_ai.labels_of(book.get("rows") or [])
        out = google_names_ai.read_missing(labels)
        # The orphan list is held for the day and a reading changes what it
        # would say, so it is dropped here rather than leaving the page
        # showing the answer from before the button was pressed — which reads
        # as a button that did nothing. The drop lives beside the write, per
        # hub/report_cache.py.
        try:
            from . import report_cache
            report_cache.invalidate("google-orphans")
        except Exception:                           # noqa: BLE001
            pass
        try:
            audit.log("google_links", "read_labels", actor=current_user() or "",
                      read=out.get("read", 0), stored=out.get("stored", 0),
                      ungrounded=out.get("ungrounded", 0))
        except Exception:                           # noqa: BLE001
            pass
        return jsonify(out)

    @app.route("/api/google/orphans")
    def api_google_orphans():
        """Google resources no client is attached to, with who they might be.

        The other half of the index: it can say what it joined, and this says
        what it could not. An index that has never been built reports that,
        never "no orphans".
        """
        gate = _require_api()
        if gate:
            return gate
        from .google_links import orphans
        plats = [p for p in
                 str(request.args.get("platforms", "")).split(",") if p.strip()]
        return jsonify(orphans(
            q=request.args.get("q", ""),
            platform=request.args.get("platform", ""),
            platforms=plats,
            include_other=str(request.args.get("include_other", "")).lower()
            in ("1", "true", "yes"),
            # A page at a time. Clamped at both ends for the reason
            # hub/webargs.py exists: ?limit=-1 was a 500 on Postgres and a
            # full dump on SQLite.
            limit=clamp_int(request.args.get("limit"), 25, 1, 200),
            offset=clamp_int(request.args.get("offset"), 0, 0, 100000)))

    @app.route("/api/google/attach", methods=["POST"])
    def api_google_attach():
        """Attach one Google resource to a client, or several at once."""
        gate = _require_api()
        if gate:
            return gate
        body = request.get_json(silent=True) or {}
        actor = current_user() or ""
        links = body.get("links")
        if links:
            from .google_links import attach_many
            out = attach_many(links, actor=actor, force=bool(body.get("force")))
        else:
            from .google_links import attach
            out = attach(str(body.get("resource_id") or ""),
                         str(body.get("client") or ""), actor=actor,
                         force=bool(body.get("force")))
        # The reports these rows come off are dropped by
        # `google_index._forget_reports()`, which runs inside `set_client()` —
        # one description of what an attachment invalidates, next to the write
        # that causes it, rather than a second copy here that drifts.
        return jsonify(out)

    @app.route("/api/google/rebuild", methods=["POST"])
    def api_google_rebuild():
        """Force a re-sweep now, rather than waiting for the three-hour job."""
        gate = _require_api()
        if gate:
            return gate
        from . import google_index
        return jsonify(google_index.build(force=True))

    @app.route("/api/backup")
    def api_backup():
        """What of the JSON on the disk is mirrored into the database.

        The disk is not part of the database backup and does not survive being
        recreated, so this answers "what would we actually lose?" — including
        the files deliberately left out because they are rebuildable.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import jsonstore
        out = jsonstore.status()
        out["restore_at_boot"] = app.config.get("HUB_JSONSTORE_RESTORE")
        return jsonify(out)

    @app.route("/api/environment")
    def api_environment():
        """Which environment variable actually supplied each setting.

        A setting here answers to several names — PEXELS_API and
        PEXELS_API_KEY, GHL_PRIVATE_TOKEN and SMART1SUITE_PRIVATE_TOKEN,
        SECRET_KEY and FLASK_SECRET_KEY and SESSION_SECRET — because this
        deployment's environment is assembled from linked Render env groups
        that each named things their own way. That resolves the key, and it
        also means nobody can tell *which* name did it. On a second deployment
        that matters twice: a variable set under a name nothing reads looks
        exactly like one that took effect, and two names holding different
        values silently resolve to whichever comes first.

        No value is ever returned — this is read on a screen and pasted into
        chats, the rule services/provider_check.py already works to.
        """
        gate = _require_api()
        if gate:
            return gate
        from .config import settings as _cfg
        return jsonify({"settings": _cfg.env_report(),
                        "problems": _cfg.placeholder_warnings()})

    @app.route("/api/oauth-redirects")
    def api_oauth_redirects():
        """Every OAuth callback this Hub sends, and where it must be registered.

        A redirect URI is matched by the provider as an exact string, hostname
        included, and half of these are built from whichever hostname the
        browser used. So a custom domain added to the service silently doubles
        what has to be registered, and the first sign of the half that was not
        is a customer meeting redirect_uri_mismatch on a consent screen.

        The origin is taken from the request rather than from PUBLIC_BASE_URL,
        because the point is to name the hostname somebody is actually on —
        reading the variable instead would report the one host that was never
        in doubt. No client id or secret is ever returned.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import oauth_redirects
        return jsonify(oauth_redirects.report(request.url_root))

    @app.route("/api/report-cache")
    def api_report_cache():
        """What the day cache is holding, and how old each entry is.

        Names, days and sizes only. No payloads: this is read into a page,
        and a report's rows carry client names — the rule
        `services/provider_check.py` works to about what a diagnostic may
        carry.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import report_cache
        return jsonify(report_cache.state())

    @app.route("/api/report-cache/clear", methods=["POST"])
    def api_report_cache_clear():
        """Empty the day cache, so the next open of each report re-runs it.

        A POST, and behind Utilities: pressing it means every QA report and
        every report-shaped tool page runs its whole build again on the next
        visit. Each report also has its own Refresh button, which is the one
        to use for a single report.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import report_cache
        name = (request.args.get("name") or "").strip()
        dropped = report_cache.invalidate(name) if name else report_cache.invalidate()
        audit.log("hub", "report_cache_cleared", actor=current_user(),
                  detail=name or "all", dropped=dropped)
        return jsonify({"ok": True, "dropped": dropped})

    @app.route("/api/integrity")
    def api_integrity():
        """Static audit for defect patterns that have each shipped before."""
        gate = _require_api()
        if gate:
            return gate
        from . import integrity
        return jsonify(integrity.run())

    @app.route("/api/quotas/warnings")
    def api_quota_warnings():
        """Just the providers needing attention — for a banner or a cron job."""
        gate = _require_api()
        if gate:
            return gate
        from . import quotas
        warns = quotas.warnings(request.args.get("month"))
        return jsonify({"warnings": warns, "count": len(warns),
                        "ok": not warns})

    @app.route("/api/client/brand")
    def api_client_brand():
        """Logos, colours and fonts for the Client 360 brand card."""
        gate = _require_api()
        if gate:
            return gate
        from .client_brand import brand_kit
        return jsonify(brand_kit(request.args.get("name", ""),
                                 request.args.get("domain", "")))

    @app.route("/api/client/brand/lookup", methods=["POST"])
    def api_client_brand_lookup():
        """Look a client's brand up live, and keep what the call paid for.

        A POST behind a button, never a page load: the lookup is billed, and
        a GET that spends money is one a reload or a prefetch fires without
        anybody asking — the rule /tools/domains' refresh already works to.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import brand_lookup
        from .client_brand import brand_kit
        body = request.get_json(silent=True) or {}
        client = str(body.get("name") or "").strip()
        domain = str(body.get("domain") or "").strip()
        if not domain:
            return jsonify({"error": "No website on this client to look a "
                                     "brand up by."}), 400
        # use_cache=False: this is the refresh button. Answering it out of the
        # cache would make a press that changes nothing look like a press that
        # did not work.
        res = brand_lookup.lookup(domain, client=client, module="client360",
                                  use_cache=False)
        # A lookup that found nothing still answers with the whole card, not
        # with a bare shell: their own website may have published the logo and
        # the palette all along, and returning `{found: False}` alone would
        # have the refresh button wipe them off the card it was pressed on —
        # the two-blocks failure hub/client_brand._merge closes, coming back
        # in through the one control that redraws the card.
        if not res.get("found"):
            kit = brand_kit(client, domain)
            kit["note"] = res.get("note", "") or kit.get("note", "")
            kit["unconfigured"] = bool(res.get("unconfigured"))
            kit["refused"] = bool(res.get("refused"))
            return jsonify(kit), 200
        kit = brand_kit(client, domain)
        kit["looked_up"] = True
        kit["note"] = res.get("note", "")
        # What the call we just paid for put into the client's gallery. Said
        # on the card rather than done silently: a file that appears in a
        # gallery with nothing announcing it is one nobody knows to look for.
        try:
            from . import client_logos
            kit["logos_filed"] = client_logos.summary(res.get("logos") or {})
        except Exception:  # noqa: BLE001
            pass
        return jsonify(kit)

    @app.route("/api/client/logos", methods=["POST"])
    def api_client_logos():
        """File every logo we already hold for this client into their gallery.

        A POST behind a button, and deliberately **not** a lookup: this reads
        the brand record already stored and the last site scan, so it costs
        nothing at Brandfetch. That matters because the scan is where the logo
        comes from for most local businesses, and there was no way to get one
        into the gallery without spending a lookup that would find nothing.

        Filing the same logo twice is impossible — see hub/client_logos.py —
        so pressing it again is safe and says what was already there.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import client_logos
        body = request.get_json(silent=True) or {}
        client = str(body.get("name") or "").strip()
        if not client:
            return jsonify({"error": "No client named."}), 400
        res = client_logos.file_logos(
            client, str(body.get("domain") or "").strip(),
            actor=current_user() or "system")
        return jsonify({**res, "summary": client_logos.summary(res)})

    @app.route("/api/client/scan-facts")
    def api_client_scan_facts():
        """What the last Insites scan knows about this client, grouped.

        Client 360 read four fields out of a 440-field audit. This is the
        rest of what is worth reading — read-only, no scan is started, and
        nothing here is billed.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import scan_facts
        return jsonify(scan_facts.facts(request.args.get("domain", "")))

    @app.route("/api/client/audit")
    def api_client_audit():
        """The same audit the Website Audit tool shows, for a client record.

        One reading of one audit, on every screen that shows it. The tool at
        `/tools/website-audit`, the prospect record, the customer-facing
        placement and the upsell report all read `hub/website_audit.py`, and
        Client 360 was the last screen still on the thin one — five collapsed
        reference rows about what a business spends, where every other screen
        shows the total, the annualised figure and what is deliberately left
        out of it. Same client, same audit, two answers depending on which
        record you opened.

        **Served from `/api/client/` rather than from the audit tool's own
        blueprint on purpose.** Client 360 is framed inside Smart 1 Suite, and
        `hub/suite_embed.EMBEDDABLE` allowlists `/api/client/` and not
        `/api/website-audit` — so a card pointed at the tool's route would
        render everywhere except inside the frame, which is the half-broken
        embed that file exists to prevent. This is the same function, not a
        second description of it.

        Read-only and unbilled: no scan is started here.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import website_audit
        return jsonify(website_audit.audit(request.args.get("domain", "")))

    @app.route("/api/client/work")
    def api_client_work():
        """Everything the Hub has made for this client, newest first."""
        gate = _require_api()
        if gate:
            return gate
        from .client_brand import work_log
        from . import client_groups
        limit = clamp_int(request.args.get("limit"), 50, 1, 200)
        name = request.args.get("name", "")
        # A grouped client reads across the whole group. Every merged row
        # carries the member it belongs to — see hub/client_groups.py.
        also = client_groups.member_names(name, request.args.get("url", ""))
        return jsonify(work_log(name, limit, also=also))

    @app.route("/api/client/orders")
    def api_client_orders():
        """The insertion orders this Hub has sent for a client.

        Not the campaigns Knack holds — those are the Products & IOs card
        above. This is what we sent, which exists from the day the order goes
        out rather than from the day somebody sets the campaign up, and is the
        only answer there is for a client whose record is otherwise empty.

        Under `/api/client/` deliberately: `hub/suite_embed.EMBEDDABLE`
        allowlists that prefix, so the card works inside the Smart 1 Suite
        frame. A route somewhere else renders on every screen except the one
        the record is framed in, and fails silently there.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import client_groups, io_records
        name = request.args.get("name", "") or request.args.get("client", "")
        # A grouped client reads across the whole group, the way the work log
        # and the invoices do: the bill is one relationship even when the
        # orders were written under two names.
        names = client_groups.member_names(name, request.args.get("url", "")) \
            or [name]
        rows, measured, error = [], True, ""
        seen = set()
        for member in [name] + [n for n in names if n != name]:
            got = io_records.listing(member)
            if not got.get("measured"):
                measured, error = False, got.get("error", "")
                continue
            for row in got["rows"]:
                if row.get("order") in seen:
                    continue
                seen.add(row.get("order"))
                rows.append(dict(row, member=member))
        rows.sort(key=lambda r: str(r.get("last_submitted_at")
                                    or r.get("submitted_at") or ""),
                  reverse=True)
        return jsonify({"orders": rows, "measured": measured, "error": error})

    @app.route("/api/client/brand/push-to-suite", methods=["POST"])
    def api_brand_push():
        """Send the brand guide into the client's Smart 1 Suite sub-account."""
        gate = _require_api()
        if gate:
            return gate
        from .client_brand import brand_guide_payload, mark_pushed
        body = request.get_json(silent=True) or {}
        client = str(body.get("name") or "")
        payload = brand_guide_payload(client, str(body.get("domain") or ""))
        if not payload.get("found"):
            return jsonify({"error": "No brand data on file for that client "
                                     "yet — run a Brandfetch lookup first."}), 400
        target = (os.environ.get("GHL_BRAND_WEBHOOK_URL") or "").strip()
        if not target:
            # Return the payload anyway so it's copy-pasteable. A missing
            # webhook shouldn't mean the work is unavailable. Nothing is
            # recorded: a payload offered for somebody to paste by hand has
            # not reached Suite, and marking it would put a green pill over
            # work nobody has done.
            return jsonify({"ok": False, "delivered": False, "payload": payload,
                            "reason": "not_configured",
                            "note": "Set GHL_BRAND_WEBHOOK_URL to deliver this "
                                    "automatically. The payload above is ready "
                                    "to paste into a Suite workflow meanwhile."})
        # Three outcomes, not two. "The variable is not set", "Suite refused
        # it" and "we could not reach Suite" send somebody to three different
        # places, and the browser used to report all of them as the first --
        # telling a rep to set a variable that is already set, the rule
        # `services/provider_check.py` works to. The refusal carries the
        # status line, because `raise_for_status()` discarding the provider's
        # own sentence is how every button came to report its own invented
        # diagnosis of one shared failure.
        reason, note = "", ""
        try:
            import requests as _rq
            r = _rq.post(target, json=payload, timeout=15)
            ok = r.ok
            if not ok:
                reason = "refused"
                note = f"Smart 1 Suite answered {r.status_code}."
        except Exception as exc:  # noqa: BLE001
            ok = False
            reason = "unreachable"
            note = f"Smart 1 Suite could not be reached ({type(exc).__name__})."
        # Recorded only where it actually landed. Client 360 draws the state
        # from this and hides the button, because "pressing it again just
        # overwrites what's there" -- and until this call existed nothing had
        # ever written the field, so the guard held for the life of one page
        # view and the button came back on every reload. The stamp travels
        # back so the card shows what the next load will read rather than a
        # second idea of when this happened.
        pushed_at = mark_pushed(client) if ok else ""
        audit.log("brand", "pushed_to_suite", actor=current_user(),
                  client=client, ok=ok, reason=reason)
        return jsonify({"ok": ok, "delivered": ok, "payload": payload,
                        "pushed_at": pushed_at, "reason": reason, "note": note})

    @app.route("/api/search")
    def api_search():
        """The top search box: clients first, then the Hub's own pages.

        Read-only and reaches no provider, so it is a GET a keystroke can
        fire. The client half is matched live through `clients_registry`
        rather than from a stored index, because a search that cannot find a
        client we signed last week is one people stop using.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import search_index
        return jsonify(search_index.search(
            request.args.get("q", ""),
            limit=clamp_int(request.args.get("limit"), 12, 1, 30)))

    @app.route("/api/client/context")
    def api_client_context():
        """Merged client record for prefilling any form in the Hub."""
        gate = _require_api()
        if gate:
            return gate
        from .client_context import context
        return jsonify(context(request.args.get("name", ""),
                               request.args.get("domain", "")))

    # NOT /sites/match: DispatcherMiddleware owns the whole /sites prefix and
    # forwards it to the Sites Admin app, so a hub route under it is never
    # reached — it 404s (or 503s when that module is down). Anything the hub
    # app serves has to live outside a mounted prefix.
    @app.route("/tools/sites-match")
    def page_sites_match():
        gate = _require_page()
        if gate:
            return gate
        return render_template("sites_match.html", user=current_user(),
                               active="sites")

    @app.route("/api/sites-match")
    def api_sites_match():
        """Propose a client for every unlinked *live* Simvoly project.

        Read-only. Live only by default: an expired or cancelled project is not
        a website anybody can visit, and its domain has often been repointed,
        so linking a client to one attributes them a site that may now belong
        to somebody else. `?include_inactive=1` shows the rest; the count of
        what is being left out is in the response either way.
        """
        gate = _require_api()
        if gate:
            return gate
        from .sites_match import suggest
        include = str(request.args.get("include_inactive", "")).lower() in (
            "1", "true", "yes")
        return jsonify(suggest(active_only=not include))

    @app.route("/api/client-urls")
    def api_client_urls():
        """Clients with no website, and what the rest of the Hub knows.

        The companion to /api/db/urls, which can say who is missing a URL but
        not do anything about it. This one reads the click-thrus on their live
        products, the Knack website registry, our Simvoly projects, their site
        scans and their Google access requests, and proposes the domain those
        agree on. It writes nothing.
        """
        gate = _require_api()
        if gate:
            return gate
        from .client_urls import missing
        return jsonify(missing(
            include_found=str(request.args.get("include_found", "")).lower()
            in ("1", "true", "yes")))

    @app.route("/api/client-urls/accept", methods=["POST"])
    def api_client_urls_accept():
        """Record one URL a human recognised — everywhere it belongs.

        This used to write the Hub's overlay and nothing else, so a rep who
        accepted a domain here found Client 360 still saying the client had no
        website. It goes through `domain_links.attach()` now, which writes the
        overlay, the client's 360 record, the Simvoly project and the Knack
        website record, and reports each one separately.
        """
        gate = _require_api()
        if gate:
            return gate
        from .domain_links import attach
        body = request.get_json(silent=True) or {}
        out = attach(str(body.get("domain") or body.get("url") or ""),
                     str(body.get("client") or ""),
                     actor=current_user() or "",
                     url=str(body.get("url") or ""),
                     force=bool(body.get("force")))
        if out.get("ok"):
            audit.log("hub", "client_url_accepted", actor=current_user(),
                      client=body.get("client", ""), detail=out["domain"])
        return jsonify(out)

    @app.route("/api/client-urls/accept-many", methods=["POST"])
    def api_client_urls_accept_many():
        """Accept several at once, each with its own outcome.

        Reviewing thirty proposals and clicking thirty confirm dialogs is how
        a reviewer stops reading them. Each result carries its own report — a
        bulk action that returns one number hides the two that failed.
        """
        gate = _require_api()
        if gate:
            return gate
        from .domain_links import attach_many
        body = request.get_json(silent=True) or {}
        links = body.get("links") or body.get("matches") or []
        return jsonify(attach_many(links, actor=current_user() or "",
                                   force=bool(body.get("force"))))

    @app.route("/api/orphan-urls")
    def api_orphan_urls():
        """Every URL the Hub holds that no client is attached to.

        The other direction from /api/client-urls: not "this client has no
        website" but "this website has no client". Same four systems, read
        rather than written, and a source that could not be read is named.
        """
        gate = _require_api()
        if gate:
            return gate
        from .domain_links import orphans
        return jsonify(orphans(q=request.args.get("q", ""),
                               limit=clamp_int(request.args.get("limit"),
                                               400, 1, 2000)))

    @app.route("/api/domain/attach", methods=["POST"])
    def api_domain_attach():
        """Attach one orphan URL, or several, to a client."""
        gate = _require_api()
        if gate:
            return gate
        body = request.get_json(silent=True) or {}
        actor = current_user() or ""
        links = body.get("links")
        if links:
            from .domain_links import attach_many
            return jsonify(attach_many(links, actor=actor,
                                       force=bool(body.get("force"))))
        from .domain_links import attach
        return jsonify(attach(str(body.get("domain") or ""),
                              str(body.get("client") or ""), actor=actor,
                              url=str(body.get("url") or ""),
                              force=bool(body.get("force"))))

    @app.route("/api/client-urls/clear", methods=["POST"])
    def api_client_urls_clear():
        """Undo one. A wrong URL has to be removable without a deploy.

        With a `domain` this removes that one site and leaves the client's
        others alone; without one it clears every site accepted for them.
        """
        gate = _require_api()
        if gate:
            return gate
        from .client_urls import clear
        body = request.get_json(silent=True) or {}
        out = clear(body.get("client", ""), str(body.get("domain") or ""))
        if out.get("ok"):
            audit.log("hub", "client_url_cleared", actor=current_user(),
                      client=body.get("client", ""))
        return jsonify(out)

    @app.route("/api/sites-match/read-names", methods=["POST"])
    def api_sites_match_read_names():
        """Read the unmatched project titles for the business inside them.

        A POST, and a button, because the call is billed: a GET that spends
        money is one a reload or a link preview fires without anybody asking —
        the rule `hub/domain_purchase.py` settled for the domain calendar and
        `hub/brand_lookup.py` for the brand card.

        The model is shown project titles and nothing else. It never sees the
        client book, so it cannot name a client; what it returns is a run of
        words out of the title, which then goes through
        `site_names.exact_matches()` against the real book exactly like a
        candidate a rule derived. Reading changes no match on its own.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import site_names_ai
        from .sites_match import suggest
        body = request.get_json(silent=True) or {}
        include = bool(body.get("include_inactive"))
        data = suggest(active_only=not include)
        titles = [r.get("site") or "" for key in
                  ("unmatched", "no_domain", "suggested")
                  for r in (data.get(key) or [])]
        out = site_names_ai.read_missing(titles)
        # The report is held for the day, and a reading changes what it would
        # say — so it is dropped here rather than leaving the page showing the
        # answer from before the button was pressed, which reads as a button
        # that did nothing. The drop lives beside the write, per
        # hub/report_cache.py.
        try:
            from . import report_cache
            report_cache.invalidate("sites-match")
        except Exception:                               # noqa: BLE001
            pass
        try:
            audit.log("sites_match", "read_names", actor=current_user() or "",
                      read=out.get("read", 0), stored=out.get("stored", 0),
                      ungrounded=out.get("ungrounded", 0))
        except Exception:                               # noqa: BLE001
            pass
        return jsonify(out)

    @app.route("/api/sites-match/apply", methods=["POST"])
    def api_sites_match_apply():
        """Write only the matches a human accepted."""
        gate = _require_api()
        if gate:
            return gate
        from .sites_match import apply as apply_matches
        body = request.get_json(silent=True) or {}
        return jsonify(apply_matches(body.get("matches") or [],
                                     actor=current_user() or ""))

    # /tools/domains, like /tools/sites-match, is a hub route under /tools —
    # which is not itself a mount. The mounts are longer prefixes
    # (/tools/io, /tools/social...), so this one reaches the hub app.
    @app.route("/tools/domains")
    def page_domain_purchase():
        gate = _require_page()
        if gate:
            return gate
        return render_template("domain_purchase.html", user=current_user(),
                               active="domains")

    # /tools/campaign-assets is a hub route under /tools, like /tools/domains
    # — the mounts are all longer prefixes, so this one reaches the hub app.
    @app.route("/tools/campaign-assets")
    def page_campaign_assets():
        gate = _require_page()
        if gate:
            return gate
        from .knack_products import (F_ASSETS_FLAG, F_ASSETS_NEEDED,
                                     F_CLARIFICATION)
        return render_template(
            "campaign_assets.html", user=current_user(),
            active="campaign-assets",
            # Handed to the page as data rather than written into its script,
            # so a renumbered field is one environment variable and not an
            # edit in two files that can disagree.
            field_ids={"clarification": F_CLARIFICATION,
                       "assets_flag": F_ASSETS_FLAG,
                       "assets_needed": F_ASSETS_NEEDED})

    @app.route("/api/campaign-assets")
    def api_campaign_assets():
        """Campaigns waiting on an asset or a clarification (object_135)."""
        gate = _require_api()
        if gate:
            return gate
        from .campaign_assets import report
        return jsonify(report(q=request.args.get("q", ""),
                              scope=request.args.get("scope", "open")))

    @app.route("/api/campaign-assets/fields")
    def api_campaign_assets_fields():
        """The three pinned ids against object_135's live schema.

        A renumbered field reads back empty on every record and looks exactly
        like a client base with nothing outstanding, so what Knack calls each
        id is answerable from the page rather than only from a diagnostic.
        """
        gate = _require_api()
        if gate:
            return gate
        from .campaign_assets import field_check
        return jsonify(field_check())

    @app.route("/api/domains/purchased")
    def api_domains_purchased():
        """Domains Smart 1 bought for a client, by renewal billing date.

        Served from the nightly snapshot of object_153 and from the cached
        QuickBooks renewal charges, not from either provider. The report says
        how old each is; forcing a fresh pull is the POST below, so a page
        load can never be the thing that pulls.
        """
        gate = _require_api()
        if gate:
            return gate
        from .domain_purchase import report
        return jsonify(report(q=request.args.get("q", "")))

    @app.route("/api/domains/do-not-renew", methods=["POST"])
    def api_domains_do_not_renew():
        """Mark one domain as not to be renewed, or clear the mark.

        Knack publishes no such field either, so like the billed tick this is
        the Hub's own — but unlike the billed tick it is never retired when
        the renewal date rolls: a domain that renewed after we said it should
        not is a finding, and clearing the mark would delete the evidence.
        """
        gate = _require_api()
        if gate:
            return gate
        from .domain_purchase import set_do_not_renew
        body = request.get_json(silent=True) or {}
        out = set_do_not_renew(str(body.get("record_id") or ""),
                               bool(body.get("do_not_renew")),
                               for_date=str(body.get("for_date") or ""),
                               reason=str(body.get("reason") or ""),
                               actor=current_user() or "")
        if out.get("ok"):
            audit.log("hub", "domain_do_not_renew", actor=current_user(),
                      detail=f"{body.get('record_id')} "
                             f"{'marked' if body.get('do_not_renew') else 'cleared'}")
        return jsonify(out)

    @app.route("/api/domains/do-not-renew")
    def api_domains_do_not_renew_report():
        """Everything marked do-not-renew, and what renewed anyway."""
        gate = _require_api()
        if gate:
            return gate
        from .domain_purchase import do_not_renew_report
        return jsonify(do_not_renew_report())

    @app.route("/api/domains/ytd")
    def api_domains_ytd():
        """Year to date, both directions: unbilled renewals and unknown charges.

        The two questions nobody could ask before QuickBooks was joined to the
        registry — a renewal that came due with no invoice behind it, and a
        Website Domain Renewal charge that matches no record here.
        """
        gate = _require_api()
        if gate:
            return gate
        from .domain_purchase import year_to_date
        try:
            year = int(request.args.get("year") or 0) or None
        except ValueError:
            year = None
        # No refresh here either: this reads the same two caches the calendar
        # does, and POST /api/domains/refresh is the one control that pulls.
        return jsonify(year_to_date(year=year))

    @app.route("/api/domains/read-descriptions", methods=["POST"])
    def api_domains_read_descriptions():
        """Read the line descriptions every matching rule failed on.

        A POST, and a button, because the call is billed: a GET that spends
        money is one a reload or a link preview fires without anybody asking —
        the rule this same page settled for its Refresh.

        Only the unmatched lines are sent. What comes back is a business name
        grounded in the description it came from, which then goes through the
        matcher's own name passes against the real registry — so it can move a
        charge from "nothing to look at" to "here is a candidate", and it
        resolves to `probable`, which `year_to_date()` counts as having no
        record here until somebody presses Link. Nothing here marks a renewal
        billed.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import invoice_names
        from .domain_purchase import year_to_date
        body = request.get_json(silent=True) or {}
        try:
            year = int(body.get("year") or 0) or None
        except (TypeError, ValueError):
            year = None
        data = year_to_date(year=year)
        # The same set `ai_names` counted: unmatched, or matched only on a
        # resemblance. A line the rules answered is never sent.
        wanted = [c.get("description") or ""
                  for c in (data.get("unrecorded") or [])]
        out = invoice_names.read_missing(wanted)
        try:
            audit.log("domain_renewals", "read_descriptions",
                      actor=current_user() or "",
                      read=out.get("read", 0), stored=out.get("stored", 0),
                      ungrounded=out.get("ungrounded", 0))
        except Exception:                               # noqa: BLE001
            pass
        return jsonify(out)

    @app.route("/api/domains/charges/link", methods=["POST"])
    def api_domains_charge_link():
        """Attach one QuickBooks charge to one website record, or clear it.

        A person saying "this charge is that domain" is the only fact in the
        matcher — it outranks every rule and survives the next read of
        QuickBooks, so a description nothing could join up is joined once
        rather than every month.
        """
        gate = _require_api()
        if gate:
            return gate
        from .domain_renewals import link_charge
        body = request.get_json(silent=True) or {}
        out = link_charge(str(body.get("key") or ""),
                          str(body.get("record_id") or ""),
                          actor=current_user() or "",
                          domain=str(body.get("domain") or ""))
        if out.get("ok"):
            audit.log("hub", "domain_charge_linked", actor=current_user(),
                      detail=f"{body.get('key')} -> "
                             f"{body.get('record_id') or '(cleared)'}")
        return jsonify(out)

    @app.route("/api/domains/refresh", methods=["POST"])
    def api_domains_refresh():
        """Pull object_153 now, for the Refresh button.

        POST rather than a `?refresh=1` on the read: this reaches Knack and
        rewrites the stored snapshot, and a GET that does that is one a
        prefetch or a reload can fire without anybody asking for it.
        """
        gate = _require_api()
        if gate:
            return gate
        from datetime import date as _date

        from .domain_purchase import refresh, report
        out = refresh(force=True)
        audit.log("hub", "domains_refresh", actor=current_user(),
                  ok=bool(out.get("ok")), count=out.get("count"),
                  error=out.get("error") or None)
        # Both halves, through the one button. Billed is read from QuickBooks
        # and cached exactly as the registry is, and a Refresh that pulled one
        # of the two would leave the page reporting a fresh timestamp over a
        # stale answer to the question it is actually asked.
        qb = {}
        try:
            from .domain_renewals import charges
            qb = charges(_date.today().year, refresh=True)
            qb = {"error": qb.get("error", ""), "count": len(qb.get("lines") or []),
                  "fetched": qb.get("fetched_at", "")}
        except Exception as exc:                        # noqa: BLE001
            qb = {"error": f"{type(exc).__name__}: {exc}", "count": 0,
                  "fetched": ""}
        # The report comes back with it, so the button is one round trip and
        # the page cannot end up showing a fresh timestamp over old rows.
        return jsonify({**report(q=request.args.get("q", ""), build=False),
                        "refresh": out, "quickbooks_refresh": qb})

    @app.route("/api/domains/billed", methods=["POST"])
    def api_domains_billed():
        """Tick or untick one renewal as billed.

        Knack publishes no billed field, so this is the Hub's own — kept
        against the renewal billing date it was ticked for, so next year's
        renewal does not inherit this year's tick.
        """
        gate = _require_api()
        if gate:
            return gate
        from .domain_purchase import set_billed
        body = request.get_json(silent=True) or {}
        out = set_billed(str(body.get("record_id") or ""),
                         bool(body.get("billed")),
                         for_date=str(body.get("for_date") or ""),
                         actor=current_user() or "")
        if out.get("ok"):
            audit.log("hub", "domain_billed", actor=current_user(),
                      detail=f"{body.get('record_id')} "
                             f"{'billed' if body.get('billed') else 'unbilled'}")
        return jsonify(out)

    @app.route("/api/db/urls")
    def api_db_urls():
        """Clients with no usable URL, and one domain filed under two names."""
        gate = _require_api()
        if gate:
            return gate
        from .client_context import url_audit
        return jsonify(url_audit())

    @app.route("/api/client/by-url")
    def api_client_by_url():
        """Resolve a client from a URL, whatever the name is filed as."""
        gate = _require_api()
        if gate:
            return gate
        from .client_context import resolve_by_url
        return jsonify(resolve_by_url(request.args.get("url", "")))

    @app.route("/api/clients/crosswalk")
    def api_clients_crosswalk():
        """Every client record in every module, grouped by one derived key.

        The join that is missing from the data. Answers the two questions no
        single module can: which records are the same client, and which ones
        carry a name with no URL behind it and therefore cannot be joined to
        anything at all.
        """
        gate = _require_api()
        if gate:
            return gate
        from .client_key import crosswalk
        from . import report_cache
        # Every module's client table, read and joined. Held for the day; the
        # accept/attach paths that change what joins to what drop it.
        return jsonify(report_cache.serve("crosswalk", crosswalk))

    @app.route("/api/client/records")
    def api_client_records():
        """Every module record belonging to one client, by name or by URL."""
        gate = _require_api()
        if gate:
            return gate
        from .client_key import for_client
        return jsonify(for_client(name=request.args.get("client", ""),
                                  url=request.args.get("url", "")))

    @app.route("/api/client/resolve")
    def api_client_resolve():
        """Who a name or URL actually is, and how confidently we know."""
        gate = _require_api()
        if gate:
            return gate
        from .client_key import resolve
        fuzzy = str(request.args.get("fuzzy", "")).lower() in ("1", "true", "yes")
        return jsonify(resolve(name=request.args.get("client", ""),
                               url=request.args.get("url", ""),
                               allow_fuzzy=fuzzy))

    @app.route("/api/db/structure")
    def api_db_structure():
        """Where client data lives, and where it can drift apart."""
        gate = _require_api()
        if gate:
            return gate
        from .client_context import structure_report
        return jsonify(structure_report())

    @app.route("/api/qb/health")
    def api_qb_health():
        """Why the QuickBooks connection may not be holding."""
        gate = _require_api()
        if gate:
            return gate
        try:
            from . import quickbooks as qb
            return jsonify(qb.health())
        except Exception as exc:  # noqa: BLE001
            return jsonify({"ok": False, "problems": [str(exc)]}), 200

    @app.route("/api/seo/llms-txt")
    def api_llms_txt_build():
        """Draft an llms.txt for a client from what the Hub already knows."""
        gate = _require_api()
        if gate:
            return gate
        from .llms_txt import build, load
        client = request.args.get("client", "")
        if request.args.get("saved") == "1":
            return jsonify({"client": client, "text": load(client)})
        return jsonify(build(client))

    @app.route("/api/seo/llms-txt", methods=["POST"])
    def api_llms_txt_save():
        gate = _require_api()
        if gate:
            return gate
        from .llms_txt import save
        body = request.get_json(silent=True) or {}
        client = str(body.get("client") or "")
        text = str(body.get("text") or "")
        if not client or not text.strip():
            return jsonify({"error": "Client and text are both required."}), 400
        if "NEED " in text:
            return jsonify({"error": "This still contains NEED placeholders. "
                                     "Fill them in first — a file with gaps is "
                                     "worse than none, because a model treats "
                                     "the whole thing as authoritative."}), 400
        return jsonify(save(client, text))

    @app.route("/llms/<slug>.txt")
    def public_llms_txt(slug):
        """Serve the approved file publicly, as plain text.

        Deliberately no login: the point is that an AI system can fetch it.
        Ideally this lives at the client's own domain root as /llms.txt — this
        URL is what you use until it can.
        """
        import re as _re
        from . import seo
        from .llms_txt import load

        def slugify(v):
            return _re.sub(r"[^a-z0-9]+", "-", str(v or "").lower()).strip("-")

        want = slugify(slug)
        try:
            from . import clients_registry
            names = [c.get("name", "") for c in clients_registry.all_clients()]
        except Exception:  # noqa: BLE001
            names = []
        for name in names:
            if name and slugify(name) == want:
                text = load(name)
                if text:
                    return app.response_class(
                        text, mimetype="text/plain; charset=utf-8")
        return app.response_class("Not found.\n", status=404,
                                  mimetype="text/plain; charset=utf-8")

    @app.route("/api/suite/blog/access")
    def api_blog_access():
        """Which blogs scopes the Suite token actually has."""
        gate = _require_api()
        if gate:
            return gate
        from .ghl_blog import check_access, BlogError
        try:
            return jsonify(check_access())
        except BlogError as exc:
            return jsonify({"ok": False, "problem": str(exc)}), 200

    @app.route("/api/suite/blog/publish-llms", methods=["POST"])
    def api_blog_publish_llms():
        """Publish a client's llms.txt to Suite as a blog post."""
        gate = _require_api()
        if gate:
            return gate
        from .ghl_blog import publish_llms_txt, BlogError
        from .llms_txt import load
        body = request.get_json(silent=True) or {}
        client = str(body.get("client") or "")
        text = str(body.get("text") or "") or load(client)
        if not client or not text.strip():
            return jsonify({"error": "Save the file before publishing it."}), 400
        if "NEED " in text:
            return jsonify({"error": "This still has NEED placeholders — fill "
                                     "them in before publishing."}), 400
        try:
            out = publish_llms_txt(client, text,
                                   post_id=str(body.get("post_id") or ""),
                                   status=str(body.get("status") or "PUBLISHED"))
        except BlogError as exc:
            return jsonify({"error": str(exc)}), 400
        # Remember the URL so Client 360 can link to it.
        try:
            from . import seo
            store = seo.load_store(client) or {}
            rec = store.get("llms_txt") or {}
            rec.update({"suite_url": out.get("url"), "post_id": out.get("post_id")})
            store["llms_txt"] = rec
            seo.save_store(client, store)
        except Exception:  # noqa: BLE001
            pass
        return jsonify(out)

    @app.route("/api/client/website-registry")
    def api_website_registry():
        """GA, GTM, platform, go-live and H&M fee from Knack object_153."""
        gate = _require_api()
        if gate:
            return gate
        from .knack_websites import enrich
        return jsonify(enrich(request.args.get("name", ""),
                              request.args.get("domain", "")))

    @app.route("/api/client/website-record")
    def api_website_record():
        """The domain record on object_153: live date, status, the domain.

        Client 360 draws this per website. The *choices* come from the live
        schema for the reason the web ticket form's do: the field ids are ours
        but a dropdown's options are Knack's, and Knack refuses the whole
        record over one value it does not publish.
        """
        gate = _require_api()
        if gate:
            return gate
        from .knack_websites import domain_record
        out = domain_record(request.args.get("name", ""),
                            request.args.get("domain", ""))
        # The renewal standing rides along with the record rather than being a
        # second fetch: a domain attached to a client is a domain somebody on
        # Client 360 wants the renewal answer for, and sending them to
        # /tools/domains to find the same row again is how a list stays
        # unactioned. Never raises — the domain record is the point of this
        # route and a QuickBooks that will not answer must not cost it.
        try:
            from .domain_purchase import status_for_record
            out["renewal"] = status_for_record(out.get("record_id", ""))
        except Exception as exc:                        # noqa: BLE001
            out["renewal"] = {"applies": False,
                              "reason": f"The renewal standing could not be "
                                        f"read ({type(exc).__name__})."}
        return jsonify(out)

    @app.route("/api/client/website-record/save", methods=["POST"])
    def api_website_record_save():
        """Write the domain record back to Knack. Says what was refused."""
        gate = _require_api()
        if gate:
            return gate
        from .knack_websites import update_record
        body = request.get_json(silent=True) or {}
        out = update_record(str(body.get("record_id") or ""),
                            body.get("values") or {},
                            actor=current_user() or "")
        if out.get("ok"):
            audit.log("hub", "website_record_saved", actor=current_user(),
                      client=str(body.get("client") or ""),
                      detail=",".join(out.get("updated") or [])[:200])
        return jsonify(out)

    @app.route("/api/client/analytics-ids")
    def api_analytics_ids():
        """GA and GTM from BOTH Knack and Google, with whether they agree."""
        gate = _require_api()
        if gate:
            return gate
        from .analytics_ids import compare
        return jsonify(compare(request.args.get("name", ""),
                               request.args.get("domain", "")))

    @app.route("/api/qa/analytics-ids")
    def api_analytics_audit():
        """Every client where the two sources disagree, or we lack access."""
        gate = _require_api()
        if gate:
            return gate
        from .analytics_ids import audit_all
        from . import report_cache
        # One Knack website record and one Google index lookup per client.
        # Dropped by `google_index._forget_reports()` when a property is
        # attached, which is the write that changes what this compares.
        return jsonify(report_cache.serve("qa:analytics-ids", audit_all))

    @app.route("/api/scans/stuck")
    def api_scans_stuck():
        """How many scans are stuck, and how long they've been there.

        Surfaced on Diagnostics because a stalled scan is an operational
        problem, not something you'd think to go looking for on the Scans
        page — it looks like work in progress until somebody counts.
        """
        gate = _require_api()
        if gate:
            return gate
        try:
            from modules.scans.app import Scan, SessionLocal
            from datetime import datetime, timedelta, timezone
            db = SessionLocal()
            try:
                rows = db.query(Scan).filter(Scan.status == "running").all()
                now = datetime.now(timezone.utc)
                buckets = {"under_15m": 0, "15m_to_1h": 0, "over_1h": 0,
                           "unresolvable": 0}
                oldest = None
                for r in rows:
                    c = r.created_at
                    if c is not None and c.tzinfo is None:
                        c = c.replace(tzinfo=timezone.utc)
                    age = (now - c).total_seconds() / 60 if c else 0
                    oldest = max(oldest or 0, age)
                    if not r.insites_report_id and age > 30:
                        buckets["unresolvable"] += 1
                    elif age < 15:
                        buckets["under_15m"] += 1
                    elif age < 60:
                        buckets["15m_to_1h"] += 1
                    else:
                        buckets["over_1h"] += 1
            finally:
                db.close()
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": f"Scans unavailable ({type(exc).__name__})."}), 200
        total = sum(buckets.values())
        return jsonify({
            "running": total, "buckets": buckets,
            "oldest_minutes": round(oldest) if oldest else 0,
            "state": ("error" if buckets["unresolvable"] else
                      "warn" if buckets["over_1h"] else "ok"),
            "advice": ("Scans with no Insites report id can never resolve — "
                       "they were started before the callback fix. Clear them "
                       "and re-run."
                       if buckets["unresolvable"] else
                       "Some scans have been running over an hour. Insites "
                       "audits normally take one to four minutes."
                       if buckets["over_1h"] else
                       "Nothing stalled."),
        })

    @app.route("/api/scheduler")
    def api_scheduler():
        """What the background jobs are doing."""
        gate = _require_api()
        if gate:
            return gate
        from . import scheduler as _sched
        out = _sched.status(app)
        out["boot_error"] = app.config.get("HUB_SCHEDULER_BOOT_ERROR")
        return jsonify(out)

    @app.route("/api/seo/schema-questions")
    def api_schema_questions():
        """The full question set, answered where possible."""
        gate = _require_api()
        if gate:
            return gate
        from .schema_questions import build
        return jsonify(build(request.args.get("client", ""),
                             use_ai=request.args.get("ai", "1") != "0"))

    @app.route("/api/seo/schema-questions/regenerate", methods=["POST"])
    def api_schema_regenerate():
        """Ask AI again for one question — the New AI button."""
        gate = _require_api()
        if gate:
            return gate
        from .schema_questions import regenerate_one
        body = request.get_json(silent=True) or {}
        return jsonify(regenerate_one(str(body.get("client") or ""),
                                      str(body.get("key") or "")))

    @app.route("/api/seo/schema-questions", methods=["POST"])
    def api_schema_answers_save():
        """Save approved and edited answers."""
        gate = _require_api()
        if gate:
            return gate
        from .schema_questions import save_answers, can_approve
        body = request.get_json(silent=True) or {}
        client = str(body.get("client") or "")
        if not client:
            return jsonify({"error": "No client given."}), 400
        out = save_answers(client, body.get("answers") or {}, current_user() or "")
        out.update(can_approve(client))
        return jsonify(out)

    @app.route("/api/client/utm")
    def api_client_utm():
        """Tracked links built for this client.

        Read from the UTM Builder's own store rather than duplicating them —
        two copies of a link is how one ends up stale and the other gets used.
        """
        gate = _require_api()
        if gate:
            return gate
        name = (request.args.get("name") or "").strip().lower()
        try:
            from modules.utm_builder.app import load_links
            rows = [r for r in (load_links() or [])
                    if str(r.get("client") or "").strip().lower() == name]
        except Exception as exc:  # noqa: BLE001
            return jsonify({"links": [], "error": f"{type(exc).__name__}"}), 200
        rows.sort(key=lambda r: str(r.get("created") or ""), reverse=True)
        return jsonify({"client": name, "count": len(rows), "links": rows[:40]})

    @app.route("/api/seo/blogs/image", methods=["POST"])
    def api_blog_image():
        """Generate, approve or delete a post's featured image."""
        gate = _require_api()
        if gate:
            return gate
        from . import blog_images as BI
        body = request.get_json(silent=True) or {}
        client = str(body.get("client") or "")
        pid = body.get("id")
        action = str(body.get("action") or "generate")
        actor = current_user() or ""
        if not client or pid is None:
            return jsonify({"error": "Client and post id are both required."}), 400
        if action == "approve":
            return jsonify(BI.approve(client, pid, actor))
        if action in ("delete", "reject"):
            return jsonify(BI.reject(client, pid, actor))
        return jsonify(BI.generate(client, pid, str(body.get("extra") or ""), actor))

    @app.route("/api/seo/blogs/image-status")
    def api_blog_image_status():
        gate = _require_api()
        if gate:
            return gate
        from . import blog_images as BI
        return jsonify(BI.status(request.args.get("client", "")))

    @app.route("/api/knack/products")
    def api_knack_products():
        """Live IO products for a client, with how fresh the data is."""
        gate = _require_api()
        if gate:
            return gate
        from .knack_products import for_client, status
        name = request.args.get("name", "")
        if not name:
            return jsonify(status())
        return jsonify(for_client(name))

    @app.route("/api/knack/products/refresh", methods=["POST"])
    def api_knack_products_refresh():
        gate = _require_api()
        if gate:
            return gate
        from .knack_products import refresh
        return jsonify(refresh())

    @app.route("/api/client/forms")
    def api_client_forms():
        """Suite forms with submissions, against the previous period."""
        gate = _require_api()
        if gate:
            return gate
        from .ghl_forms import summary
        return jsonify(summary(request.args.get("name", ""),
                               request.args.get("location", ""),
                               request.args.get("period", "this_month")))

    @app.route("/creative")
    def page_creative():
        """Creative tools, mirroring the Tools index."""
        gate = _require_page()
        if gate:
            return gate
        return render_template("creative.html", user=current_user(),
                               active="tools")

    @app.route("/api/scans/click-thru-domains")
    def api_click_thru_domains():
        """Root domains taken from click-thru URLs on live products.

        These are clients the bulk scanner used to skip for having no website
        on file, while their live campaigns pointed at one the whole time.
        """
        gate = _require_api()
        if gate:
            return gate
        from .knack_products import scan_domains
        return jsonify(scan_domains(request.args.get("client", "")))

    @app.route("/api/io/prefill")
    def api_io_prefill():
        """Start an IO from a client, their last IO, or a proposal."""
        gate = _require_api()
        if gate:
            return gate
        from . import io_prefill
        client = request.args.get("client", "")
        mode = request.args.get("mode", "new")
        if mode == "renewal":
            return jsonify(io_prefill.from_last_io(client))
        if mode == "creative":
            return jsonify(io_prefill.creative_for(client))
        if mode == "existing":
            # What a New IO for this client might be replacing. Asked before
            # the interview builds a second order that bills beside the first.
            return jsonify(io_prefill.open_ios(client))
        return jsonify(io_prefill.from_client(client))

    @app.route("/api/io/from-proposal", methods=["POST"])
    def api_io_from_proposal():
        """Read a proposal — uploaded now, or already on the client — for an IO.

        Works without a client, so a prospect's proposal can start an IO
        before they exist in the system.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import io_prefill
        client = (request.form.get("client") or "").strip()
        text, name = "", ""
        up = request.files.get("file")
        if up and up.filename:
            name = up.filename
            raw = up.read(8 * 1024 * 1024)
            text = _read_document(raw, name)
        else:
            body = request.get_json(silent=True) or {}
            client = client or str(body.get("client") or "")
            name = str(body.get("filename") or "")
            text = str(body.get("text") or "")
            if not text and client and name:
                text = _proposal_text_for(client, name)
        if not text.strip():
            return jsonify({"error": "Couldn't read any text from that "
                                     "proposal. If it's a scanned PDF there's "
                                     "no text layer to read."}), 400
        result = io_prefill.from_proposal(client, text, name)

        # Every product the reader found, classified against the rate card and
        # carrying the question it still needs answered. The conversion flow
        # used to receive the same names as a sentence -- "not on the rate
        # card, so not selected" -- and a product the client agreed to would
        # quietly fail to reach the document that bills for it.
        try:
            from . import product_intake
            months = 0
            try:
                months = int(float(str((result.get("fields") or {})
                                       .get("term_months") or 0)))
            except (TypeError, ValueError):
                months = 0
            rows = product_intake.read_products(
                (result.get("fields") or {}).get("products_detail") or [],
                months=months or 1)
            # The card matches what it can by name; the model is asked about
            # the rest. It suggests and never decides -- see the rules at the
            # foot of hub/product_intake.py. A model that is off or slow costs
            # the ordering of a candidate list and nothing else.
            rows = product_intake.ai_match(rows)
            result["product_intake"] = rows
            result["intake_summary"] = product_intake.summary(rows)
            result["consulting_product"] = product_intake.CONSULTING["product"]
        except Exception as exc:                            # noqa: BLE001
            # A reader that works and an intake that does not is still worth
            # returning -- but never silently. CLAUDE.md: guard boot-time
            # failures, and record them.
            app.logger.warning("product intake failed: %s", exc)
            result["product_intake_error"] = str(exc)

        # Flights the document disagrees with itself about, as questions the
        # interview asks up front rather than as a note nobody reads.
        try:
            result["flight_questions"] = io_prefill.flight_questions(
                result.get("fields") or {})
        except Exception as exc:                            # noqa: BLE001
            app.logger.warning("flight questions failed: %s", exc)
            result["flight_questions"] = []
        return jsonify(result)

    @app.route("/api/products/classify", methods=["POST"])
    def api_products_classify():
        """Match free-text product names against the rate card.

        Both builders ask this: the IO while converting a proposal it just
        read, the Proposal Builder when a rep types a product that is not a
        catalogue pick. One matcher, so the two cannot disagree about what a
        client was sold.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import product_intake
        body = request.get_json(silent=True) or {}
        names = body.get("products") or body.get("names") or []
        if isinstance(names, str):
            names = [names]
        months = body.get("months") or 1
        try:
            months = max(1, int(months))
        except (TypeError, ValueError):
            months = 1
        rows = product_intake.read_products(names, months=months)
        return jsonify({"ok": True, "products": rows,
                        "summary": product_intake.summary(rows),
                        "consulting": product_intake.CONSULTING,
                        "basis_labels": product_intake.BASIS_LABEL})

    @app.route("/api/spec/<source>")
    def api_spec(source):
        """A campaign spec from a client, their last IO, or a proposal.

        One shape shared by the Proposal Builder and the IO Builder, so a
        proposal converts by loading rather than by retyping.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import campaign_spec as CS
        client = request.args.get("client", "")
        if source == "last-io":
            spec = CS.from_last_io(client)
            if not spec:
                return jsonify({"error": "No previous IO for that client."}), 404
        elif source == "proposal":
            spec = CS.from_proposal_text(request.args.get("text", ""), client)
        else:
            spec = CS.from_client(client)
        out = spec.to_dict()
        out["ready_for_io"] = spec.ready_for_io()
        return jsonify(out)

    @app.route("/api/spec/to-io", methods=["POST"])
    def api_spec_to_io():
        """Convert a spec into what the IO intake consumes."""
        gate = _require_api()
        if gate:
            return gate
        from . import campaign_spec as CS
        body = request.get_json(silent=True) or {}
        spec = CS.CampaignSpec.from_dict(body.get("spec") or {})
        return jsonify(CS.to_io_payload(spec))

    @app.route("/api/client/suite-match")
    def api_suite_match():
        """Suite sub-accounts that look like this client.

        The card offered "Add to Smart 1 Suite" whether or not an account
        already existed under a slightly different name — which is how a
        client ends up with two sub-accounts and their history split across
        both. Search first, offer to attach, and only offer to create when
        nothing plausible comes back.
        """
        gate = _require_api()
        if gate:
            return gate
        name = (request.args.get("name") or "").strip()
        if not name:
            return jsonify({"matches": [], "searched": ""})
        import re as _re

        def norm(v):
            v = _re.sub(r"\b(llc|inc|ltd|co|corp|company|the|dba)\b", " ",
                        str(v or "").lower())
            return _re.sub(r"[^a-z0-9]+", "", v)

        want = norm(name)
        # Try the distinctive part too — "Icon Solar Power, LLC" should find
        # a sub-account called just "Icon Solar".
        terms = [name] + ([" ".join(name.split()[:2])] if len(name.split()) > 2 else [])
        seen, matches = set(), []
        try:
            from modules.suite_panel.app import ghl, _env
            for term in terms:
                data = ghl("/locations/search",
                           query={"companyId": _env("GHL_COMPANY_ID"),
                                  "limit": "10", "query": term}) or {}
                for loc in (data.get("locations") or []):
                    lid = loc.get("id") or loc.get("_id")
                    if not lid or lid in seen:
                        continue
                    seen.add(lid)
                    ln = loc.get("name") or ""
                    n = norm(ln)
                    exact = n == want
                    close = bool(n) and (n in want or want in n)
                    if exact or close:
                        matches.append({
                            "id": lid, "name": ln,
                            "website": loc.get("website") or "",
                            "confidence": "exact" if exact else "close",
                            "why": ("Name matches once LLC/Inc are ignored."
                                    if exact else
                                    f'"{ln}" looks like a variation of this client.'),
                        })
        except Exception as exc:  # noqa: BLE001
            return jsonify({"matches": [], "error": f"{type(exc).__name__}",
                            "searched": name}), 200
        matches.sort(key=lambda m: 0 if m["confidence"] == "exact" else 1)
        return jsonify({
            "searched": name, "matches": matches, "count": len(matches),
            "note": ("Attach one of these rather than creating a second "
                     "sub-account — a duplicate splits the client's history."
                     if matches else
                     "Nothing in Suite looks like this client. Search the full "
                     "list before creating one."),
        })

    @app.route("/sales/leads")
    def page_leads():
        """One panel for every lead, whatever produced it."""
        gate = _require_page()
        if gate:
            return gate
        return render_template("leads.html", user=current_user(), active="leads")

    @app.route("/api/leads")
    def api_leads():
        gate = _require_api()
        if gate:
            return gate
        from . import leads
        return jsonify(leads.listing(
            days=clamp_int(request.args.get("days"), 30, 1, 730),
            source=request.args.get("source", ""),
            page=request.args.get("page", ""),
            undelivered_only=request.args.get("undelivered") == "1"))

    @app.route("/api/leads/capture", methods=["POST"])
    def api_leads_capture():
        """Where every landing page and calculator posts.

        Unauthenticated on purpose — these come from public pages. It stores
        before it forwards, so a Suite outage can't destroy a lead.
        """
        from . import leads
        body = request.get_json(silent=True) or request.form.to_dict() or {}
        src = str(body.get("source") or "").strip()

        ip = leads.client_ip(request)
        allowed, retry_after = leads.rate_check(ip)
        if not allowed:
            # Recorded, because the number that stops a script is also the
            # number that could turn away a busy office sharing one address.
            # If real submissions start showing up here, raise
            # LEADS_RATE_LIMIT — a turned-away lead costs more than spam.
            audit.log("leads", "rate_limited", ip=ip, source=src[:60],
                      page=str(body.get("page") or "")[:120])
            return jsonify({
                "ok": False,
                "error": "Too many submissions from this connection. "
                         "Please try again shortly.",
            }), 429, {"Retry-After": str(retry_after)}

        if not src:
            return jsonify({"ok": False, "error": "source is required."}), 400
        fields = body.get("fields") if isinstance(body.get("fields"), dict) else {
            k: v for k, v in body.items()
            if k not in ("source", "page", "pdf_url", "client", "meta")}
        if not (fields.get("email") or fields.get("phone")):
            return jsonify({"ok": False,
                            "error": "An email or phone is required."}), 400
        return jsonify(leads.capture_and_deliver(
            src, str(body.get("page") or ""), fields,
            str(body.get("pdf_url") or ""), str(body.get("client") or ""),
            body.get("meta") if isinstance(body.get("meta"), dict) else None))

    @app.route("/api/leads/convert", methods=["POST"])
    def api_leads_convert():
        """Tie a prospect to the client account they became.

        A link, not a creation: a client here is anyone with a product in
        Knack, which is what billing reads. Writing an account from this
        endpoint would produce a client the Hub shows and no invoice ever
        mentions, so the account is still created in Knack and this records
        which lead it came from.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import clients_registry, leads
        body = request.get_json(silent=True) or request.form.to_dict() or {}
        lead_id = str(body.get("id") or "").strip()
        client = str(body.get("client") or "").strip()
        if not lead_id or not client:
            return jsonify({"ok": False,
                            "error": "A lead and a client are required."}), 400

        # The client has to exist, or "converted" would mean nothing checkable
        # and the name could drift from the account it is supposed to name.
        known = clients_registry.find_client(client)
        if known is None:
            return jsonify({
                "ok": False,
                "error": f"“{client}” is not in the client registry yet. "
                         f"Create the account in Knack first — that is what "
                         f"billing reads — then convert this lead to it.",
            }), 404

        row = leads.mark_converted(lead_id, str(known.get("name") or client),
                                   actor=current_user() or "")
        if row is None:
            return jsonify({"ok": False,
                            "error": "That lead could not be found."}), 404
        return jsonify({"ok": True, "lead": row})

    @app.route("/api/leads/duplicates")
    def api_leads_duplicates():
        """Rows that look like one prospect arriving more than once.

        A proposal list, never a merge. Grouping is on evidence that
        identifies a business exactly — one email address, one website — and
        an exact company name on its own is offered as *possible* and
        nothing more, because two franchises of one brand carry one name and
        are two businesses.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import leads
        return jsonify(leads.merge_candidates(
            days=clamp_int(request.args.get("days"), 365, 1, 730)))

    @app.route("/api/leads/merge", methods=["POST"])
    def api_leads_merge():
        """Fold one or more leads into another.

        The survivor keeps its own details and fills only its blanks, the
        merged rows are kept rather than deleted so where each came from
        survives, and nothing is re-delivered — two delivered rows mean the
        Suite already holds two contacts, and the answer to that is to say so
        rather than to write a third.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import leads
        body = request.get_json(silent=True) or {}
        ids = body.get("from")
        if isinstance(ids, str):
            ids = [ids]
        result = leads.merge(str(body.get("into") or ""),
                             ids if isinstance(ids, list) else [],
                             actor=current_user() or "")
        return jsonify(result), (200 if result.get("ok") else 400)

    @app.route("/api/leads/ghl/preflight")
    def api_leads_ghl_preflight():
        """Check lead delivery config against the live Suite API. Reads only.

        `?find=1` also lists the sub-accounts whose name matches, so the right
        location id can be copied out rather than hunted for in Suite.
        """
        gate = _require_api()
        if gate:
            return gate
        from .ghl_contacts import preflight
        return jsonify(preflight(find=request.args.get("find") == "1"))

    @app.route("/api/leads/retry", methods=["POST"])
    def api_leads_retry():
        gate = _require_api()
        if gate:
            return gate
        from . import leads
        return jsonify(leads.retry_undelivered())

    @app.route("/api/rate-card")
    def api_rate_card():
        """The rate card, so the proposal quotes what the IO enforces."""
        gate = _require_api()
        if gate:
            return gate
        from . import rate_card as rc
        term = request.args.get("q", "")
        if term:
            return jsonify({"products": rc.search(term)})
        return jsonify({"products": rc.products(),
                        "categories": rc.categories(),
                        "drift": rc.check_drift(),
                        # So a caller can tell "no products" from "couldn't
                        # read the card" — they look the same otherwise.
                        "source": rc.status()})

    @app.route("/api/rate-card/plan", methods=["POST"])
    def api_rate_card_plan():
        """Cost a set of products: delivery per line, plus the IO's guardrails.

        This is what makes the proposal show a live breakdown instead of a
        blank page until Generate.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import rate_card as rc
        items = (request.get_json(silent=True) or {}).get("items") or []
        lines, monthly = [], 0.0
        for i in items:
            # The line carries its own category, and it is what tells a
            # "Behavioral" on location lookback from the one on mobile.
            p = rc.find(str(i.get("product") or ""),
                        str(i.get("category") or "")) or {}
            budget = float(i.get("monthly") or 0)
            monthly += budget
            lines.append({**i, "listed_rate": p.get("listed_rate", ""),
                          "category": p.get("category", ""),
                          "requirements": p.get("requirements", ""),
                          "timeline": p.get("timeline", ""),
                          "delivery": rc.estimate_delivery(p, budget)})
        checks = rc.guardrails(items)
        return jsonify({
            "lines": lines,
            "monthly_total": round(monthly, 2),
            "annual_total": round(monthly * 12, 2),
            "guardrails": checks,
            "blocked": any(c["level"] == "block" for c in checks),
            "note": ("This plan can't be written as an IO as it stands."
                     if any(c["level"] == "block" for c in checks) else
                     "Within the rate card."),
        })

    @app.route("/api/qa/dashboard/<action>", methods=["POST"])
    def api_qa_dashboard(action):
        """Add a dashboard URL, or skip a client that doesn't need one."""
        gate = _require_api()
        if gate:
            return gate
        from . import qa
        body = request.get_json(silent=True) or {}
        client = str(body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        actor = current_user() or ""
        # Each of these three takes a client off the No Dashboards list — or,
        # for unskip, puts one back — so each drops the day's stored copy of
        # that report and of Active Clients, which carries the same dashboard
        # column. Skip and unskip do it inside `qa`, beside the write; adding
        # a URL writes through `knack_api`, which has no business knowing what
        # a QA report is, so that one is dropped here.
        if action == "skip":
            audit.log("qa", "dashboard_skipped", actor=actor, client=client,
                      reason=str(body.get("reason") or ""))
            return jsonify(qa.skip_dashboard(
                client, actor, str(body.get("reason") or "")))
        if action == "unskip":
            audit.log("qa", "dashboard_unskipped", actor=actor, client=client)
            return jsonify(qa.unskip_dashboard(client))
        if action == "add":
            from . import knack_api
            url = str(body.get("url") or "").strip()
            out = knack_api.set_dashboard_url(client, url)
            audit.log("qa", "dashboard_added", actor=actor, client=client,
                      ok=out.get("ok"), updated=out.get("updated"))
            if out.get("ok"):
                qa.forget("no-dashboards", "active-clients")
            return jsonify(out)
        return jsonify({"error": "Unknown action."}), 400

    @app.route("/api/qa/io-reconcile/<action>", methods=["POST"])
    def api_qa_io_reconcile(action):
        """Settle an insertion order that is never going to appear in Knack.

        A staff decision about a campaign, so it records who made it: a mark
        nobody can attribute is one nobody can revisit. It writes a small Hub
        overlay and nothing else — not Knack, not Smart 1 Suite, not the quote
        the order came from.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import io_reconcile, qa
        body = request.get_json(silent=True) or {}
        order = str(body.get("order") or "").strip()
        if not order:
            return jsonify({"ok": False, "error": "order is required."}), 400
        actor = current_user() or ""
        if action == "settle":
            out = io_reconcile.settle(
                order, reason=str(body.get("reason") or "other"),
                note=str(body.get("note") or ""), actor=actor)
        elif action == "unsettle":
            out = {"ok": bool(io_reconcile.unsettle(order)), "order": order}
            if not out["ok"]:
                out["error"] = "That order was not settled."
        else:
            return jsonify({"ok": False, "error": "Unknown action."}), 400
        if out.get("ok"):
            # The press takes a row off this report, so the day's stored copy
            # goes with it — otherwise the row is still there on the next open
            # and the button reads as having done nothing.
            qa.forget("io-not-in-knack")
            audit.log("qa", f"io_{action}", actor=actor, order=order,
                      reason=str(body.get("reason") or "") or None)
        return jsonify(out)

    @app.route("/api/qa/dashboard-skips")
    def api_qa_dashboard_skips():
        gate = _require_api()
        if gate:
            return gate
        from . import qa
        return jsonify(qa.skipped_dashboards())

    @app.route("/api/qa/form-summary/<opp_id>")
    def api_qa_form_summary(opp_id):
        """Everything the submitter actually filled in, for one request.

        The report shows one line per request because a queue has to be
        scannable. This returns the whole form for the moment someone needs
        the detail, without sending them into GoHighLevel to find it.
        """
        gate = _require_api()
        if gate:
            return gate
        try:
            from modules.suite_panel.app import ghl
            data = ghl(f"/opportunities/{opp_id}") or {}
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": f"Couldn't reach Smart 1 Suite "
                                     f"({type(exc).__name__})."}), 200
        opp = data.get("opportunity") or data
        contact = opp.get("contact") or {}
        fields = []
        for cf in (opp.get("customFields") or []):
            label = str(cf.get("name") or cf.get("id") or "").strip()
            value = cf.get("fieldValue")
            if isinstance(value, list):
                value = ", ".join(str(v) for v in value)
            value = str(value or "").strip()
            if label and value:
                fields.append({"label": label, "value": value[:2000]})
        return jsonify({
            "name": opp.get("name", ""),
            "contact": {"name": contact.get("name", ""),
                        "email": contact.get("email", ""),
                        "phone": contact.get("phone", ""),
                        "company": contact.get("companyName", "")},
            "status": opp.get("status", ""),
            "created": str(opp.get("createdAt") or "")[:10],
            "fields": fields,
            "note": ("" if fields else
                     "This request has no custom form fields recorded — the "
                     "submitter may have used a different form."),
        })

    # ---- Landing Page Maker -------------------------------------------
    #
    # These sit on the HUB app, and /sales is deliberately not a mount --
    # /sales/builder and /sales/proposals are, but the prefix itself is not,
    # so /sales/landing reaches this app. Adding a "/sales" mount later would
    # make every route below unreachable without erroring; /api/integrity
    # checks for exactly that.
    @app.route("/sales/landing")
    def page_landing_maker():
        gate = _require_page()
        if gate:
            return gate
        from . import landing_maker as lm
        return render_template("landing_maker.html", user=current_user(),
                               active="landing", directions=lm.DIRECTIONS)

    @app.route("/sales/landing/p/<slug>")
    def page_landing_preview(slug):
        """The built page itself. Public, and deliberately ungated.

        It is a landing page: the people it is built for are prospects on a
        client's campaign, not staff with a Hub login. It returns the stored
        HTML verbatim, and because the hub app is the dispatcher default
        rather than something wrapped in HubBar, no staff sidebar is injected
        into it -- which is what keeps internal navigation off a page that may
        be pasted onto a client's own domain.
        """
        from . import landing_maker as lm
        row = lm.get(slug)
        if not row:
            return "No such landing page.", 404
        return row.get("page_html", ""), 200, {"Content-Type": "text/html"}

    @app.route("/api/landing")
    def api_landing_list():
        gate = _require_api()
        if gate:
            return gate
        from . import landing_maker as lm
        return jsonify(lm.listing(request.args.get("client", ""),
                                  request.args.get("q", "")))

    @app.route("/api/landing/proposals")
    def api_landing_proposals():
        """Both kinds of proposal on one client, for the picker.

        The page asks for the client first, because "which proposal?" is only
        answerable once you know whose. A global list of every proposal in the
        Hub is the wrong question and gets longer every week.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import landing_maker as lm
        return jsonify(lm.proposals_for(request.args.get("client", "")))

    @app.route("/api/landing", methods=["POST"])
    def api_landing_create():
        gate = _require_api()
        if gate:
            return gate
        from . import landing_maker as lm
        # An upload arrives as multipart, everything else as JSON. Reading
        # request.form for the multipart case matters: get_json() is empty
        # there, so a file upload would otherwise lose the goal and offer.
        body = request.get_json(silent=True) or {}
        text = str(body.get("text") or "")
        up = request.files.get("file") if request.files else None
        if up and up.filename:
            text = _read_document(up.read(8 * 1024 * 1024), up.filename)
            body = {**request.form.to_dict(), **body}
        return jsonify(lm.create(
            proposal_id=str(body.get("proposal_id") or ""),
            uploaded_id=str(body.get("uploaded_id") or ""),
            client=str(body.get("client") or ""), text=text,
            kind=str(body.get("kind") or "client"),
            website=str(body.get("website") or ""),
            direction=str(body.get("direction") or "trust"),
            goal=str(body.get("goal") or ""), offer=str(body.get("offer") or ""),
            promoting=str(body.get("promoting") or ""),
            actor=current_user() or ""))

    @app.route("/api/landing/goals")
    def api_landing_goals():
        """The page goals the maker offers, from the one list that defines them.

        Served rather than written into the template, so the choices, the
        copy prompt and the rendered form cannot disagree about what a goal
        is -- the drift that a second hand-kept copy of a list guarantees.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import landing_spec
        return jsonify({"goals": landing_spec.goal_choices(),
                        "default": landing_spec.DEFAULT_GOAL})

    @app.route("/api/landing/offer-check")
    def api_landing_offer_check():
        """Whether an offer is usable as written, before the page is built.

        Said while the rep is still typing rather than after a page has been
        generated around a promise it cannot keep.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import landing_spec
        state, reason = landing_spec.offer_state(request.args.get("offer", ""))
        return jsonify({"state": state, "reason": reason,
                        "usable": state == landing_spec.READ})

    @app.route("/api/landing/<page_id>/revise", methods=["POST"])
    def api_landing_revise(page_id):
        """Rewrite a built page against an instruction, keeping the old one."""
        gate = _require_api()
        if gate:
            return gate
        from . import landing_maker as lm
        body = request.get_json(silent=True) or {}
        return jsonify(lm.revise(page_id, str(body.get("instructions") or ""),
                                 current_user() or ""))

    @app.route("/api/landing/<page_id>", methods=["DELETE"])
    def api_landing_delete(page_id):
        gate = _require_api()
        if gate:
            return gate
        from . import landing_maker as lm
        return jsonify(lm.remove(page_id, current_user() or ""))

    @app.route("/api/landing/<page_id>", methods=["POST"])
    def api_landing_save(page_id):
        gate = _require_api()
        if gate:
            return gate
        from . import landing_maker as lm
        body = request.get_json(silent=True) or {}
        return jsonify(lm.update_html(page_id, str(body.get("html") or ""),
                                      current_user() or ""))

    @app.route("/api/providers")
    def api_providers():
        """Every provider, configured or not, with what breaks when it isn't.

        This is the answer to "is Cloudinary actually set?" — a question that
        went unanswered long enough for the whole v1.6.0 tool set to run
        degraded in production without anyone noticing.
        """
        gate = _require_api()
        if gate:
            return gate
        from .config import settings as _cfg
        rows = _cfg.status()
        return jsonify({
            "providers": rows,
            "missing_required": _cfg.missing_required(),
            "ok": not _cfg.missing_required(),
            "degraded": [r["name"] for r in rows if r["state"] == "warn"],
        })

    # ---------------- clients: one list from every source ----------------
    @app.route("/api/clients/search")
    def api_clients_search():
        gate = _require_api()
        if gate:
            return gate
        from . import clients_registry
        rows = clients_registry.search_clients(request.args.get("q", ""),
                                               limit=clamp_int(request.args.get("limit"), 12, 1, 500))
        return jsonify({"clients": rows})

    @app.route("/api/clients/house", methods=["GET", "POST"])
    def api_house_clients_hub():
        gate = _require_api()
        if gate:
            return gate
        from . import clients_registry
        if request.method == "GET":
            return jsonify({"clients": clients_registry.house_clients()})
        body = request.get_json(silent=True) or {}
        if body.get("delete"):
            ok = clients_registry.delete_house_client(str(body.get("slug") or ""))
            clients_registry.all_clients(refresh=True)
            return jsonify({"ok": ok, "clients": clients_registry.house_clients()})
        try:
            row = clients_registry.add_house_client(
                body.get("name", ""), body.get("url", ""), body.get("notes", ""),
                actor=current_user() or "")
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        clients_registry.all_clients(refresh=True)
        audit.log("hub", "house_client_added", actor=current_user(),
                  detail=row["name"])
        return jsonify({"ok": True, "client": row,
                        "clients": clients_registry.house_clients()})

    def _hub_user():
        """The signed-in user as a rich object, for the help/demo layer."""
        try:
            from . import identity
            u = identity.user_from_environ(request.environ)
            if u:
                return u
            name = current_user()
            return identity.User(email="", name=name, via="password") if name else None
        except Exception:  # noqa: BLE001
            return None

    def _login_cookie(user, nxt="/"):
        from . import identity
        if not nxt.startswith("/"):
            nxt = "/"
        resp = make_response(redirect(nxt))
        secure = (os.environ.get("NODE_ENV") == "production"
                  or os.environ.get("FLASK_ENV") == "production")
        # Both cookies: the rich one for v7 features, the legacy one so every
        # existing @requires_login check keeps working untouched.
        resp.set_cookie(identity.COOKIE_NAME, identity.issue_cookie(user),
                        max_age=identity.SESSION_TTL_SECONDS, httponly=True,
                        samesite="Lax", secure=secure)
        resp.set_cookie(auth.COOKIE_NAME, auth.issue_cookie_value(user.name),
                        max_age=auth.SESSION_TTL_SECONDS, httponly=True,
                        samesite="Lax", secure=secure)
        # And the SameSite=None companion, so a page framed inside Smart 1
        # Suite is not anonymous. It grants strictly less than the two above:
        # see hub/embed.py.
        from . import suite_embed as _embed
        _embed.issue_cookie(resp, user.name, secure)
        return resp

    @app.route("/auth/google")
    def auth_google():
        from . import identity
        if not identity.google_configured():
            return redirect("/login?error=google_not_configured")
        state = identity.new_state()
        redirect_uri = request.url_root.rstrip("/") + "/auth/google/callback"
        resp = make_response(redirect(identity.authorize_url(redirect_uri, state)))
        # State is held in a short-lived cookie and compared on the way back —
        # without it the callback accepts a code an attacker supplies.
        resp.set_cookie("s1_oauth_state", state, max_age=600, httponly=True,
                        samesite="Lax")
        resp.set_cookie("s1_oauth_next", request.args.get("next", "/"),
                        max_age=600, httponly=True, samesite="Lax")
        return resp

    @app.route("/auth/google/callback")
    def auth_google_callback():
        from . import identity
        sent = request.cookies.get("s1_oauth_state") or ""
        got = request.args.get("state") or ""
        if not sent or sent != got:
            audit.log("auth", "login_rejected", reason="state_mismatch")
            return render_template("login.html", next="/",
                                   error="That sign-in link expired. Try again."), 400
        code = request.args.get("code") or ""
        if not code:
            return redirect("/login")
        try:
            user = identity.complete_google_login(
                code, request.url_root.rstrip("/") + "/auth/google/callback")
        except identity.LoginRejected as exc:
            return render_template("login.html", next="/", error=str(exc)), 403
        nxt = request.cookies.get("s1_oauth_next") or "/"
        resp = _login_cookie(user, nxt)
        resp.delete_cookie("s1_oauth_state")
        resp.delete_cookie("s1_oauth_next")
        return resp

    @app.route("/auth/demo", methods=["POST"])
    def auth_demo():
        from . import identity
        try:
            user = identity.complete_demo_login(
                request.form.get("code") or "",
                request.form.get("name") or "")
        except identity.LoginRejected as exc:
            return render_template("login.html", next="/", error=str(exc)), 403
        return _login_cookie(user, "/")

    @app.route("/login", methods=["GET", "POST"])
    def login():
        """The single sign-in page.

        Tries a real user account first, then falls back to the legacy shared
        PANEL_PASSWORD. The fallback stays until every account is migrated —
        removing it while somebody still depends on it locks them out of their
        own tool with no way back in.

        Unlike /signin, this deliberately tells you when an email has no
        account and points at sign-up. That does leak which addresses are
        registered, which is normally worth avoiding — but sign-up is already
        restricted to @smart1marketing.com, so the only thing an attacker
        learns is which colleagues have signed up yet. For an internal tool
        that trade is worth making; being told "wrong email or password" when
        you simply haven't registered is how people give up.
        """
        from . import identity
        google_on = (os.environ.get("HUB_GOOGLE_LOGIN", "").lower()
                     in {"1", "true", "yes", "on"}) and identity.google_configured()

        def page(error=None, offer_signup=False, last_email="", code=200):
            return render_template(
                "login.html", next=request.form.get("next") or request.args.get("next", "/"),
                error=error, offer_signup=offer_signup, last_email=last_email,
                google_enabled=google_on), code

        if request.method == "GET":
            if current_user():
                return redirect(request.args.get("next") or "/")
            return page()[0]

        # Last hop, not the first: the client-supplied first entry in
        # X-Forwarded-For is spoofable, and this exact mistake was flagged in
        # three separate apps during the suite audit. One helper now, because
        # it was written out longhand at four call sites and one of them had
        # it backwards.
        ip = auth.client_ip(request.headers, request.remote_addr or "")
        wait = auth.throttle_check(ip)
        if wait:
            return page(f"Too many attempts. Try again in "
                        f"{max(1, wait // 60)} minute(s).", code=429)

        email = (request.form.get("email") or "").strip().lower()
        password = request.form.get("password") or ""
        nxt = request.form.get("next") or "/"
        if not nxt.startswith("/"):
            nxt = "/"

        # ---- 1. the shared password always works, checked FIRST ----
        #
        # Order matters and this is the bug that locked Todd out of his own
        # Hub: the three founding super admins are SEEDED as rows before
        # anyone sets a password, so by_email() finds them, the account branch
        # takes over, and the shared password is never reached. Checking the
        # shared password first means it is a genuine way back in rather than
        # one that silently stops working the moment accounts are seeded.
        shared_ok = bool(auth.panel_password()) and auth.check_password(password)

        # ---- 2. real user account ----
        account = None
        if not shared_ok:
            try:
                from . import users as _users
                from .users_routes import _login_response
                account = _users.by_email(email) if email else None
                if account is not None and account.password_hash:
                    try:
                        user = _users.authenticate(email, password)
                        auth.throttle_reset(ip)
                        return _login_response(user, nxt)
                    except _users.UserError as exc:
                        auth.throttle_fail(ip, email)
                        return page(str(exc), last_email=email, code=401)
            except ImportError:
                account = None

        # ---- 3. legacy shared password (emergency access) ----
        if shared_ok:
            auth.throttle_reset(ip)
            actor = email or "Shared login"
            audit.log("hub", "login_shared_password", actor=actor, ip=ip)
            resp = make_response(redirect(nxt))
            _secure = (os.environ.get("NODE_ENV") == "production"
                       or os.environ.get("FLASK_ENV") == "production")
            resp.set_cookie(
                auth.COOKIE_NAME, auth.issue_cookie_value(actor),
                max_age=auth.SESSION_TTL_SECONDS, httponly=True, samesite="Lax",
                secure=_secure)
            from . import suite_embed as _embed
            _embed.issue_cookie(resp, actor, _secure)
            return resp

        # ---- 3. no account, and not the shared password ----
        # The address goes in too: one IP working down the staff list is the
        # attack a per-attempt counter cannot see, because fourteen addresses
        # at one guess each never reaches six on any of them.
        auth.throttle_fail(ip, email)
        audit.log("hub", "login_failed", actor=email or "?", ip=ip)
        if email and account is None:
            return page(f"There's no account for {email} yet.",
                        offer_signup=True, last_email=email, code=401)
        return page("That email and password don't match.",
                    last_email=email, code=401)

    @app.route("/logout")
    def logout():
        resp = make_response(redirect("/login"))
        resp.delete_cookie(auth.COOKIE_NAME)
        # The embed companion authenticates on its own, so leaving it behind
        # would keep a framed page signed in after a logout that looked like
        # it worked.
        from . import suite_embed as _embed
        _embed.clear_cookie(resp)
        return resp

    # ---------------- access level: General or Admin ----------------
    def current_account():
        """The signed-in *account* row, or None for a shared-password session.

        Re-read every request rather than trusted from the cookie: a role
        change or a suspension has to take effect on the next click, not
        whenever the cookie happens to expire.
        """
        try:
            from .users_routes import current_account as _acct
            return _acct()
        except Exception:  # noqa: BLE001 — never 500 a page over the gate
            return None

    _UNREAD = object()

    def viewer_is_admin(account=_UNREAD) -> bool:
        """Does this request get the Utilities section?

        Three cases, and the middle one is the decision worth stating:

          * an account -> its own role decides;
          * a shared-password session -> Admin, because PANEL_PASSWORD is the
            emergency door and an emergency door that cannot reach Diagnostics
            leads nowhere. `hub/access.py` says the rest of it, including that
            the way to close the door is to clear the variable;
          * nobody signed in -> False, though no gate reaches here: the login
            redirect runs first.

        `account` is for a caller that has already read the row: it is
        deliberately not cached per request, so asking again is a second query
        for an answer just fetched. Passing it in keeps the rule in one place
        rather than having each such caller restate `account.is_admin` and
        quietly disagree about what None means.
        """
        if account is _UNREAD:
            account = current_account()
        if account is not None:
            return bool(account.is_admin)
        return bool(current_user())

    # Paths a signed-in account may reach while it still owes a password
    # change. Everything else redirects to /account until the starting
    # password is gone -- including the API routes, or a page would render its
    # shell and then fail every fetch inside it.
    _PASSWORD_GATE_OPEN = ("/account", "/signout", "/logout", "/login",
                           "/reset", "/forgot", "/assets/", "/static/",
                           "/hub-", "/favicon.ico", "/robots.txt",
                           "/llms.txt", "/api/version")

    @app.before_request
    def _password_change_gate():
        """A starting password is valid for exactly one sign-in.

        Without this, `must_change_password` is a note on an admin panel:
        every roster account would keep `Smart12026!` indefinitely, on a Hub
        whose password is written down in a repository. The flag has to stop
        something, and this is the something.

        A JSON path answers 403 with the redirect named in the body, rather
        than serving a redirect a `fetch()` would follow into an HTML login
        page and report as malformed data.
        """
        path = request.path or "/"
        if any(path.startswith(p) for p in _PASSWORD_GATE_OPEN):
            return None
        account = current_account()
        if account is None or not account.must_change_password:
            return None
        from . import access
        if access.wants_json(path, request.headers.get("Accept", "")):
            return jsonify({"error": "Set a new password before using the Hub.",
                            "redirect": "/account"}), 403
        return redirect("/account?first=1")

    @app.before_request
    def _utilities_gate():
        """General Access sees everything except Utilities.

        One list in `hub/access.py`, checked here on every request, rather
        than a decorator each new Utilities route has to remember. The
        alternative shipped once already: a whole module answered 200 to
        anyone with the URL because it never passed the guard the tiles beside
        it did.
        """
        from . import access
        path = request.path or "/"
        if not access.is_utility(path):
            return None
        if not current_user():
            return None                 # not signed in: the login gate answers
        if viewer_is_admin():
            if current_account() is None:
                audit.log("auth", "shared_password_utility", path=path)
            return None
        account = current_account()
        audit.log("auth", "utility_refused",
                  actor=account.email if account else current_user(), path=path)
        if access.wants_json(path, request.headers.get("Accept", "")):
            return jsonify({
                "error": f"{access.SECTION_LABEL} is for admin accounts.",
                "level": "General"}), 403
        return render_template("users_not_admin.html",
                               section=access.SECTION_LABEL,
                               account=account, user=current_user(),
                               support_email=_support_email(),
                               support_name=_support_name()), 403

    @app.before_request
    def _record_presence():
        """Note that this person was seen, for the headcount on the dashboard.

        There is no session table — signing in issues a signed cookie and the
        server keeps nothing — so "who is logged in" is answered by "who has
        been seen lately", and `hub/presence.py` is where that is written down
        along with why. Throttled to one write per person per minute per
        worker, so this is a dict lookup on almost every request.

        Registered after the gates rather than before them on purpose: a
        person refused a Utilities page is still signed in and still at their
        desk. It never returns a response and never raises — a headcount is a
        nice-to-have and no route depends on it.
        """
        try:
            name = current_user()
            if name:
                from . import presence
                presence.touch_display(name)
        except Exception:  # noqa: BLE001 — presence must never cost a page
            pass
        return None

    def _support_name() -> str:
        from .user_directory import SUPPORT_NAME
        return SUPPORT_NAME

    def _support_email() -> str:
        from .user_directory import SUPPORT_EMAIL
        return SUPPORT_EMAIL

    # ---------------- keeping crawlers out ----------------
    @app.route("/robots.txt")
    def robots_txt():
        """Deliberately no login: a robots.txt behind a login is not read.

        The header in `hub/no_crawl.NoIndex` is what actually removes a page
        from an index; this is the half the well-behaved crawler asks for
        first, and it names the AI crawlers individually because several of
        them -- Google-Extended and Applebot-Extended among them -- honour
        only their own token and ignore the wildcard.
        """
        from . import no_crawl
        resp = make_response(no_crawl.robots_txt())
        resp.mimetype = "text/plain"
        return resp

    @app.route("/llms.txt")
    def hub_llms_txt():
        """The Hub's own llms.txt, which says the opposite of a client's.

        `hub/llms_txt.py` builds one FOR a client, where being read is the
        point. This one is about this host, and a model that fetches it and
        finds a 404 learns nothing about whether it was welcome.
        """
        from . import no_crawl
        resp = make_response(no_crawl.llms_txt())
        resp.mimetype = "text/plain"
        return resp

    def _require_page():
        """Redirect helper for HTML pages."""
        if not current_user():
            return redirect("/login?next=" + request.path)
        return None

    # ---------------- shell pages ----------------
    @app.route("/")
    def dashboard():
        gate = _require_page()
        if gate:
            return gate
        from . import partner as partner_pages
        return render_template("dashboard.html", user=current_user(),
                               modules=MODULES, active="dashboard",
                               partner_tiles=partner_pages.tiles())

    @app.route("/client360")
    def client360():
        gate = _require_page()
        if gate:
            return gate
        return render_template("client360.html", user=current_user(), modules=MODULES,
                               active="c360", q=request.args.get("q", ""))

    @app.route("/tools")
    def tools():
        gate = _require_page()
        if gate:
            return gate
        return render_template("tools.html", user=current_user(), modules=MODULES, active="tools")

    # ---------------- SEO section ----------------
    @app.route("/seo")
    def seo_home():
        gate = _require_page()
        if gate:
            return gate
        return render_template("seo.html", user=current_user(), modules=MODULES, active="seo")

    @app.route("/seo/client")
    def seo_client_page():
        gate = _require_page()
        if gate:
            return gate
        name = (request.args.get("name") or "").strip()
        if not name:
            return redirect("/seo")
        return render_template("seo_client.html", user=current_user(), modules=MODULES,
                               active="seo", client=name)

    @app.route("/seo/webmaster")
    def seo_webmaster_page():
        gate = _require_page()
        if gate:
            return gate
        return render_template("seo_webmaster.html", user=current_user(),
                               modules=MODULES, active="seo")

    @app.route("/api/seo/tasks")
    def api_seo_tasks():
        """What has been raised for this client, and whether it can be."""
        gate = _require_api()
        if gate:
            return gate
        from . import seo_tasks, knack_api as _k
        client = (request.args.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        return jsonify({
            "counts": seo_tasks.status(client),
            "enabled": seo_tasks.enabled(),
            "knack_configured": _k.configured(),
            "due_field": seo_tasks.due_field(),
            "rules": {"faq_days": seo_tasks.DUE_DAYS_FAQ,
                      "schema_days": seo_tasks.DUE_DAYS_SCHEMA,
                      "blog_lead_days": seo_tasks.BLOG_LEAD_DAYS},
        })

    @app.route("/api/seo/webmaster")
    def api_seo_webmaster():
        """The roster only. Numbers arrive per row from /google — see
        hub/seo.webmaster_roster for why this route makes no Google call."""
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        try:
            return jsonify({"clients": seo.webmaster_roster()})
        except Exception as exc:  # noqa: BLE001
            return jsonify({"clients": [], "error": str(exc)})

    @app.route("/api/seo/clients")
    def api_seo_clients():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        try:
            return jsonify({"clients": seo.seo_clients()})
        except Exception as exc:  # noqa: BLE001
            return jsonify({"clients": [], "error": str(exc)})

    @app.route("/api/seo/detail")
    def api_seo_detail():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        name = (request.args.get("name") or "").strip()
        try:
            return jsonify(seo.client_detail(name, full=bool(request.args.get("full"))))
        except Exception as exc:  # noqa: BLE001
            errors.log_exception("seo-detail", exc, path=request.path,
                                 actor=current_user() or "")
            return jsonify({"client": name, "error": str(exc)})

    @app.route("/api/seo/scan", methods=["POST"])
    def api_seo_scan():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        url = (body.get("url") or "").strip()
        client = (body.get("client") or "").strip()
        if not url:
            return jsonify({"error": "No URL provided."}), 400
        try:
            out = seo.scan_schema(url)
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": f"Could not scan the site: {exc}"})
        if client:                      # cache the result on the client store
            import datetime as _dt
            from . import dates as _dates
            _now = _dt.datetime.now()
            out["at"] = _dates.fmt(_now) + _now.strftime(" %I:%M %p")
            store = seo.load_store(client)
            store["last_scan"] = out
            seo.save_store(client, store)
        audit.log("hub", "seo_scan", actor=current_user(), detail=url)
        return jsonify(out)

    @app.route("/api/seo/sitemap", methods=["POST"])
    def api_seo_sitemap():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        url = (body.get("url") or "").strip()
        if not client or not url:
            return jsonify({"error": "client and url are required."}), 400
        try:
            pages = seo.sitemap_pages(url)
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": f"Could not read the sitemap: {exc}"})
        store = seo.load_store(client)
        store["sitemap"] = pages
        store["site_url"] = url
        seo.save_store(client, store)
        done = set(store.get("pages", {}))
        return jsonify({"total": len(pages), "generated": len(done),
                        "remaining": [p for p in pages if p not in done][:10],
                        "pages": pages[:50]})

    @app.route("/api/seo/generate", methods=["POST"])
    def api_seo_generate():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        store = seo.load_store(client)
        urls = body.get("urls")
        if not urls:
            done = set(store.get("pages", {}))
            urls = [p for p in store.get("sitemap", []) if p not in done][:10]
        if not urls:
            return jsonify({"pages": [], "questions": store.get("questions", []),
                            "done": True})
        try:
            out = seo.generate_for_pages(client, urls[:10])
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": str(exc)})
        store = seo.load_store(client)
        remaining = [p for p in store.get("sitemap", []) if p not in store.get("pages", {})]
        out["remaining"] = len(remaining)
        out["total"] = len(store.get("sitemap", []))
        audit.log("hub", "seo_generate", actor=current_user(),
                  detail=f"{client}: {len(urls)} pages")
        # One ticket per page that actually got schema, not per page asked for.
        from . import seo_tasks
        done = [pg.get("url") for pg in (out.get("pages") or []) if pg.get("url")]
        out["tasks"] = seo_tasks.for_pages(client, done or urls[:10],
                                           kind="schema", actor=current_user())
        return jsonify(out)

    @app.route("/api/seo/page", methods=["POST"])
    def api_seo_page():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        url = (body.get("url") or "").strip()
        if not client or not url:
            return jsonify({"error": "client and url are required."}), 400
        store = seo.load_store(client)
        page = store.setdefault("pages", {}).get(url)
        if page is None:
            return jsonify({"error": "Unknown page."}), 404
        if "schema" in body:
            sch = body["schema"]
            if isinstance(sch, str):
                try:
                    sch = json.loads(sch)
                except ValueError as exc:
                    return jsonify({"error": f"Schema is not valid JSON: {exc}"}), 400
            page["schema"] = sch
        if "approved" in body:
            page["approved"] = bool(body["approved"])
        if "posted" in body:
            page["posted"] = bool(body["posted"])
        seo.save_store(client, store)
        approved = sum(1 for p in store["pages"].values() if p.get("approved"))
        return jsonify({"ok": True, "approved_total": approved})

    @app.route("/api/seo/business", methods=["POST"])
    def api_seo_business():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        store = seo.load_store(client)
        if isinstance(body.get("business_info"), dict):
            store.setdefault("business_info", {}).update(body["business_info"])
        if isinstance(body.get("answers"), dict):
            store.setdefault("answers", {}).update(
                {k: v for k, v in body["answers"].items() if str(v).strip()})
            store["questions"] = [q for q in store.get("questions", [])
                                  if q not in store["answers"]]
        seo.save_store(client, store)
        return jsonify({"ok": True, "questions": store.get("questions", [])})

    @app.route("/api/seo/setup", methods=["POST"])
    def api_seo_setup():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        store = seo.load_store(client)
        setup = store.setdefault("setup", {})
        for k in ("access_method", "access_url", "login", "password",
                  "webmaster_status", "blogs_enabled", "blogs_per_month",
                  "blogs_frequency", "completed", "skipped_steps", "notes"):
            if k in body:
                setup[k] = body[k]
        seo.save_store(client, store)
        audit.log("hub", "seo_setup_saved", actor=current_user(), detail=client)
        return jsonify({"ok": True})

    @app.route("/api/seo/pages")
    def api_seo_pages():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        name = (request.args.get("name") or "").strip()
        store = seo.load_store(name)
        pages = list(store.get("pages", {}).values())
        remaining = [p for p in store.get("sitemap", []) if p not in store.get("pages", {})]
        return jsonify({"pages": pages, "total": len(store.get("sitemap", [])),
                        "remaining": len(remaining)})

    @app.route("/api/seo/checks", methods=["POST"])
    def api_seo_checks():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        store = seo.load_store(client)
        checks = store.setdefault("checks", {})
        for k in ("schema", "listings"):
            if k in body:
                checks[k] = bool(body[k])
        if "setup" in body:
            store.setdefault("setup", {})["completed"] = bool(body["setup"])
        seo.save_store(client, store)
        audit.log("hub", "seo_checks", actor=current_user(), detail=client)
        return jsonify({"ok": True, "status": seo.client_status(store)})

    @app.route("/api/client/social")
    def api_client_social():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        name = (request.args.get("name") or "").strip()
        domain = (request.args.get("domain") or "").strip()
        social = seo.get_social(name, domain) if name else {}
        # Merge anything a site scan found. Brandfetch returns what a brand
        # publishes about itself — usually two or three profiles. A scan reads
        # the client's own pages and routinely finds more: a TikTok in the
        # footer, a LinkedIn nobody registered with Brandfetch.
        #
        # Saved values win. Someone who corrected a URL by hand should not
        # have it overwritten by the next scan.
        found_by_scan = []
        try:
            from modules.scans.app import latest_payload_for_domain
            from modules.scans.reports import social_profiles
            payload = latest_payload_for_domain(domain or name)
            for key, url in (social_profiles(payload or {}) or {}).items():
                if not str(social.get(key) or "").strip():
                    social[key] = url
                    found_by_scan.append(key)
        except Exception:  # noqa: BLE001
            pass
        return jsonify({"social": social, "from_scan": found_by_scan,
                        "note": (f"{len(found_by_scan)} profile(s) came from "
                                 f"the last site scan rather than Brandfetch."
                                 if found_by_scan else "")})

    @app.route("/api/client/social", methods=["POST"])
    def api_client_social_set():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        social = seo.set_social(client, body.get("social") or {})
        audit.log("hub", "client_social_saved", actor=current_user(), detail=client)
        return jsonify({"ok": True, "social": social})

    # ------------- the Smart 1 Suite app frame (a CLIENT, not a rep)
    #
    # The path is "/suite-app" and deliberately NOT "/suite/app": /suite is a
    # dispatcher-mounted module, and a hub route under a mounted prefix never
    # receives the request — it 404s with every template valid and every link
    # resolving. This is the first trap CLAUDE.md names and it has bitten four
    # times now; /api/integrity's high-severity check is what caught this one.
    #
    # hub/suite_embed.py solved the staff case: a rep already has a Hub session
    # in that browser. A client has none and must never be given one, so
    # identity comes from HighLevel's SSO handshake instead — the framed page
    # asks its parent for the user payload, HighLevel replies encrypted under
    # the app's SSO key, and hub/suite_sso.py decrypts it here to learn which
    # sub-account this person is in.
    #
    # These two routes are deliberately the whole client-facing surface. They
    # prove who is looking and then hand them their *existing* content link —
    # the pages behind it are already client-facing, already scoped to one
    # client and already tested, so no new place to see data is created here.
    # Somewhere for one client to be shown another client's record is exactly
    # what this handshake exists to prevent, and the smallest way to build it
    # is not to build one.
    @app.route("/suite-app")
    def suite_app_frame():
        from . import suite_sso
        # Not configured is said in words rather than drawn as a broken frame.
        # This page is only ever reached from inside Suite, so the reader is
        # whoever configured the menu link.
        return render_template("suite_app.html",
                               configured=suite_sso.configured(),
                               why_not=suite_sso.why_not())

    @app.route("/suite-app/session", methods=["POST"])
    def suite_app_session():
        """Turn an SSO payload into somewhere this client may go.

        No client name is taken from the request — the location id inside the
        decrypted payload is the only thing identity comes from, which is the
        whole security model and the reason this route accepts nothing else.
        """
        from . import suite_sso
        body = request.get_json(silent=True) or {}
        found = suite_sso.identify(str(body.get("payload") or ""))
        if not found["ok"]:
            # The four refusals are named for whoever has to fix one, and none
            # of them says anything a prober could tune a guess against: an
            # unreadable payload is unreadable whatever was wrong with it.
            return jsonify({"ok": False, "state": found["state"],
                            "error": found["detail"]}), 403
        from modules.social_planner import links as social_links
        target = social_links.link(found["client"], found.get("client_url", ""),
                                   "approve", request.host_url)
        audit.log("hub", "suite_sso_session", actor="suite",
                  client=found["client"], location=found["location_id"],
                  user=found["user"].get("email", ""))
        return jsonify({"ok": True, "client": found["client"], "url": target})

    # ------------- social content: requests, ideas, the client's own link
    #
    # The Social Media card above is the client's profile URLs, which is a
    # different question from "is anybody at this client asking us for
    # anything". Before this the record said nothing at all about the work:
    # a client could have three requests overdue, a link nobody had sent them
    # and four posts sitting unanswered, and none of it was on the one screen
    # a rep opens. hub/social_status.py answers it, and the dashboard
    # scoreboard reads the same module so the two cannot disagree.
    @app.route("/api/client/social-content")
    def api_client_social_content():
        gate = _require_api()
        if gate:
            return gate
        from . import social_status
        name = (request.args.get("name") or "").strip()
        url = (request.args.get("url") or "").strip()
        if not name:
            return jsonify({"measured": False, "error": "No client was named."})
        return jsonify(social_status.for_client(name, url,
                                                request.host_url))

    @app.route("/api/social/scoreboard")
    def api_social_scoreboard():
        """Who is waiting on us, for the dashboard.

        Deliberately not behind `access.UTILITY_PREFIXES`: this is the
        workload of the people reading the dashboard, and the presence
        headcount already showed what happens when a figure everybody sees is
        served by a path most accounts are refused — the panel rendered a
        confident green nothing for eleven of fourteen people.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import social_status
        return jsonify(social_status.scoreboard())

    @app.route("/api/sales/prospects")
    def api_sales_prospects():
        """Who is waiting to be called, for the dashboard.

        The prospect queue lives on `/qa/prospect-queue`, and a queue nobody
        is told has anything in it is the failure it was built to undo one
        step later — the note `hub/social_status.py` makes about there being
        no mailer here, so the honest route is putting the number where people
        already look.

        Read from the queue's own day cache rather than rebuilt, so this tile
        and that report cannot answer the same question differently, and so a
        page that loads on every visit does not walk the lead store to do it.

        Not behind `access.UTILITY_PREFIXES`, for the same reason the two
        scoreboards below give: this is the work of the people reading the
        dashboard, and a figure everybody sees served by a path most accounts
        are refused renders a confident nothing for eleven of the fourteen.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import prospect_queue
        return jsonify(prospect_queue.scoreboard())

    @app.route("/api/sales/scoreboard")
    def api_sales_scoreboard():
        """What the pipeline is worth and what needs chasing, for the dashboard.

        Not behind `access.UTILITY_PREFIXES`, for the reason the social
        scoreboard above gives: this is the work of the people reading the
        dashboard, and a figure everybody sees that is served by a path most
        accounts are refused renders a confident nothing for eleven of the
        fourteen.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import sales_status
        return jsonify(sales_status.scoreboard())

    # ------------- attached Google accounts (shared: SEO page + Client 360)
    @app.route("/api/client/links")
    def api_client_links():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        name = (request.args.get("name") or "").strip()
        return jsonify({"attached": seo.get_links(name) if name else {}})

    @app.route("/api/client/links", methods=["POST"])
    def api_client_links_set():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        kind = (body.get("kind") or "").strip()
        if not client or kind not in seo.LINK_KINDS:
            return jsonify({"error": f"client and kind ({'|'.join(seo.LINK_KINDS)}) are required."}), 400
        att = seo.set_link(client, kind, body.get("data"),
                           remove=(body.get("remove") or "").strip())
        audit.log("hub", "client_account_attached", actor=current_user(),
                  detail=f"{client}: {kind}")
        return jsonify({"ok": True, "attached": att})

    @app.route("/api/client/profile")
    def api_client_profile():
        gate = _require_api()
        if gate:
            return gate
        from . import seo, client_groups, scan_facts
        name = (request.args.get("name") or "").strip()
        if not name:
            return jsonify({"profile": {}})
        prof = seo.get_profile(name)
        # What their own website publishes, offered into the fields nobody has
        # filled in. Suggested, never saved: the profile is what a person
        # typed and it wins from the moment they press Save — the overlay rule
        # hub/client_urls.py works to. See hub/scan_facts.contact_suggestions.
        prof["suggested"] = scan_facts.contact_suggestions(
            prof, (request.args.get("domain") or "").strip())
        # Notes are shared across a group; contacts, address and category are
        # not. Two brands of one parent company have their own front desk, and
        # showing one company's phone number under the other's name is a
        # confidently wrong answer of the kind grouping exists to avoid.
        others = [n for n in client_groups.member_names(name)
                  if n and n.strip().lower() != name.lower()]
        if others:
            notes = [dict(n) for n in (prof.get("notes") or [])]
            for other in others:
                for n in (seo.get_profile(other).get("notes") or []):
                    row = dict(n)
                    row["member"] = other
                    notes.append(row)
            notes.sort(key=lambda n: str(n.get("time") or ""), reverse=True)
            prof["notes"] = notes[:200]
            prof["group"] = others
        return jsonify({"profile": prof})

    @app.route("/api/client/profile", methods=["POST"])
    def api_client_profile_set():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        prof = seo.set_profile(client, body)
        audit.log("hub", "client_profile_saved", actor=current_user(), detail=client)
        return jsonify({"ok": True, "profile": prof})

    # ------------- client groups: one company, several client records
    # Why this exists, and every rule it enforces: hub/client_groups.py.
    @app.route("/api/client/group")
    def api_client_group():
        gate = _require_api()
        if gate:
            return gate
        from . import client_groups
        name = (request.args.get("name") or "").strip()
        url = (request.args.get("url") or "").strip()
        return jsonify({"roster": client_groups.roster(name, url) if name else {},
                        "groups": client_groups.groups()})

    @app.route("/api/client/group/add", methods=["POST"])
    def api_client_group_add():
        gate = _require_api()
        if gate:
            return gate
        from . import client_groups
        body = request.get_json(silent=True) or {}
        res = client_groups.add_member(
            str(body.get("parent") or ""), str(body.get("member") or ""),
            parent_url=str(body.get("parent_url") or ""),
            member_url=str(body.get("member_url") or ""),
            actor=current_user() or "")
        if res.get("error"):
            return jsonify(res), 400
        audit.log("hub", "client_grouped", actor=current_user(),
                  client=str(body.get("parent") or ""),
                  detail=f"{body.get('member')} grouped under {body.get('parent')}")
        return jsonify(res)

    @app.route("/api/client/group/remove", methods=["POST"])
    def api_client_group_remove():
        gate = _require_api()
        if gate:
            return gate
        from . import client_groups
        body = request.get_json(silent=True) or {}
        name = str(body.get("client") or "")
        res = client_groups.remove_member(name, str(body.get("url") or ""),
                                          actor=current_user() or "")
        if res.get("error"):
            return jsonify(res), 400
        audit.log("hub", "client_ungrouped", actor=current_user(),
                  client=name, detail=f"{name} removed from its group")
        return jsonify(res)

    @app.route("/api/client/notes", methods=["POST"])
    def api_client_notes():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        text = (body.get("text") or "").strip()
        if not client or not text:
            return jsonify({"error": "client and text are required."}), 400
        prof = seo.add_note(client, text, author=current_user() or "")
        audit.log("hub", "client_note_added", actor=current_user(), detail=client)
        return jsonify({"ok": True, "profile": prof})

    @app.route("/api/client/tickets")
    def api_client_tickets():
        gate = _require_api()
        if gate:
            return gate
        from . import knack_api
        name = (request.args.get("name") or "").strip()
        website = (request.args.get("website") or "").strip()
        if not knack_api.configured():
            return jsonify({"configured": False, "tickets": []})
        try:
            return jsonify({"configured": True,
                            "tickets": knack_api.list_tickets(name, website)})
        except Exception as exc:  # noqa: BLE001
            errors.log_exception("knack-tickets", exc, path=request.path,
                                 actor=current_user() or "")
            return jsonify({"configured": True, "tickets": [], "error": str(exc)})

    @app.route("/api/client/requests/triage", methods=["POST"])
    def api_client_request_triage():
        """Read the description and propose the choices it already answers.

        One route for all three objects, because `ticket_form_fields()`,
        `campaign_form_fields()` and `ad_copy.form_fields()` hand back the
        same shape and three copies of this would be three descriptions of
        what a suggestion is. `hub/static/knack-form.js` draws the control
        once for the same reason; the third form went a release without it
        because only the browser half was shared and this half knew two
        kinds.

        A POST, and only ever into the fields the caller says are **empty**: a
        value somebody chose is the better source and is never offered over,
        and the gate is here rather than in the browser because a rule the
        form keeps while the endpoint breaks it is not a rule. Nothing is
        written — the suggestion is drawn dotted and one press keeps it.

        An unrecognised kind is **refused by name** rather than falling
        through to whichever branch is last. It used to read `if ticket ...
        else campaign`, so a typo in the caller answered with the campaign
        change form's dropdowns against an ad copy request's prose — every
        suggestion then either dropped for not being one of the field's
        options or, worse, kept for a field of the same name on a different
        object. Both look like a button that half works.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import ad_copy, knack_api, request_triage
        body = request.get_json(silent=True) or {}
        kind = str(body.get("kind") or "ticket")
        # The second half of each pair is the tag hub/ai.py bills the call
        # under -- it rides in the `ai` log as `tool`, and spend_by_module()
        # splits the provider-spend audit on it. Left at "tickets" for all
        # three, every triage call this Hub ever makes reads as the ticket
        # form's, which is a confident wrong answer on the one page that
        # says what the models are costing us.
        READERS = {
            "ticket": (lambda: knack_api.ticket_form_fields("create"),
                       "tickets"),
            "support": (lambda: knack_api.campaign_form_fields("support"),
                        "campaign_support"),
            "change": (lambda: knack_api.campaign_form_fields("change"),
                       "campaign_support"),
            "adcopy": (ad_copy.form_fields, "ad_copy"),
        }
        if kind not in READERS:
            return jsonify({"ok": False, "suggestions": {}, "unusable": 0,
                            "error": f"There is no form called {kind!r}.",
                            "note": ""})
        if not knack_api.configured():
            return jsonify({"ok": False, "suggestions": {}, "unusable": 0,
                            "error": "Knack is not configured, so the "
                                     "options a field accepts cannot be read.",
                            "note": ""})
        if kind == "adcopy":
            # Discovered from its field ids rather than pinned, and a
            # discovery that failed already knows how to say so — better than
            # reading an object with no fields and reporting nothing to
            # suggest, which is what "we could not find the form" looks like
            # from the outside.
            _obj, why = ad_copy.resolve()
            if not _obj:
                return jsonify({"ok": False, "suggestions": {}, "unusable": 0,
                                "error": why, "note": ""})
        reader, module = READERS[kind]
        try:
            fields = reader()
        except Exception as exc:                        # noqa: BLE001
            # Named, not answered with an empty list: "this field accepts
            # nothing" and "we could not read what it accepts" are different,
            # and only the second is somebody's to fix.
            return jsonify({"ok": False, "suggestions": {}, "unusable": 0,
                            "error": f"The live object could not be read "
                                     f"({type(exc).__name__}).", "note": ""})
        out = request_triage.suggest(
            body.get("text") or "", fields, body.get("empty") or [],
            module=module)
        return jsonify(out)

    @app.route("/api/client/tickets/fields")
    def api_client_ticket_fields():
        """The controls a ticket form should draw, from the live object.

        The form asks for this rather than carrying its own copy of the field
        list: the ids are ours, but the dropdown choices are Knack's, and a
        form that guesses a choice writes a value Knack refuses — which loses
        the whole ticket, not the one field.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import knack_api
        scope = "manage" if (request.args.get("scope") or "") == "manage" else "create"
        if not knack_api.configured():
            return jsonify({"configured": False, "fields": []})
        try:
            fields = knack_api.ticket_form_fields(scope)
            # A ticket raised from a client record should not ask for the
            # website that record has held since their last site scan. One
            # reader, shared with the campaign request and the ad copy form:
            # three copies of this mapping is how two of them come to offer a
            # different answer for one client. Offered into empty fields only,
            # and nothing is written until the ticket is created.
            from .client_context import offer_into
            values, notes = offer_into(fields, {},
                                       request.args.get("client", ""),
                                       request.args.get("url", ""))
            return jsonify({"configured": True, "scope": scope,
                            "object": knack_api.TICKETS_OBJECT,
                            "fields": fields,
                            "values": values, "notes": notes})
        except Exception as exc:  # noqa: BLE001
            errors.log_exception("knack-tickets", exc, path=request.path,
                                 actor=current_user() or "")
            return jsonify({"configured": True, "fields": [], "error": str(exc)})

    @app.route("/api/client/tickets", methods=["POST"])
    def api_client_tickets_create():
        gate = _require_api()
        if gate:
            return gate
        from . import knack_api
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        subject = (body.get("subject") or "").strip()
        if not client or not subject:
            return jsonify({"error": "client and subject are required."}), 400
        if not knack_api.configured():
            return jsonify({"error": "Knack isn't configured — set KNACK_APP_ID and "
                                     "KNACK_API_KEY, then redeploy."}), 400
        try:
            rec = knack_api.create_ticket(
                client, (body.get("website") or "").strip(), subject,
                (body.get("description") or "").strip(),
                author=current_user() or "",
                requested_by=(body.get("requested_by") or "").strip(),
                values=body.get("values") or {})
        except Exception as exc:  # noqa: BLE001
            errors.log_exception("knack-tickets", exc, path=request.path,
                                 actor=current_user() or "")
            return jsonify({"error": str(exc)})
        audit.log("hub", "web_ticket_created", actor=current_user(),
                  detail=f"{client}: {subject[:60]}")
        # `rejected` is reported, never swallowed: a ticket that was created
        # with half its fields missing must not read as a clean success.
        return jsonify({"ok": True, "id": rec.get("id"),
                        "written": rec.get("written") or [],
                        "rejected": rec.get("rejected") or []})

    @app.route("/api/client/tickets/update", methods=["POST"])
    def api_client_tickets_update():
        """Manage Ticket: edit an existing ticket's fields.

        The record id travels in the body rather than the path so the URL
        stays a literal tools/linkcheck.py can verify — a path built by
        concatenation is one nothing checks (see CLAUDE.md).
        """
        gate = _require_api()
        if gate:
            return gate
        from . import knack_api
        body = request.get_json(silent=True) or {}
        record_id = (body.get("id") or "").strip()
        values = body.get("values") or {}
        if not record_id or not isinstance(values, dict) or not values:
            return jsonify({"error": "id and values are required."}), 400
        if not knack_api.configured():
            return jsonify({"error": "Knack isn't configured — set KNACK_APP_ID and "
                                     "KNACK_API_KEY, then redeploy."}), 400
        try:
            res = knack_api.update_ticket(record_id, values)
        except Exception as exc:  # noqa: BLE001
            errors.log_exception("knack-tickets", exc, path=request.path,
                                 actor=current_user() or "")
            return jsonify({"error": str(exc)})
        if res.get("ok"):
            audit.log("hub", "web_ticket_updated", actor=current_user(),
                      detail=f"{record_id}: {', '.join(sorted(values))[:80]}")
        return jsonify(res)

    @app.route("/api/knack/campaign-fields")
    def api_knack_campaign_fields():
        """The controls a campaign request form should draw, from the live
        object — plus the field mapping shown before anything is written.

        The form asks for this rather than carrying its own copy: the support
        ids are ours, but the dropdown choices, the connection records and the
        field types are Knack's, and a form that guesses one writes a value
        Knack refuses — which loses the whole request, not the one field.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import knack_api
        kind = (request.args.get("kind") or "").strip()
        if kind not in ("change", "support"):
            return jsonify({"error": "kind must be change or support."}), 400
        if not knack_api.configured():
            return jsonify({"configured": False, "fields": []})
        try:
            info = knack_api.campaign_field_map(kind)
            fields = knack_api.campaign_form_fields(kind)
            # Same one reader as the web ticket above.
            from .client_context import offer_into
            values, notes = offer_into(fields, {},
                                       request.args.get("client", ""),
                                       request.args.get("url", ""))
            return jsonify({"configured": True, "kind": kind, **info,
                            "fields": fields,
                            "values": values, "notes": notes})
        except Exception as exc:  # noqa: BLE001
            errors.log_exception("knack-campaign", exc, path=request.path,
                                 actor=current_user() or "")
            return jsonify({"configured": True, "fields": [], "error": str(exc)})

    @app.route("/api/knack/people")
    def api_knack_people():
        """Names from object_161 + object_109 for Requested By dropdowns."""
        gate = _require_api()
        if gate:
            return gate
        from . import knack_api
        if not knack_api.configured():
            return jsonify({"configured": False, "names": []})
        try:
            return jsonify({"configured": True, "names": knack_api.people_names()})
        except Exception as exc:  # noqa: BLE001
            errors.log_exception("knack-people", exc, path=request.path,
                                 actor=current_user() or "")
            return jsonify({"configured": True, "names": [], "error": str(exc)})

    @app.route("/api/client/campaign-request", methods=["POST"])
    def api_client_campaign_request():
        gate = _require_api()
        if gate:
            return gate
        from . import knack_api
        body = request.get_json(silent=True) or {}
        kind = (body.get("kind") or "").strip()
        client = (body.get("client") or "").strip()
        subject = (body.get("subject") or "").strip()
        if kind not in ("change", "support") or not client or not subject:
            return jsonify({"error": "kind (change|support), client and subject are required."}), 400
        if not knack_api.configured():
            return jsonify({"error": "Knack isn't configured — set KNACK_APP_ID and "
                                     "KNACK_API_KEY, then redeploy."}), 400
        values = body.get("values")
        if values is not None and not isinstance(values, dict):
            return jsonify({"error": "values must be an object."}), 400
        try:
            rec = knack_api.create_campaign_request(
                kind, client, (body.get("campaign") or "").strip(),
                (body.get("io") or "").strip(), subject,
                (body.get("description") or "").strip(),
                author=current_user() or "",
                requested_by=(body.get("requested_by") or "").strip(),
                values=values or {})
        except Exception as exc:  # noqa: BLE001
            errors.log_exception("knack-campaign", exc, path=request.path,
                                 actor=current_user() or "")
            return jsonify({"error": str(exc)})
        audit.log("hub", f"campaign_{kind}_request", actor=current_user(),
                  detail=f"{client}: {subject[:60]}")
        # What reached Knack and what did not. A request created with half its
        # fields missing must not read as a clean success.
        return jsonify({"ok": True, "id": rec.get("id"),
                        "written": rec.get("written") or [],
                        "rejected": rec.get("rejected") or []})

    # ------------------------------------------------- Ad Copy Request
    # Ad Copy used to be a Campaign Change Request with its subject already
    # written: four boxes, and a rep retyping the client, the campaign, the
    # order number and the media partner from the record on the screen
    # behind it. hub/ad_copy.py is the object itself — pinned ids, the
    # controls read off the live schema, and everything the client's own
    # insertion orders can already answer.
    @app.route("/api/client/ad-copy/fields")
    def api_ad_copy_fields():
        gate = _require_api()
        if gate:
            return gate
        from . import ad_copy
        acct = current_account()
        try:
            data = ad_copy.form(
                request.args.get("client", ""),
                # The seller and the confirmation address are the signed-in
                # *account's*, not a box to retype — and deliberately not
                # `current_user()`, which answers "Shared login" for a
                # PANEL_PASSWORD session. That is a true statement about the
                # session and a wrong one in the Seller Name box, which the
                # campaign team reads as a person. With no account behind the
                # session the form falls back to the rep named on this
                # client's own orders, and says so when there isn't one.
                user_name=(getattr(acct, "name", "") or ""),
                user_email=(getattr(acct, "email", "") or ""))
        except Exception as exc:  # noqa: BLE001
            errors.log_exception("ad-copy-fields", exc, path=request.path,
                                 actor=current_user() or "")
            return jsonify({"configured": True, "error": str(exc),
                            "fields": [], "values": {}, "options": {},
                            "notes": []})
        return jsonify(data)

    @app.route("/api/client/ad-copy", methods=["POST"])
    def api_ad_copy_create():
        gate = _require_api()
        if gate:
            return gate
        from . import ad_copy, knack_api
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        values = body.get("values") or {}
        if not client or not isinstance(values, dict):
            return jsonify({"error": "client and values are required."}), 400
        if not knack_api.configured():
            return jsonify({"error": "Knack isn't configured — set KNACK_APP_ID "
                                     "and KNACK_API_KEY, then redeploy."}), 400
        try:
            rec = ad_copy.create(client, values, author=current_user() or "")
        except Exception as exc:  # noqa: BLE001
            errors.log_exception("ad-copy", exc, path=request.path,
                                 actor=current_user() or "")
            return jsonify({"error": str(exc)})
        # Logged under `ad_copy`, not `hub`: this is work filed against a
        # client, and `client_brand.work_log()` skips a module its own table
        # cannot name — which reads on the record as a client nobody has done
        # anything for. `WORK_KINDS` carries the entry.
        audit.log("ad_copy", "request_created", actor=current_user(),
                  client=client,
                  detail=f"{client}: {str(values.get('campaign') or '')[:60]}")
        return jsonify({"ok": True, "id": rec.get("id"),
                        "written": rec.get("written") or [],
                        "rejected": rec.get("rejected") or []})

    @app.route("/api/client/website-hosted", methods=["POST"])
    def api_client_website_hosted():
        """Whether Smart 1 Marketing hosts this site.

        Stored through the existing website-override mechanism rather than a
        parallel store: overrides are already merged into every website dict
        on read, are already scoped per domain, and are already documented as
        hub-side corrections that never write back to Knack. A second store
        would need its own merge step and would drift.
        """
        gate = _require_api()
        if gate:
            return gate
        body = request.get_json(silent=True) or {}
        client = str(body.get("client") or "").strip()
        domain = str(body.get("domain") or "").strip()
        value = str(body.get("s1m_hosted") or "").strip().lower()
        if value not in ("", "yes", "no"):
            return jsonify({"error": "Value must be yes, no or blank."}), 400
        if not client:
            return jsonify({"error": "No client given."}), 400
        try:
            from . import seo
            seo.set_website_override(client, domain, {"s1m_hosted": value})
            audit.log("hub", "s1m_hosted_set", actor=current_user(),
                      client=client, domain=domain, value=value or "cleared")
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": f"Could not save ({type(exc).__name__})."}), 500
        return jsonify({"ok": True, "s1m_hosted": value})

    @app.route("/api/client/website-platform", methods=["POST"])
    def api_client_website_platform():
        """Hub-only correction of a website record's platform — Knack untouched."""
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        domain = (body.get("domain") or "").strip()
        platform = (body.get("platform") or "").strip()
        if not client or not domain or not platform:
            return jsonify({"error": "client, domain and platform are required."}), 400
        try:
            seo.set_website_override(client, domain, {"platform": platform})
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        audit.log("hub", "website_platform_corrected", actor=current_user(),
                  detail=f"{client}: {domain} -> {platform}")
        return jsonify({"ok": True})

    @app.route("/api/websites/search")
    def api_websites_search():
        """Search the websites inventory — used to attach extra website
        records to a client (hub-only, never written back to Knack)."""
        gate = _require_api()
        if gate:
            return gate
        q = (request.args.get("q") or "").strip().lower()
        if not q:
            return jsonify({"results": []})
        out = []
        for w in knack_data.websites():
            hay = " ".join(str(w.get(k) or "") for k in ("name", "domain")).lower()
            if q in hay:
                out.append({"name": w.get("name"), "domain": w.get("domain"),
                            "platform": w.get("platform"), "status": w.get("status")})
            if len(out) >= 8:
                break
        return jsonify({"results": out})

    # ---------------- SEO blogs ----------------
    @app.route("/api/seo/blogs")
    def api_seo_blogs():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        name = (request.args.get("name") or "").strip()
        store = seo.load_store(name)
        blogs = store.get("blogs", {})
        # Archived posts stay in the store but leave the working list.
        _posts = [p for p in blogs.get("posts", []) if not p.get("archived")]
        return jsonify({"posts": _posts,
                        "archived": sum(1 for p in blogs.get("posts", [])
                                        if p.get("archived")),
                        "focus": blogs.get("focus", ""),
                        "questions": blogs.get("questions", []),
                        "settings": seo.blog_settings(name, store),
                        "site_url": seo.client_site_url(name, store),
                        "frequency": store.get("setup", {}).get("blogs_frequency", ""),
                        "per_month": store.get("setup", {}).get("blogs_per_month", "")})

    @app.route("/api/seo/blogs/plan", methods=["POST"])
    def api_seo_blogs_plan():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        try:
            out = seo.blog_plan(client, (body.get("focus") or "").strip(),
                                clamp_int(body.get("months"), 3, 1, 24),
                                (body.get("start") or "").strip())
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": str(exc)})
        audit.log("hub", "seo_blog_plan", actor=current_user(),
                  detail=f"{client}: {len(out['posts'])} posts")
        from . import seo_tasks
        out["tasks"] = seo_tasks.for_posts(client, out.get("posts") or [],
                                           actor=current_user())
        return jsonify(out)

    @app.route("/api/seo/blogs/write", methods=["POST"])
    def api_seo_blogs_write():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        ids = [int(i) for i in (body.get("ids") or []) if str(i).isdigit()]
        if not client or not ids:
            return jsonify({"error": "client and ids are required."}), 400
        try:
            out = seo.blog_write(client, ids)
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": str(exc)})
        audit.log("hub", "seo_blog_write", actor=current_user(),
                  detail=f"{client}: {len(out['written'])} posts")
        return jsonify(out)

    @app.route("/api/seo/blogs/update", methods=["POST"])
    def api_seo_blogs_update():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        pid = body.get("id")
        store = seo.load_store(client)
        post = next((p for p in store.get("blogs", {}).get("posts", [])
                     if p["id"] == pid), None)
        if not client or post is None:
            return jsonify({"error": "Unknown client or post."}), 404
        if isinstance(body.get("title"), str) and body["title"].strip():
            post["title"] = body["title"].strip()
        if isinstance(body.get("content"), str):
            post["content"] = body["content"]
            if body["content"].strip():
                post["status"] = "written"
        if "posted" in body:
            post["posted"] = bool(body["posted"])
        if "archived" in body:
            # Archive hides a post from the working list without deleting it.
            # A deleted post takes its written content and its image with it,
            # which is rarely what "I'm done with this" means.
            post["archived"] = bool(body["archived"])
        from . import blog_spec
        if "categories" in body or "tags" in body:
            known = seo.blog_settings(client, store)["categories"]
            tax = blog_spec.clamp_taxonomy(
                body.get("categories", post.get("categories")),
                body.get("tags", post.get("tags")), known)
            post["categories"], post["tags"] = tax["categories"], tax["tags"]
            blogs = store.setdefault("blogs", {})
            blogs["categories"] = blog_spec.merge_categories(known, tax["categories"])
        if isinstance(body.get("slug"), str) and body["slug"].strip():
            post["slug"] = blog_spec.slugify_title(body["slug"])
        # Editing the copy re-runs the never-mention check rather than leaving
        # the flag from the version that has just been replaced — stale either
        # way is wrong, and a cleared flag on copy that still says it is worse.
        if isinstance(body.get("content"), str):
            post["flags"] = blog_spec.scan_forbidden(
                post.get("content", "") + " " + str(post.get("meta_description") or ""),
                seo.blog_settings(client, store)["avoid"])
        if isinstance(body.get("answers"), dict):
            blogs = store.setdefault("blogs", {})
            blogs.setdefault("answers", {}).update(
                {k: v for k, v in body["answers"].items() if str(v).strip()})
            blogs["questions"] = [q for q in blogs.get("questions", [])
                                  if q not in blogs["answers"]]
        seo.save_store(client, store)
        return jsonify({"ok": True})

    @app.route("/api/seo/blogs/answers", methods=["POST"])
    def api_seo_blogs_answers():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        store = seo.load_store(client)
        blogs = store.setdefault("blogs", {})
        if isinstance(body.get("answers"), dict):
            blogs.setdefault("answers", {}).update(
                {k: v for k, v in body["answers"].items() if str(v).strip()})
            blogs["questions"] = [q for q in blogs.get("questions", [])
                                  if q not in blogs["answers"]]
        seo.save_store(client, store)
        return jsonify({"ok": True, "questions": blogs.get("questions", [])})

    @app.route("/api/seo/blogs/tag", methods=["POST"])
    def api_seo_blogs_tag():
        """Fill in categories and tags on posts planned before they existed."""
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        ids = [int(i) for i in (body.get("ids") or []) if str(i).isdigit()]
        try:
            out = seo.blog_tag_posts(client, ids or None)
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": str(exc)})
        audit.log("hub", "seo_blog_tag", actor=current_user(), client=client,
                  detail=f"{out['tagged']} posts")
        return jsonify(out)

    @app.route("/api/seo/blogs/settings", methods=["POST"])
    def api_seo_blogs_settings():
        """The default author, the guardrail text and the never-mention list."""
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        updates = {k: body[k] for k in
                   ("author", "guidance", "avoid", "categories", "approved_only")
                   if k in body}
        settings = seo.save_blog_settings(client, updates)
        audit.log("hub", "seo_blog_settings", actor=current_user(),
                  client=client, detail=", ".join(sorted(updates)) or "no change")
        return jsonify({"ok": True, "settings": settings})

    @app.route("/api/seo/blogs/topics", methods=["POST"])
    def api_seo_blogs_topics():
        """Load the topic list the client already approved.

        Takes an uploaded PDF/DOCX/text file, or pasted text. The parsed list
        is returned in full rather than as a count: a thirty-topic document
        that parsed into three needs to be seen to be caught.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = ((request.form.get("client") if request.form else "")
                  or body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        action = ((request.form.get("action") if request.form else "")
                  or body.get("action") or "load").strip()
        if action == "clear":
            return jsonify(seo.clear_approved_topics(client))
        append = str((request.form.get("append") if request.form else "")
                     or body.get("append") or "").lower() in ("1", "true", "yes")
        up = request.files.get("file") if request.files else None
        text, filename = "", ""
        if up and up.filename:
            filename = up.filename
            text = _read_document(up.read(8 * 1024 * 1024), filename)
            if not text.strip():
                return jsonify({"error": "Couldn't read any text from that "
                                         "file. A scanned PDF has no text "
                                         "layer to read."}), 400
        else:
            text = str(body.get("text") or (request.form.get("text") if request.form else "") or "")
            filename = str(body.get("filename") or "pasted list")
        if not text.strip():
            return jsonify({"error": "Upload a document or paste the topics."}), 400
        out = seo.set_approved_topics(client, text, filename, append=append)
        audit.log("hub", "seo_blog_topics", actor=current_user(), client=client,
                  detail=f"{out['found']} topics from {filename}")
        return jsonify(out)

    @app.route("/api/seo/publish/instructions", methods=["POST"])
    def api_seo_publish_instructions():
        """The prompt the rep pastes into Claude in Chrome, for one CMS.

        Neither CMS has a write API we can use — see hub/cms_publish.py — so a
        browser agent driving the admin IS the publishing path, and this
        endpoint's job is to hand it everything it needs in one block of text.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import cms_publish, seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        cms = (body.get("cms") or "").strip()
        kind = (body.get("kind") or "blogs").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        if cms not in cms_publish.CMS_KEYS:
            return jsonify({"error": f"Unknown CMS '{cms}'."}), 400
        if kind not in cms_publish.KINDS:
            return jsonify({"error": f"Unknown content kind '{kind}'."}), 400
        store = seo.load_store(client)
        site = seo.client_site_url(client, store)
        wanted = [str(u) for u in (body.get("urls") or [])]
        settings = None

        if kind == "schema":
            pages = store.get("pages", {})
            types = {r["url"]: r["types"] for r in seo.schema_pages_table(client)}
            chosen = [dict(pages[u], url=u, types=types.get(u, []))
                      for u in wanted if u in pages]
        elif kind == "faqs":
            from . import faq as _faq
            chosen = []
            for url in wanted:
                page = _faq.get_page(client, url)
                if page is None:
                    continue
                # The accordion travels with the questions: it is what
                # actually goes on the page, and it carries its own FAQPage
                # schema so the agent is not asked to paste two things.
                chosen.append({"url": url, "questions": page.get("questions", []),
                               "html": _faq.accordion_html(client, [url])})
        elif kind == "alt":
            from . import alt_text
            chosen = alt_text.selected_pages(client, wanted or None)
        else:
            ids = [int(i) for i in (body.get("ids") or []) if str(i).isdigit()]
            posts = {p["id"]: p for p in store.get("blogs", {}).get("posts", [])}
            chosen = [posts[i] for i in ids if i in posts]
            settings = seo.blog_settings(client, store)

        out = cms_publish.instructions(cms, kind, chosen, client=client,
                                       site_url=site, settings=settings,
                                       placement=str(body.get("placement") or ""))
        audit.log("hub", "seo_publish_instructions", actor=current_user(),
                  client=client, detail=f"{kind} → {cms}: {len(out.get('items', []))}")
        return jsonify(out)

    # ---------------- SEO alt text ----------------
    @app.route("/api/seo/alt")
    def api_seo_alt():
        gate = _require_api()
        if gate:
            return gate
        from . import alt_text
        return jsonify(alt_text.load((request.args.get("name") or "").strip()))

    @app.route("/api/seo/alt/scan", methods=["POST"])
    def api_seo_alt_scan():
        """Read the first N sitemap pages and list every image on them."""
        gate = _require_api()
        if gate:
            return gate
        from . import alt_text
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        limit = clamp_int(body.get("pages"), alt_text.DEFAULT_PAGES,
                          1, alt_text.MAX_PAGES)
        try:
            out = alt_text.scan(client, limit,
                                [str(u) for u in (body.get("urls") or [])])
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": str(exc)})
        audit.log("hub", "seo_alt_scan", actor=current_user(), client=client,
                  detail=f"{len(out['pages'])} pages, {out['total_images']} images")
        return jsonify(out)

    @app.route("/api/seo/alt/write", methods=["POST"])
    def api_seo_alt_write():
        gate = _require_api()
        if gate:
            return gate
        from . import alt_text
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        try:
            out = alt_text.rewrite(client, [str(u) for u in (body.get("urls") or [])])
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": str(exc)})
        if not out.get("error"):
            audit.log("hub", "seo_alt_write", actor=current_user(), client=client,
                      detail=f"{out.get('written', 0)} images")
        return jsonify(out)

    @app.route("/api/seo/alt/update", methods=["POST"])
    def api_seo_alt_update():
        gate = _require_api()
        if gate:
            return gate
        from . import alt_text
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        img = alt_text.set_alt(client, str(body.get("url") or ""),
                               str(body.get("src") or ""),
                               str(body.get("alt") or ""))
        if img is None:
            return jsonify({"error": "That image is not in the last scan."}), 404
        return jsonify({"ok": True, "image": img})

    @app.route("/seo/alt/<slug>/code.html")
    def seo_alt_code(slug):
        """The rewritten alt text as markup, old tag and new."""
        gate = _require_page()
        if gate:
            return gate
        from . import alt_text, seo
        match = next((c for c in seo.seo_clients() if c["slug"] == slug), None)
        name = match["client"] if match else slug.replace("-", " ")
        raw = (request.args.get("urls") or "").strip()
        urls = [u for u in raw.split("\n") if u.strip()] if raw else None
        body = alt_text.code_view(name, urls)
        resp = make_response(body)
        resp.headers["Content-Type"] = "text/plain; charset=utf-8"
        resp.headers["Content-Disposition"] = f'attachment; filename="{slug}-alt-text.html"'
        return resp

    @app.route("/seo/blogs/<slug>.doc")
    def seo_blogs_doc(slug):
        gate = _require_page()
        if gate:
            return gate
        from . import seo
        match = next((c for c in seo.seo_clients() if c["slug"] == slug), None)
        name = match["client"] if match else slug.replace("-", " ")
        raw = (request.args.get("ids") or "").strip()
        ids = None
        if raw and raw.lower() != "all":
            ids = [int(x) for x in raw.split(",") if x.strip().isdigit()]
        body = seo.blogs_doc(name, ids)
        resp = make_response(body)
        resp.headers["Content-Type"] = "application/msword"
        resp.headers["Content-Disposition"] = f'attachment; filename="{slug}-blogs.doc"'
        return resp

    @app.route("/seo/blogs/<slug>/view")
    def seo_blogs_view(slug):
        gate = _require_page()
        if gate:
            return gate
        from . import seo
        match = next((c for c in seo.seo_clients() if c["slug"] == slug), None)
        name = match["client"] if match else slug.replace("-", " ")
        raw = (request.args.get("ids") or "").strip()
        ids = None
        if raw and raw.lower() != "all":
            ids = [int(x) for x in raw.split(",") if x.strip().isdigit()]
        return app.response_class(seo.blogs_doc(name, ids), mimetype="text/html")

    @app.route("/api/seo/compiled")
    def api_seo_compiled():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        name = (request.args.get("name") or "").strip()
        return jsonify(seo.compiled_json(name))

    @app.route("/seo/download/<slug>.<fmt>")
    def seo_download(slug, fmt):
        gate = _require_page()
        if gate:
            return gate
        from . import seo
        match = next((c for c in seo.seo_clients() if c["slug"] == slug), None)
        name = match["client"] if match else slug.replace("-", " ")
        if fmt == "html":
            body = seo.compiled_html(name)
            resp = make_response(body)
            resp.headers["Content-Type"] = "text/plain; charset=utf-8"
        else:
            resp = make_response(json.dumps(seo.compiled_json(name), indent=1))
            resp.headers["Content-Type"] = "application/json"
        resp.headers["Content-Disposition"] = f'attachment; filename="{slug}-schema.{fmt}"'
        return resp

    def _client_from_slug(slug: str) -> str:
        from . import seo
        match = next((c for c in seo.seo_clients() if c["slug"] == slug), None)
        return match["client"] if match else slug.replace("-", " ")

    def _urls_param() -> list[str]:
        """?urls=<one per line> — which saved pages a download covers."""
        raw = request.args.get("urls") or ""
        return [u.strip() for u in raw.split("\n") if u.strip()]

    # ---------------- business info: what we know, then GMB ----------------
    @app.route("/api/seo/enrich", methods=["POST"])
    def api_seo_enrich():
        """Complete a client's business info — Hub records first, then a
        Google Business Profile lookup for whatever is still blank."""
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        if not client:
            return jsonify({"error": "client is required."}), 400
        try:
            out = seo.enrich_business_info(client, force=bool(body.get("force")))
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": str(exc)}), 500
        if out.get("filled"):
            audit.log("hub", "seo_business_enriched", actor=current_user(),
                      detail=client, fields=", ".join(out["filled"]))
        return jsonify(out)

    # ---------------- saved schema pages: table, edit, delete ----------------
    @app.route("/api/seo/schema/pages")
    def api_seo_schema_pages():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        name = (request.args.get("name") or "").strip()
        return jsonify({"pages": seo.schema_pages_table(name)})

    @app.route("/api/seo/schema/page", methods=["POST"])
    def api_seo_schema_page_update():
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        url = (body.get("url") or "").strip()
        if not client or not url:
            return jsonify({"error": "client and url are required."}), 400
        if body.get("delete"):
            ok = seo.delete_page(client, url)
            return jsonify({"ok": ok, "pages": seo.schema_pages_table(client)})
        page = seo.update_page_meta(client, url, body)
        if page is None:
            return jsonify({"error": "That page is not saved for this client."}), 404
        return jsonify({"ok": True, "page": page,
                        "pages": seo.schema_pages_table(client)})

    @app.route("/api/seo/schema/detail")
    def api_seo_schema_detail():
        """Full JSON-LD for one saved page — powers the view/edit modal."""
        gate = _require_api()
        if gate:
            return gate
        from . import seo
        name = (request.args.get("name") or "").strip()
        url = (request.args.get("url") or "").strip()
        page = seo.load_store(name).get("pages", {}).get(url)
        if not page:
            return jsonify({"error": "Not found"}), 404
        return jsonify({"page": page})

    @app.route("/seo/schema/<slug>/download.<fmt>")
    def seo_schema_download(slug, fmt):
        """Schema for one page, the selected pages, or all of them."""
        gate = _require_page()
        if gate:
            return gate
        from . import seo
        name = _client_from_slug(slug)
        urls = _urls_param()
        wanted = {u.rstrip("/") for u in urls}
        suffix = "page" if len(urls) == 1 else ("selected" if urls else "all")
        store = seo.load_store(name)
        if fmt == "doc":
            body = seo.schema_doc(name, urls or None)
            ctype, ext = "application/msword", "doc"
        elif fmt == "json":
            picked = {u: p.get("schema") for u, p in store.get("pages", {}).items()
                      if not urls or u.rstrip("/") in wanted}
            body = json.dumps({"client": name, "pages": len(picked),
                               "schemas": picked}, indent=1)
            ctype, ext = "application/json", "json"
        else:
            out = [f"<!-- JSON-LD schema for {name} — generated by Smart 1 Hub -->"]
            for u, p in store.get("pages", {}).items():
                if urls and u.rstrip("/") not in wanted:
                    continue
                out.append(f"\n<!-- ===== {u} ===== -->")
                out.append('<script type="application/ld+json">')
                out.append(json.dumps(p.get("schema"), indent=1))
                out.append("</script>")
            body = "\n".join(out)
            ctype, ext = "text/plain; charset=utf-8", "html"
        resp = make_response(body)
        resp.headers["Content-Type"] = ctype
        resp.headers["Content-Disposition"] = \
            f'attachment; filename="{slug}-schema-{suffix}.{ext}"'
        return resp

    # ---------------------------- FAQ Builder ----------------------------
    @app.route("/api/seo/faq/generate", methods=["POST"])
    def api_faq_generate():
        gate = _require_api()
        if gate:
            return gate
        from . import faq
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        url = (body.get("url") or "").strip()
        if not client or not url:
            return jsonify({"error": "client and url are required."}), 400
        count = clamp_int(body.get("count"), 6, 1, 50)
        avoid = body.get("avoid") if isinstance(body.get("avoid"), list) else []
        try:
            return jsonify(faq.generate(client, url, count=count, avoid=avoid))
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": str(exc)}), 500

    @app.route("/api/seo/faq/pages")
    def api_faq_pages():
        gate = _require_api()
        if gate:
            return gate
        from . import faq
        name = (request.args.get("name") or "").strip()
        return jsonify({"pages": faq.list_pages(name)})

    @app.route("/api/seo/faq/save", methods=["POST"])
    def api_faq_save():
        gate = _require_api()
        if gate:
            return gate
        from . import faq
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        url = (body.get("url") or "").strip()
        if not client or not url:
            return jsonify({"error": "client and url are required."}), 400
        try:
            page = faq.save_page(client, url, body.get("questions") or [],
                                 added_to_site=body.get("added_to_site", ""),
                                 title=body.get("title", ""),
                                 style=body.get("style"))
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        audit.log("hub", "faq_page_saved", actor=current_user(), detail=client,
                  url=url, questions=len(page.get("questions", [])))
        # The FAQ exists in the Hub; somebody still has to put it on the page.
        # Raising the ticket must not be able to fail the save that earned it.
        from . import seo_tasks
        task = seo_tasks.for_faq(client, url, body.get("title", ""),
                                 actor=current_user())
        return jsonify({"ok": True, "page": page, "pages": faq.list_pages(client),
                        "task": task})

    @app.route("/api/seo/faq/page", methods=["POST"])
    def api_faq_page_update():
        gate = _require_api()
        if gate:
            return gate
        from . import faq
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        url = (body.get("url") or "").strip()
        if not client or not url:
            return jsonify({"error": "client and url are required."}), 400
        if body.get("delete"):
            ok = faq.delete_page(client, url)
            return jsonify({"ok": ok, "pages": faq.list_pages(client)})
        page = faq.update_page(client, url, body)
        if page is None:
            return jsonify({"error": "That FAQ page is not saved for this client."}), 404
        return jsonify({"ok": True, "page": page, "pages": faq.list_pages(client)})

    @app.route("/seo/faq/<slug>/download.<fmt>")
    def seo_faq_download(slug, fmt):
        """html = embeddable accordion, doc = customer review document,
        json = FAQPage schema only."""
        gate = _require_page()
        if gate:
            return gate
        from . import faq
        name = _client_from_slug(slug)
        urls = _urls_param()
        suffix = "page" if len(urls) == 1 else ("selected" if urls else "all")
        if fmt == "html":
            body = faq.accordion_html(name, urls or None,
                                      standalone=request.args.get("standalone") == "1")
            ctype, ext = "text/plain; charset=utf-8", "html"
        elif fmt == "json":
            body = faq.schema_html(name, urls or None)
            ctype, ext = "text/plain; charset=utf-8", "txt"
        else:
            body = faq.review_doc(name, urls or None)
            ctype, ext = "application/msword", "doc"
        resp = make_response(body)
        resp.headers["Content-Type"] = ctype
        resp.headers["Content-Disposition"] = \
            f'attachment; filename="{slug}-faq-{suffix}.{ext}"'
        return resp

    # ------------------ uploaded proposals (Client 360) ------------------
    @app.route("/api/proposal-to-io", methods=["POST"])
    def api_proposal_to_io():
        """Read a proposal we already sent and pull the IO fields out of it.

        The proposal is the agreement; the insertion order is what makes it
        real. Retyping one into the other is where the two drift apart — the
        proposal promises a package at a price and the IO ends up saying
        something slightly different, and nobody notices until billing.

        This extracts what it can and says how confident it is. It does not
        create the IO: the builder still asks for everything, and a field this
        could not find arrives empty rather than guessed, because a wrong
        number that looks filled in is worse than a blank one.
        """
        gate = _require_api()
        if gate:
            return gate

        body = request.get_json(silent=True) or {}
        url = str(body.get("url") or "").strip()
        client = str(body.get("client") or "").strip()
        if not url.startswith("https://"):
            return jsonify({"error": "That proposal has no readable file."}), 400

        # Only our own storage. This fetches a URL, so without the check it is
        # an SSRF hole that reads anything the server can reach.
        from urllib.parse import urlparse
        host = (urlparse(url).hostname or "").lower()
        if not (host.endswith("cloudinary.com") or host.endswith("res.cloudinary.com")
                or host == (urlparse(request.url_root).hostname or "").lower()):
            return jsonify({"error": "That file isn't stored with us, so it "
                                     "can't be read."}), 400

        import requests as _rq
        try:
            r = _rq.get(url, timeout=20)
            r.raise_for_status()
            raw = r.content
        except Exception as exc:                        # noqa: BLE001
            return jsonify({"error": f"Couldn't fetch the proposal "
                                     f"({type(exc).__name__})."}), 502

        text = ""
        if raw[:5] == b"%PDF-":
            try:
                import io as _io
                from pypdf import PdfReader
                pages = PdfReader(_io.BytesIO(raw)).pages
                # First 12 pages: the plan and pricing are near the front, and
                # a 60-page appendix is cost without content.
                text = "\n".join((p.extract_text() or "") for p in pages[:12])
            except Exception as exc:                    # noqa: BLE001
                return jsonify({"error": f"That PDF couldn't be read "
                                         f"({type(exc).__name__})."}), 422
        else:
            try:
                text = raw.decode("utf-8", "ignore")
            except Exception:                           # noqa: BLE001
                text = ""

        text = " ".join(text.split())[:24000]
        if len(text) < 40:
            return jsonify({"error": "There's no readable text in that file — "
                                     "it may be a scan rather than a document.",
                            "fields": {}, "questions": []}), 200

        from . import ai as _ai, rate_card as _rc
        if not _ai.ready():
            return jsonify({"error": "AI isn't configured, so the proposal "
                                     "can't be read automatically.",
                            "fields": {}, "questions": []}), 200

        # The rate card is given to the model so a line it reads as "OTT" is
        # matched against what we actually sell, rather than invented.
        catalog = [p.get("label", "") for p in (_rc.products() or [])][:120]

        schema_hint = {
            "client": "business name the proposal is addressed to",
            "monthly_total": "total monthly media spend as a number, or null",
            "term_months": "campaign length in months, or null",
            "start_date": "YYYY-MM-DD if stated, else null",
            "end_date": "YYYY-MM-DD if stated, else null",
            "products": "array of {product, monthly, notes} — product must be "
                        "one of the catalog labels, or the closest match",
            "geography": "markets or radius named, else null",
            "notes": "anything a trafficker would need that has no field",
        }
        try:
            out = _ai.chat_json(
                [{"role": "system", "content":
                  "You read media proposals and extract the facts needed to "
                  "write an insertion order. Never invent a number: if the "
                  "proposal does not state something, return null for it. "
                  "Match products to the supplied catalog; if nothing is a "
                  "reasonable match, use the proposal's own wording and say so "
                  "in notes. Return JSON only."},
                 {"role": "user", "content":
                  f"Catalog: {catalog}\n\nReturn JSON with these keys: "
                  f"{schema_hint}\n\nProposal text:\n{text}"}],
                module="io_builder", purpose="proposal_to_io")
        except Exception as exc:                        # noqa: BLE001
            return jsonify({"error": f"The proposal couldn't be read "
                                     f"({type(exc).__name__})."}), 502

        fields = out if isinstance(out, dict) else {}
        if client and not fields.get("client"):
            fields["client"] = client

        # What the IO needs and the proposal did not say. These become the
        # questions the builder asks, so the gap is explicit rather than a
        # blank field someone has to notice.
        asks = []
        if not fields.get("start_date"):
            asks.append({"key": "start_date", "q": "What date does the campaign start?"})
        if not fields.get("term_months"):
            asks.append({"key": "term_months", "q": "How many months does it run?"})
        if not fields.get("monthly_total"):
            asks.append({"key": "monthly_total", "q": "What is the monthly media spend?"})
        if not (fields.get("products") or []):
            asks.append({"key": "products", "q": "Which products should the IO carry?"})
        if not fields.get("geography"):
            asks.append({"key": "geography", "q": "Which markets or radius does it cover?"})

        # Run what was extracted past the same guardrails the IO enforces, so a
        # proposal that promises something unwritable is caught here rather
        # than at the end of the builder.
        checks = []
        try:
            items = [{"product": p.get("product", ""),
                      "monthly": p.get("monthly") or 0}
                     for p in (fields.get("products") or [])]
            if items:
                checks = _rc.guardrails(items)
        except Exception:                               # noqa: BLE001
            checks = []

        return jsonify({"fields": fields, "questions": asks,
                        "guardrails": checks,
                        "note": ("Read from the proposal. Anything it didn't "
                                 "state is left blank rather than guessed.")})

    @app.route("/api/client/proposals")
    def api_client_proposals():
        gate = _require_api()
        if gate:
            return gate
        from . import proposals, client_groups
        name = (request.args.get("client") or "").strip()
        if not name:
            return jsonify({"proposals": []})
        items = proposals.list_proposals(name)
        others = [n for n in client_groups.member_names(name)
                  if n and n.strip().lower() != name.lower()]
        for other in others:
            for row in proposals.list_proposals(other):
                row = dict(row)
                row["member"] = other
                items.append(row)
        if others:
            items.sort(key=lambda i: (str(i.get("date_sent") or ""),
                                      str(i.get("uploaded_at") or "")), reverse=True)
        return jsonify({"proposals": items, "group": others,
                        "cloudinary": proposals.cloudinary_ready()})

    @app.route("/api/client/proposals/upload", methods=["POST"])
    def api_client_proposals_upload():
        gate = _require_api()
        if gate:
            return gate
        from . import proposals
        client = (request.form.get("client") or "").strip()
        upload = request.files.get("file")
        if not client:
            return jsonify({"error": "client is required."}), 400
        if upload is None or not upload.filename:
            return jsonify({"error": "Choose a PDF or Word file to upload."}), 400
        try:
            record = proposals.add_proposal(
                client, upload.filename, upload.read(),
                date_sent=request.form.get("date_sent", ""),
                title=request.form.get("title", ""),
                note=request.form.get("note", ""),
                actor=current_user() or "",
                value=request.form.get("value", ""),
                term=request.form.get("term", "monthly"),
                status=request.form.get("status", "sent"))
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": f"Upload failed: {exc}"}), 500
        audit.log("hub", "proposal_uploaded", actor=current_user(), detail=client,
                  name=record["filename"], date_sent=record["date_sent"])

        # A proposal that reaches a client is a deal in progress, so it opens an
        # opportunity in Smart 1 Suite. The contact is looked up first and only
        # created when the uploader supplies details -- an opportunity attached
        # to a contact invented from a business name is one no salesperson can
        # act on, and it duplicates the real contact next time anyone searches.
        #
        # The file is already saved by this point. If Suite says it needs a
        # contact, the response says so and the upload still stands; the
        # uploader answers and posts to /api/client/proposals/opportunity.
        suite = _proposal_opportunity(client, record, {
            "name": request.form.get("contact_name", ""),
            "email": request.form.get("contact_email", ""),
            "phone": request.form.get("contact_phone", ""),
        })
        if suite.get("ok"):
            proposals.update_proposal(client, record["id"], {
                "opportunity_id": suite.get("opportunity_id", "")})
            record = next((i for i in proposals.list_proposals(client)
                           if i.get("id") == record["id"]), record)
        return jsonify({"ok": True, "proposal": record, "suite": suite,
                        "proposals": proposals.list_proposals(client)})

    def _proposal_opportunity(client, record, contact):
        """File one uploaded proposal as a Smart 1 Suite opportunity."""
        try:
            from . import suite_opportunity
        except Exception as exc:  # noqa: BLE001
            return {"ok": False, "reason": f"Suite helper unavailable ({type(exc).__name__})."}
        # The client's own website, from the registry. Looked up by name and
        # left blank when there is no match -- canonical_domain() reads a URL,
        # not a client name, and feeding it one produces a plausible-looking
        # domain that belongs to nobody.
        website = ""
        try:
            from . import clients_registry
            website = (clients_registry.find_client(client) or {}).get("url", "")
        except Exception:  # noqa: BLE001
            pass
        return suite_opportunity.push_proposal(
            client=client,
            title=f"{client} — {record.get('title') or 'Marketing Proposal'}",
            value=float(record.get("value") or 0), contact=contact,
            website=website, pdf_url=record.get("url", ""),
            opportunity_id=str(record.get("opportunity_id") or ""),
            source="Smart 1 Hub — Client 360")

    @app.route("/api/client/proposals/opportunity", methods=["POST"])
    def api_client_proposals_opportunity():
        """Open (or retry) the Suite opportunity for an already-uploaded proposal.

        Separate from the upload so a missing contact costs one extra click
        rather than a re-upload, and so a proposal filed while Suite was
        misconfigured can be pushed later without touching the file.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import proposals
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        pid = (body.get("id") or "").strip()
        if not client or not pid:
            return jsonify({"error": "client and id are required."}), 400
        record = next((i for i in proposals.list_proposals(client)
                       if i.get("id") == pid), None)
        if record is None:
            return jsonify({"error": "Not found"}), 404
        suite = _proposal_opportunity(client, record, body.get("contact") or {})
        if suite.get("ok"):
            proposals.update_proposal(client, pid, {
                "opportunity_id": suite.get("opportunity_id", "")})
            audit.log("hub", "proposal_opportunity", actor=current_user(),
                      detail=client, opportunity=suite.get("opportunity_id", ""))
        return jsonify({"ok": suite.get("ok", False), "suite": suite,
                        "proposals": proposals.list_proposals(client)})

    @app.route("/api/client/proposals/suite-status")
    def api_client_proposals_suite_status():
        """Whether an upload will actually reach Smart 1 Suite."""
        gate = _require_api()
        if gate:
            return gate
        try:
            from . import suite_opportunity
            return jsonify(suite_opportunity.status())
        except Exception as exc:  # noqa: BLE001
            return jsonify({"ok": False,
                            "problems": [f"Suite helper unavailable ({type(exc).__name__})."]})

    @app.route("/api/client/proposals/update", methods=["POST"])
    def api_client_proposals_update():
        gate = _require_api()
        if gate:
            return gate
        from . import proposals
        body = request.get_json(silent=True) or {}
        client = (body.get("client") or "").strip()
        pid = (body.get("id") or "").strip()
        if not client or not pid:
            return jsonify({"error": "client and id are required."}), 400
        if body.get("delete"):
            ok = proposals.delete_proposal(client, pid)
            return jsonify({"ok": ok, "proposals": proposals.list_proposals(client)})
        hit = proposals.update_proposal(client, pid, body)
        if hit is None:
            return jsonify({"error": "Not found"}), 404
        return jsonify({"ok": True, "proposal": hit,
                        "proposals": proposals.list_proposals(client)})

    @app.route("/api/client/proposals/file/<path:name>")
    def api_client_proposals_file(name):
        """Serves proposals kept on disk when Cloudinary isn't configured."""
        gate = _require_page()
        if gate:
            return gate
        from . import proposals
        path = proposals.local_file_path(name)
        if not path:
            return jsonify({"error": "Not found"}), 404
        return send_from_directory(os.path.dirname(path), os.path.basename(path))

    @app.route("/qa")
    def qa_home():
        gate = _require_page()
        if gate:
            return gate
        from . import qa
        seen, groups = [], {}
        for key, meta in qa.REPORTS.items():
            g = meta.get("group", "Reports")
            if g not in groups:
                groups[g] = []
                seen.append(g)
            groups[g].append((key, meta))
        # Two audits live outside REPORTS because they're modules with their
        # own pages, not table-returning functions. They still belong here —
        # somebody looking for "what's wrong" shouldn't have to know which
        # kind of thing each one is.
        # Six of these are whole tools rather than table-returning functions,
        # and every one of them answers "what is wrong / what do we owe" —
        # which is the question this page exists for and is not what the
        # Tools page is for. They were on Tools under "Client Work", a group
        # whose name described where the work came from rather than what the
        # screen is for, and a report nobody thinks to look for is a report
        # nobody works. Each keeps its own URL, so every existing link and
        # every Client 360 crumb still resolves.
        extras = qa.EXTRAS
        for g, key, meta in extras:
            if g not in groups:
                groups[g] = []
                seen.append(g)
            groups[g].append((key, meta))
        return render_template("qa.html", user=current_user(), modules=MODULES,
                               active="qa", groups=[(g, groups[g]) for g in seen])

    @app.route("/qa/<key>")
    def qa_report(key):
        gate = _require_page()
        if gate:
            return gate
        from . import qa
        meta = qa.REPORTS.get(key)
        if not meta:
            return redirect("/qa")
        return render_template("qa_report.html", user=current_user(), modules=MODULES,
                               active="qa", key=key, title=meta["title"])

    @app.route("/api/qa/<key>")
    def api_qa(key):
        """Today's answer for one report — the stored run, or the first one.

        A GET never re-runs a report that has already answered today. That is
        the whole point: these builds walk a year of QuickBooks invoices, the
        GoHighLevel pipeline and the Google index, and they were doing it on
        every open, every Back button and every refresh of the tab. Re-running
        is the POST below, because a GET that rebuilds is one a prefetch or a
        link preview fires without anybody asking.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import qa
        if key not in qa.REPORTS:
            return jsonify({"error": f"Unknown report: {key}"}), 404
        try:
            out = qa.run_cached(key,
                                month=(request.args.get("month") or "").strip())
        except Exception as exc:  # noqa: BLE001 — reports must degrade gracefully
            out = {"columns": [], "rows": [], "error": str(exc)}
        out.setdefault("key", key)
        out.setdefault("title", qa.REPORTS[key]["title"])
        audit.log("hub", "qa_report", actor=current_user(), detail=key,
                  cached=bool((out.get("cache") or {}).get("from_cache")))
        return jsonify(out)

    @app.route("/api/qa/<key>/refresh", methods=["POST"])
    def api_qa_refresh(key):
        """Run this report again now, and keep what it returns for the day."""
        gate = _require_api()
        if gate:
            return gate
        from . import qa
        if key not in qa.REPORTS:
            return jsonify({"error": f"Unknown report: {key}"}), 404
        try:
            out = qa.run_cached(key,
                                month=(request.args.get("month") or "").strip(),
                                force=True)
        except Exception as exc:  # noqa: BLE001 — reports must degrade gracefully
            out = {"columns": [], "rows": [], "error": str(exc)}
        out.setdefault("key", key)
        out.setdefault("title", qa.REPORTS[key]["title"])
        audit.log("hub", "qa_report_refreshed", actor=current_user(), detail=key)
        return jsonify(out)

    @app.route("/api/qa/accounting/status", methods=["POST"])
    def api_qa_accounting_status():
        gate = _require_api()
        if gate:
            return gate
        from . import qa
        body = request.get_json(silent=True) or {}
        opp_id = (body.get("id") or "").strip()
        stage_id = (body.get("stage_id") or "").strip()
        status = (body.get("status") or "").strip().lower()
        if not opp_id or (not stage_id and not status):
            return jsonify({"error": "id and stage_id or status are required."}), 400
        if status and status not in qa.GHL_STATUSES:
            return jsonify({"error": f"status must be one of {', '.join(qa.GHL_STATUSES)}."}), 400
        try:
            qa.set_accounting_stage(opp_id, stage_id, status)
        except Exception as exc:  # noqa: BLE001
            errors.log_exception("qa-accounting", exc, path=request.path,
                                 actor=current_user() or "")
            return jsonify({"error": str(exc)})
        audit.log("hub", "accounting_stage_changed", actor=current_user(),
                  detail=f"{opp_id} -> {stage_id}")
        return jsonify({"ok": True})

    @app.route("/api/qa/invoice-off/assign", methods=["POST"])
    def api_qa_invoice_assign():
        gate = _require_api()
        if gate:
            return gate
        from . import qa
        body = request.get_json(silent=True) or {}
        customer = (body.get("customer") or "").strip()
        partner = (body.get("partner") or "").strip()
        if not customer or not partner:
            return jsonify({"error": "customer and partner are required."}), 400
        qa.assign_invoice_partner(customer, partner)
        audit.log("hub", "qa_invoice_assigned", actor=current_user(),
                  detail=f"{customer} -> {partner}")
        return jsonify({"ok": True})

    @app.route("/activity")
    def activity():
        gate = _require_page()
        if gate:
            return gate
        return render_template("activity.html", user=current_user(), modules=MODULES, active="activity")

    @app.route("/status")
    def status():
        gate = _require_page()
        if gate:
            return gate
        return render_template("status.html", user=current_user(), modules=MODULES, active="status")

    # ---------------- Clients app (prebuilt Knack lookup) ----------------
    # The React bundle was built with absolute /static/... and /data/... URLs,
    # so the hub serves those two prefixes straight from clients_app/.
    @app.route("/clients")
    @app.route("/clients/")
    def clients_index():
        gate = _require_page()
        if gate:
            return gate
        from .sidebar import render_sidebar
        _admin = viewer_is_admin()
        with open(os.path.join(CLIENTS_APP, "index.html"), "rb") as fh:
            body = fh.read()
        # Two stylesheets, both after the bundle's own <link> so equal rules
        # here win. theme.css is the shared typography-and-colour layer every
        # module gets; clients-theme.css is this page's alone, because the
        # bundle carries the old near-black-and-lime identity in class names
        # (.kpi, .badge, .tabs) too ordinary to restyle globally. It is scoped
        # to the body class added below for the same reason.
        snippet = (b'<link rel="stylesheet" href="/assets/theme.css">'
                   b'<link rel="stylesheet" href="/assets/clients-theme.css">')
        if b"</head>" in body:
            body = body.replace(b"</head>", snippet + b"</head>", 1)
        # Not data-module: that is what hub-demo.js floats "Walk me through
        # this" onto, and this page has no walkthrough written for it.
        if b"<body>" in body:
            body = body.replace(b"<body>", b'<body class="s1-clients">', 1)
        bar = render_sidebar("clients", is_admin=_admin)
        # Deep links from Client 360: /clients?q=<client> auto-fills and runs
        # the React app's search (native value setter so React sees the input).
        autosearch = b"""<script>
(function(){
  var q=new URLSearchParams(location.search).get('q'); if(!q) return;
  var tries=0;
  var t=setInterval(function(){
    tries++;
    var input=document.querySelector('input[placeholder^="Client, IO"]')||
              document.querySelector('input[type="search"]');
    if(input){
      clearInterval(t);
      var setter=Object.getOwnPropertyDescriptor(window.HTMLInputElement.prototype,'value').set;
      setter.call(input,q);
      input.dispatchEvent(new Event('input',{bubbles:true}));
      input.focus();
    } else if(tries>60){clearInterval(t);}
  },250);
})();
</script>"""
        addition = autosearch + bar
        # The last </body>, not the first — see the note in wsgi.py's HubBar.
        # A page that builds a printable document as a JavaScript string
        # carries its own </body> inside a template literal, and injecting
        # there breaks the page's script instead of ending the document.
        _cut = body.rfind(b"</body>")
        body = (body[:_cut] + addition + body[_cut:]) if _cut >= 0 else body + addition
        return app.response_class(body, mimetype="text/html")

    @app.route("/static/<path:filename>")
    def clients_static(filename):
        return send_from_directory(os.path.join(CLIENTS_APP, "static"), filename)

    @app.route("/data/<path:filename>")
    def clients_data(filename):
        if not current_user():
            return jsonify({"error": "Not authenticated."}), 401
        return send_from_directory(os.path.join(CLIENTS_APP, "data"), filename)

    # ---------------- QuickBooks connect / lookup ----------------
    @app.route("/qb/connect")
    def qb_connect():
        gate = _require_page()
        if gate:
            return gate
        from . import quickbooks as qb
        if not qb.configured():
            return ("QuickBooks is not configured — set QB_CLIENT_ID and QB_CLIENT_SECRET "
                    "(from developer.intuit.com) and redeploy. <a href='/status'>Status</a>", 400)
        return redirect(qb.authorize_url(request))

    @app.route("/qb/callback")
    def qb_callback():
        gate = _require_page()
        if gate:
            return gate
        from . import quickbooks as qb
        ok, msg = qb.handle_callback(request)
        audit.log("hub", "quickbooks_connected" if ok else "quickbooks_connect_failed",
                  actor=current_user(), detail=msg)
        return redirect("/status?qb=" + ("connected" if ok else "error"))

    @app.route("/qb/disconnect", methods=["POST", "GET"])
    def qb_disconnect():
        gate = _require_page()
        if gate:
            return gate
        from . import quickbooks as qb
        qb.disconnect()
        audit.log("hub", "quickbooks_disconnected", actor=current_user())
        return redirect("/status")

    @app.route("/api/qb/invoices")
    def api_qb_invoices():
        gate = _require_api()
        if gate:
            return gate
        from . import quickbooks as qb
        q = (request.args.get("q") or "").strip()
        cid = (request.args.get("customer_id") or "").strip()
        client = (request.args.get("client") or "").strip()
        if client:
            # Grouped clients bill as one company, so the card reads every
            # member: each one's attached customers where it has them, a name
            # search where it does not, deduplicated by customer id.
            from . import seo, client_groups
            entries = []
            for member in client_groups.member_names(client) or [client]:
                att = seo.get_links(member).get("qb") or []
                if not isinstance(att, list):
                    att = [att]
                entries.append({"client": member,
                                "ids": [str(a.get("id")) for a in att
                                        if isinstance(a, dict) and a.get("id")]})
            try:
                return jsonify(qb.lookup_for_clients(entries))
            except Exception as exc:  # noqa: BLE001 — Client 360 must degrade
                return jsonify({"configured": qb.configured(),
                                "connected": qb.connected(),
                                "customers": [], "error": str(exc)})
        if not q and not cid:
            return jsonify({"configured": qb.configured(), "connected": qb.connected(), "customers": []})
        try:
            return jsonify(qb.lookup(q, customer_id=cid or None))
        except Exception as exc:  # noqa: BLE001 — Client 360 must degrade gracefully
            return jsonify({"configured": qb.configured(), "connected": qb.connected(),
                            "customers": [], "error": str(exc)})

    @app.route("/api/qb/customers")
    def api_qb_customers():
        """Customer search for the C360 'attach QuickBooks customer' flow."""
        gate = _require_api()
        if gate:
            return gate
        from . import quickbooks as qb
        q = (request.args.get("q") or "").strip()
        if not q:
            return jsonify({"customers": []})
        try:
            return jsonify({"customers": qb.find_customers(q, limit=6)})
        except Exception as exc:  # noqa: BLE001
            return jsonify({"customers": [], "error": str(exc)})

    @app.route("/favicon.ico")
    def favicon():
        return ("", 204)

    @app.route("/health")
    @app.route("/healthz")
    def health():
        return jsonify({"status": "ok"})

    @app.route("/login/health")
    def login_health():
        """Why sign-in isn't working — readable WITHOUT signing in.

        That is the whole point: every other diagnostic in the Hub sits behind
        the login, which is useless when the login is the thing that's broken.
        Boot failures were being stored in app.config and never surfaced, so a
        users table that failed to create looked identical to a wrong password.

        Reports booleans and error *types* only. No secrets, no password, no
        token, no email addresses.
        """
        import traceback
        out = {"version": None, "panel_password_set": bool(auth.panel_password()),
               "google_button_enabled": (os.environ.get("HUB_GOOGLE_LOGIN", "").lower()
                                         in {"1", "true", "yes", "on"})}
        try:
            from . import version as _v
            out["version"] = _v.label()
        except Exception:  # noqa: BLE001
            pass

        out["db_boot_error"] = app.config.get("HUB_DB_BOOT_ERROR") or None
        out["users_registered"] = app.config.get("HUB_USERS_REGISTERED", None)
        if out["users_registered"] is False:
            out["signup_available"] = False
        out["users_boot_error"] = app.config.get("HUB_USERS_BOOT_ERROR") or None

        # Can we actually reach the users table? This is the failure that makes
        # /signup return a 500 with nothing to go on.
        try:
            from .users import User
            out["users_table"] = "ok"
            out["user_count"] = User.query.count()
            out["super_admins_seeded"] = User.query.filter_by(
                role="super_admin").count()
            out["super_admins_with_password"] = User.query.filter(
                User.role == "super_admin", User.password_hash != "").count()
        except Exception as exc:  # noqa: BLE001
            out["users_table"] = f"{type(exc).__name__}"
            out["users_table_detail"] = str(exc)[:200]
            out["user_count"] = None

        try:
            from .extensions import database_url
            url = database_url()
            out["database"] = ("postgres" if url.startswith("postgres")
                               else "sqlite" if url.startswith("sqlite") else "other")
            if url.startswith("sqlite"):
                path = url.replace("sqlite:///", "")
                out["sqlite_path"] = path
                out["sqlite_dir_writable"] = os.access(os.path.dirname(path) or ".", os.W_OK)
        except Exception as exc:  # noqa: BLE001
            out["database"] = f"error: {type(exc).__name__}"

        # Plain-English verdict, so nobody has to interpret the booleans.
        problems = []
        if app.config.get("HUB_USERS_REGISTERED") is False:
            problems.append(
                "The user-accounts blueprint failed to register, so /signup "
                "and /diagnostics/users return 404. Reason: "
                + str(app.config.get("HUB_USERS_BOOT_ERROR", "unknown"))
                + " — if it names flask_sqlalchemy, Flask-SQLAlchemy is "
                  "missing from requirements.txt.")
        try:
            from .config import settings as _cfg
            for w in _cfg.placeholder_warnings():
                problems.append(w["detail"])
        except Exception:  # noqa: BLE001
            pass
        if out.get("users_table") != "ok":
            problems.append("The user accounts table isn't reachable, so /signup "
                            "will fail. Usually DATABASE_URL is unset and the "
                            "disk fallback isn't writable.")
        if not out["panel_password_set"]:
            problems.append("PANEL_PASSWORD is not set, so the shared-password "
                            "fallback can't work either.")
        if out.get("db_boot_error"):
            problems.append("The database failed at boot: " + str(out["db_boot_error"])[:160])
        if not problems and out.get("super_admins_with_password", 0) == 0:
            problems.append("No super admin has set a password yet — go to "
                            "/signup and register todd@smart1marketing.com.")
        out["problems"] = problems
        out["ok"] = not problems
        return jsonify(out)

    # ---------------- JSON APIs (hub-side) ----------------
    def _require_api():
        if not current_user():
            return jsonify({"error": "Not authenticated."}), 401
        return None

    @app.route("/api/summary")
    def api_summary():
        gate = _require_api()
        if gate:
            return gate
        try:
            data = knack_data.summary()
        except Exception as exc:  # noqa: BLE001 — dashboard must never 500
            data = {"error": str(exc)}
        try:
            from . import seo
            seo_rows = seo.seo_clients()
            data["seo_clients"] = len(seo_rows)
            data["seo_billing_monthly"] = round(sum(c["billing"] for c in seo_rows))
        except Exception:  # noqa: BLE001 — SEO totals are additive, never break the dashboard
            data.setdefault("seo_clients", None)
            data.setdefault("seo_billing_monthly", None)
        return jsonify(data)

    @app.route("/api/presence")
    def api_presence():
        """How many people are around, and who.

        A route of its own rather than a key on `/api/status`, and that is not
        tidiness: `/api/status` is in `access.UTILITY_PREFIXES`, so for the
        eleven General accounts it answers 403 — the headcount would have been
        admin-only while sitting on everybody's dashboard, reading as zero.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import presence
        data = presence.active()
        data["line"] = presence.summary_line(data)
        return jsonify(data)

    @app.route("/api/housekeeping")
    def api_housekeeping():
        """What needs filling in across the Hub, and the page it shows on.

        A Utilities path (`hub/access.py`), which is the whole point: these
        are to-dos for whoever can open the Users panel, and until now the
        only place any of them appeared was underneath a dashboard card
        everybody reads and three people can act on.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import housekeeping
        return jsonify(housekeeping.findings())

    @app.route("/api/celebrations")
    def api_celebrations():
        """Whose birthday and whose work anniversary, this month and today.

        One request, two answers: the dashboard block reads the month and the
        popup reads `today`. Splitting them into two routes would mean two
        reads of the profile table on every page load for the same rows.

        `me` is what the popup greets by name. It is resolved from the signed-
        in *account* where there is one, and only falls back to the cookie's
        display name — two people on this roster are called Todd, and the
        shared PANEL_PASSWORD session carries a name with no account behind
        it at all.
        """
        gate = _require_api()
        if gate:
            return gate
        from . import celebrations
        try:
            data = celebrations.this_month()
            data["today_list"] = celebrations.today()
        except Exception as exc:  # noqa: BLE001 — never break the dashboard
            return jsonify({"error": str(exc), "birthdays": [],
                            "anniversaries": []})
        email, name = "", current_user() or ""
        account, account_read = None, False
        try:
            from .users_routes import current_account
            account, account_read = current_account(), True
            if account is not None:
                email, name = (account.email or ""), (account.name or name)
        except Exception:  # noqa: BLE001 — a session with no account row
            pass
        data["me"] = celebrations.mine(data["today_list"], email=email,
                                       name=name)
        data["me_name"] = name
        # Who is missing a date is a job for whoever can open the Users panel,
        # and eleven of the fourteen accounts are answered 403 by it. A
        # General account is still told the list is not the whole roster —
        # that was the point of the sentence — and is told none of the rest:
        # the counts, the names and the link are on /diagnostics, where the
        # person who can act on them is looking. hub/housekeeping.py decides,
        # so the template cannot describe it a second, different way.
        # A row we could not read is not an admin: the gaps are withheld
        # rather than shown, because the failure mode of guessing the other
        # way is a to-do published to somebody who cannot act on it, which is
        # the thing being fixed. A shared-password session reads as an admin
        # here exactly as it does everywhere else — `account_read` is true and
        # the row is legitimately None.
        if not (account_read and viewer_is_admin(account)):
            from . import housekeeping
            data["not_recorded"] = housekeeping.withheld(
                data.get("not_recorded") or {})
        return jsonify(data)

    @app.route("/api/c360")
    def api_c360():
        gate = _require_api()
        if gate:
            return gate
        q = request.args.get("q", "")
        try:
            groups = knack_data.search_client(q)
        except Exception as exc:  # noqa: BLE001
            return jsonify({"groups": [], "error": str(exc)})
        return jsonify({"groups": groups})

    @app.route("/api/c360/sites")
    def api_c360_sites():
        """Best-effort search of the Simvoly admin's Postgres inventory."""
        gate = _require_api()
        if gate:
            return gate
        q = (request.args.get("q") or "").strip()
        dsn = os.environ.get("DATABASE_URL", "")
        if not q or not dsn:
            return jsonify({"results": [], "configured": bool(dsn)})
        try:
            import psycopg2
            import psycopg2.extras
            conn = psycopg2.connect(dsn)
            try:
                with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                    cur.execute(
                        """
                        SELECT p.project_id, p.name, p.status,
                               (SELECT w.domain FROM websites w
                                 WHERE w.project_id = p.project_id AND w.domain IS NOT NULL
                                 LIMIT 1) AS domain
                          FROM projects p
                         WHERE p.name ILIKE %s
                            OR p.project_id IN (
                                 SELECT w2.project_id FROM websites w2
                                  WHERE w2.name ILIKE %s OR w2.domain ILIKE %s)
                         ORDER BY p.name LIMIT 6
                        """,
                        (f"%{q}%", f"%{q}%", f"%{q}%"),
                    )
                    rows = [dict(r) for r in cur.fetchall()]
            finally:
                conn.close()
            return jsonify({"results": rows, "configured": True})
        except Exception as exc:  # noqa: BLE001
            return jsonify({"results": [], "configured": True, "error": str(exc)})

    @app.route("/api/errors")
    def api_errors():
        gate = _require_api()
        if gate:
            return gate
        limit = clamp_int(request.args.get("limit"), 50, 1, 300)
        return jsonify({"errors": errors.read(limit=limit)})

    @app.route("/api/errors/clear", methods=["POST"])
    def api_errors_clear():
        gate = _require_api()
        if gate:
            return gate
        errors.clear()
        audit.log("hub", "error_log_cleared", actor=current_user())
        return jsonify({"ok": True})

    @app.route("/api/activity")
    def api_activity():
        gate = _require_api()
        if gate:
            return gate
        limit = clamp_int(request.args.get("limit"), 300, 1, 1000)
        module = request.args.get("module") or None
        return jsonify({"entries": audit.read(limit=limit, module=module)})

    @app.route("/api/status")
    def api_status():
        gate = _require_api()
        if gate:
            return gate
        # Every credential below is read through hub.config rather than
        # os.environ. A self-test is the one page where reading one spelling of
        # a setting is worst: it reports a key as missing that the tool beside
        # it is happily using, and somebody goes and sets a second copy of a
        # variable that was never the problem.
        from .config import settings as _cfg
        checks = []

        def add(name, status_, message):
            checks.append({"name": name, "status": status_, "message": message})

        # --- core config ---
        pw = auth.panel_password()
        if not pw:
            add("Panel password", "error", "PANEL_PASSWORD is not set — nobody can log in.")
        elif pw in ("change-me", "change-me-to-something-strong"):
            add("Panel password", "warn", "Still set to a placeholder — change it.")
        else:
            add("Panel password", "ok", "Configured.")
        # Through hub.config, which knows all three spellings. Read directly
        # this row said "not set" on a deployment carrying FLASK_SECRET_KEY and
        # nothing else — a self-test reporting a fault that is not there is as
        # expensive as one missing a fault that is.
        _secret_set = bool(_cfg.secret_key)
        add("Session secret", "ok" if _secret_set else "warn",
            "Configured — logins survive restarts." if _secret_set
            else "Not set — everyone is logged out on every restart/redeploy.")

        # --- Knack data ---
        #
        # This row used to read the *file's* mtime and print it as "Refreshed
        # Xh ago", warning past 48 hours. `data_age_hours()`'s own docstring
        # already said why that is wrong: in a Docker deploy every file is
        # written at image build time, so it measures **time since the last
        # deploy** and not since the last data refresh.
        #
        # That is wrong in both directions, which is what makes it worse than
        # nothing. A months-old export reads as "refreshed 2h ago" for the
        # first two days after any deploy; and a container simply left up for
        # a week warns that the data needs refreshing when nothing about the
        # data has changed. Either way the row is not about the data, and it
        # is read as though it is.
        #
        # `export_state()` is the honest signal and is already shared with the
        # dashboard and hub/housekeeping.py: the month the export was
        # generated *for*, against the calendar.
        state = knack_data.export_state()
        age = knack_data.data_age_hours()
        if age is None:
            add("Smart 1 Team data", "error", "clients_app/data/products.json not found.")
        elif state["stale"]:
            add("Smart 1 Team data", "warn",
                f"The committed products export is for {state['label']}, and "
                f"it is now {state['current_label']}. It is only read when "
                "Knack cannot be reached, but that is when it matters.")
        elif not state["period"]:
            # Neither stale nor current: it carries no month at all, so
            # nothing can say how old it is. Named rather than passed off as
            # fresh — the whole failure this row had.
            add("Smart 1 Team data", "warn",
                "The committed products export carries no month, so how old "
                "it is cannot be measured. It is the fallback for when Knack "
                "cannot be reached.")
        else:
            # The site count says which source answered. Products and websites
            # each prefer the live Knack object and fall back to the committed
            # export, and a count with no source on it reads as live whichever
            # it was — which is the whole reason the export went unnoticed.
            _wsrc = knack_data.websites_source()
            add("Smart 1 Team data", "ok",
                f"Export is current ({state['label']}) · "
                f"{len(knack_data.products())} product rows · "
                f"{len(knack_data.websites())} sites "
                f"({'live from Knack' if _wsrc == 'knack' else 'committed export'}).")

        # --- GHL ---
        token, company = _cfg.ghl_token, _cfg.ghl_company_id
        if not token or not company:
            add("GoHighLevel API", "error",
                f"{_cfg.spellings('ghl_token')} and/or "
                f"{_cfg.spellings('ghl_company_id')} is not set.")
        else:
            try:
                r = _rq.get(
                    "https://services.leadconnectorhq.com/locations/search",
                    params={"companyId": company, "limit": "1"},
                    headers={"Authorization": f"Bearer {token}",
                             "Version": os.environ.get("GHL_API_VERSION", "2021-07-28"),
                             "Accept": "application/json"},
                    timeout=12,
                )
                add("GoHighLevel API", "ok" if r.ok else "error",
                    "Token is valid and can read sub-accounts." if r.ok else f"Token check failed (HTTP {r.status_code}).")
            except Exception as exc:  # noqa: BLE001
                add("GoHighLevel API", "error", f"Could not reach GHL: {exc}")

        # --- Simvoly ---
        skey = _cfg.simvoly_key
        if not skey:
            add("Smart 1 Sites Platform API", "warn",
                f"{_cfg.spellings('simvoly_key')} is not set — Smart 1 Sites module runs limited/mock.")
        else:
            base = os.environ.get("SIMVOLY_API_BASE_URL", "https://api.smart1sites.com").rstrip("/")
            try:
                r = _rq.get(f"{base}/api/v1/plans", headers={"X-CLIENT-KEY": skey, "Accept": "application/json"}, timeout=12)
                add("Smart 1 Sites Platform API", "ok" if r.ok else "error",
                    "Key is valid — plan catalog reachable." if r.ok else f"Key check failed (HTTP {r.status_code}).")
            except Exception as exc:  # noqa: BLE001
                add("Smart 1 Sites Platform API", "error", f"Could not reach Smart 1 Sites API: {exc}")
        add("Sites database", "ok" if os.environ.get("DATABASE_URL") else "warn",
            "DATABASE_URL configured." if os.environ.get("DATABASE_URL")
            else "DATABASE_URL not set — Sites inventory won't persist.")

        # --- Brandfetch ---
        bkey = _cfg.brandfetch_key
        if not bkey:
            add("Brandfetch API", "skipped", "Not configured — auto-fill from website is disabled (optional).")
        else:
            try:
                r = _rq.get("https://api.brandfetch.io/v2/brands/brandfetch.com",
                            headers={"Authorization": f"Bearer {bkey}"}, timeout=12)
                if r.ok:
                    add("Brandfetch API", "ok", "Key is valid.")
                elif r.status_code == 429:
                    add("Brandfetch API", "warn", "Key valid, but quota exhausted right now.")
                else:
                    add("Brandfetch API", "error", f"Key check failed (HTTP {r.status_code}).")
            except Exception as exc:  # noqa: BLE001
                add("Brandfetch API", "error", f"Could not reach Brandfetch: {exc}")

        # --- Google OAuth ---
        gid, gsec = os.environ.get("GOOGLE_CLIENT_ID"), os.environ.get("GOOGLE_CLIENT_SECRET")
        add("Google OAuth app", "ok" if gid and gsec else "warn",
            "Client ID + secret configured. Manage connected accounts in the Google module."
            if gid and gsec else "GOOGLE_CLIENT_ID / GOOGLE_CLIENT_SECRET not set — Google module disabled.")

        # --- Smart 1 Ads ---
        # Asked of the module rather than read off the environment, for the
        # reason the video library's check gives: "no credentials", "credentials
        # but nobody has authorised" and "connected" are three different jobs,
        # and only the module can tell them apart. The developer token is the
        # one with lead time on it -- Google approves it -- so it is named.
        try:
            from modules.ads_builder import google_ads as _ads_ga, store as _ads_store
            ads_st = _ads_ga.connection_status(_ads_store)
            if ads_st["missing"]:
                # "Not connected" was the whole message and it read as "the
                # tool is down". Only the API half is: generating a campaign,
                # approving it and handing it over as a Google Ads Editor
                # import all work with none of these set.
                add("Smart 1 Ads", "skipped",
                    "Google Ads API not configured — " + ", ".join(ads_st["missing"])
                    + " not set. The generator, the approval hub and the Ads Editor "
                    "export at /tools/ads work without them; what is unavailable is "
                    "reading live campaigns and deploying through the API. "
                    "GOOGLE_ADS_DEVELOPER_TOKEN is applied for in the Google Ads "
                    "manager account under Tools → API Center.")
            elif not ads_st["connected"]:
                add("Smart 1 Ads", "warn",
                    "Credentials set but no account authorized yet — open "
                    "/tools/ads/settings and click Connect Google Ads.")
            else:
                add("Smart 1 Ads", "ok",
                    f"Connected via {ads_st['refresh_token_source']} · API "
                    f"{ads_st['api_version']}"
                    + (f" · MCC {ads_st['login_customer_id']}"
                       if ads_st["login_customer_id"] else "") + ".")
        except Exception as _ads_exc:  # noqa: BLE001
            add("Smart 1 Ads", "warn", f"Could not be checked: {_ads_exc}")

        # --- QuickBooks ---
        from . import quickbooks as qb
        if not qb.configured():
            add("QuickBooks", "skipped",
                "QB_CLIENT_ID / QB_CLIENT_SECRET not set — invoice lookup disabled (optional).")
        elif not qb.connected():
            add("QuickBooks", "warn",
                "App configured but no company connected yet — use the Connect QuickBooks button below.")
        else:
            add("QuickBooks", "ok", "Connected — client invoice lookup active on Client 360.")

        # --- Sales section (Proposal + Sales Builder) ---
        add("OpenAI API", "ok" if _cfg.openai_key else "skipped",
            "Configured — AI proposal generation enabled." if _cfg.openai_key
            else "OPENAI_API_KEY not set — proposal generation falls back to templates (optional).")
        add("Cloudinary", "ok" if (os.environ.get("CLOUDINARY_URL") or "").startswith("cloudinary://") else "warn",
            "Configured — proposal PDFs and logs persist to Cloudinary."
            if (os.environ.get("CLOUDINARY_URL") or "").startswith("cloudinary://")
            else "CLOUDINARY_URL not set — proposals persist to the local disk only.")

        # --- Display Ad Builder (second process in this container) ---
        # Asked here rather than inferred from config, because "the token is
        # set" and "the renderer is answering" are different failures with
        # different fixes, and only one of them is visible from the outside.
        try:
            from hub import ad_builder_proxy
            ab = ad_builder_proxy.status()
            add("Display Ad Builder", "ok" if ab.get("ok") else "warn", ab.get("detail", ""))
        except Exception as _ab_exc:  # noqa: BLE001
            add("Display Ad Builder", "warn",
                f"Could not be checked: {_ab_exc}")

        # --- Video background library ---
        # Asked of the library itself rather than inferred from the two keys it
        # needs, for the reason the tool's own status card exists: an empty
        # search has three different causes -- Cloudinary unset, indexing never
        # started, or genuinely no match -- and a page that shows them alike
        # sends someone looking for a bug in the search.
        try:
            from hub import video_library as _vl
            vs = _vl.status()
            if not vs.get("cloudinary"):
                add("Video background library", "error",
                    "CLOUDINARY_URL is not set — the footage library cannot be "
                    "read at all.")
            elif vs.get("missing_folders"):
                # Checked before the indexing question, because it outranks it:
                # indexing a folder Cloudinary does not have will report
                # "nothing waiting" for ever and look like a quiet, healthy
                # tool. Named as its own state rather than folded into the
                # counts, which would be truthfully zero and useless.
                add("Video background library", "error",
                    "The library is scoped to folders this Cloudinary account "
                    "does not have: "
                    + ", ".join(vs["missing_folders"])
                    + ". Nothing can be found until either the names or "
                      "CLOUDINARY_URL is corrected — this is not an empty "
                      "library.")
            elif not vs.get("indexing_started"):
                add("Video background library", "warn",
                    "Indexing has never run, so nothing is searchable yet. "
                    "Indexing is forward-only: start it at "
                    "/tools/video-backgrounds/.")
            else:
                # None means "could not be counted", which is not zero and must
                # not print as one.
                indexed = vs.get("indexed_count")
                total = vs.get("library_count")
                counts = (f"{indexed} of {total} clips indexed"
                          if indexed is not None and total is not None
                          else "counts unavailable from Cloudinary right now")
                add("Video background library",
                    "warn" if not vs.get("openai") else "ok",
                    f"Indexing started {vs['cutoff']} · {counts}."
                    + (" OPENAI_API_KEY is not set, so new clips cannot be "
                       "described — what is already indexed stays searchable."
                       if not vs.get("openai") else
                       " Only our own Cloudinary footage is indexed; free "
                       "stock is searched live in Commercial Builder.")
                    + " Scope: " + ", ".join(vs.get("folders") or []) + ".")
        except Exception as _vl_exc:  # noqa: BLE001
            add("Video background library", "warn",
                f"Could not be checked: {_vl_exc}")

        # --- binaries for the PDF optimizer ---
        gs, qpdf = shutil.which("gs"), shutil.which("qpdf")
        if gs and qpdf:
            try:
                v = subprocess.run([gs, "--version"], capture_output=True, text=True, timeout=10).stdout.strip()
            except Exception:  # noqa: BLE001
                v = "?"
            add("Ghostscript / qPDF", "ok", f"gs {v} · qpdf present — PDF optimizer ready.")
        else:
            add("Ghostscript / qPDF", "error", "Missing gs/qpdf — PDF optimizer will fail (Docker image installs these).")

        # --- persistent disk ---
        add("Persistent disk", "ok" if os.path.isdir("/var/data") else "warn",
            "/var/data mounted — audit log & tokens survive deploys." if os.path.isdir("/var/data")
            else "/var/data not mounted — audit log and Google tokens are ephemeral.")

        return jsonify({"checks": checks})

    # Shared SQLAlchemy instance. Must run BEFORE any module blueprint is
    # registered, because their models bind to it at import time and their
    # create_all() runs during registration.
    try:
        from .extensions import init_db
        init_db(app)
    except Exception as _db_exc:  # noqa: BLE001
        try:
            errors.log_exception("hub", _db_exc)
        except Exception:  # noqa: BLE001
            pass

    # ---- The hub app's own prospect-facing pages -------------------------
    #
    # A blueprint-registered module is not a dispatcher-mounted one, and the
    # two are protected by different code. `bare_prefixes` in wsgi.py keeps
    # the chrome off a mounted module's public routes and `hub/embed.py`
    # gives it the marketing-site frame-ancestors; a blueprint on the hub app
    # passes through neither. The Media Calculators are exactly that: the
    # pages behind smart1marketing.com/ims and the four calculator pages are
    # blueprint routes, so every rule written for /land/<tool>/embed missed
    # them and BOTH halves failed at once --- the staff sidebar was injected
    # into a page a prospect reads, and `_embed_policy` below answered the
    # marketing site's iframe with the Smart 1 Suite refusal, in plain text,
    # 403. A prospect saw "This Hub page is not available inside Smart 1
    # Suite." where the calculator should be, and nothing errored at either
    # end.
    #
    # Read from the module's own declaration rather than restated here, for
    # the reason modules/ads_builder gives wsgi.py: the mount and the module
    # must not be able to disagree about what is public.
    try:
        from modules.calculators import public_paths as _calc_public_paths
        PUBLIC_EMBED_PREFIXES = tuple(_calc_public_paths())
    except Exception:  # noqa: BLE001 -- a module that will not import must
        PUBLIC_EMBED_PREFIXES = ()      # not take the hub's chrome with it

    # ---------------- sidebar for blueprint-registered pages ----------------
    # Modules mounted through DispatcherMiddleware get the sidebar injected by
    # HubBar in wsgi.py. Modules registered as blueprints on this app do not,
    # and they don't extend base.html either — so Tickets, Calculators, Image
    # Picker, Page Images and Google Access rendered with no navigation at all.
    #
    # Editing five modules' templates would fix today and break again the next
    # time one is added. Injecting on the way out covers every one of them, and
    # anything registered later, automatically.
    #
    # "/sales/landing/p/" is a built landing page, served to a prospect on a
    # client's campaign and often pasted onto the client's own domain. The
    # staff sidebar, the help layer and the feedback tab must never appear on
    # one: it leaks internal navigation to a third party and makes the page
    # look broken. The maker itself, at /sales/landing, is a staff page and
    # keeps its chrome -- which is why this is the longer prefix and not
    # "/sales/landing".
    CHROMELESS = ("/login", "/signup", "/reset", "/signin", "/account",
                  # The forgotten-password page and the admin-only refusal both
                  # render on _users_base.html, which is a bare card with no
                  # <body> the injector would recognise -- and injecting the
                  # staff nav into a page somebody reached because they cannot
                  # sign in is the wrong thing to show them anyway.
                  "/forgot",
                  # Plain text, and not HTML, so the injector would skip them
                  # regardless. Named so it is a decision rather than a
                  # coincidence of the mimetype check.
                  "/robots.txt", "/llms.txt",
                  "/connect", "/api/", "/assets/", "/hub-", "/static/",
                  "/sales/landing/p/",
                  # The Smart 1 Suite app frame. A *client* opens this inside
                  # their own sub-account and has no Hub account at all, so
                  # the staff sidebar, help layer and feedback tab must not be
                  # injected into it — the same reason the landing pages above
                  # are here, one audience further out.
                  # NOT "/suite/…": that prefix is a mounted module, and a hub
                  # route under a mount never receives the request. /api/integrity
                  # has a high-severity check for exactly that, and it caught
                  # this one before it shipped.
                  "/suite-app",
                  # The display-ad proof. A client opens this to approve or
                  # send back a set of banners, so it must not arrive wearing
                  # the staff sidebar, the help layer and a feedback tab --
                  # the same reason the built landing pages above are here.
                  # The longer prefix, so the builder itself at
                  # /tools/display-ads keeps its chrome.
                  "/tools/display-ads/proof/",
                  # The commercial review link. A client opens this to approve
                  # a finished cut or send it back, so it must not arrive
                  # wearing the staff sidebar and a feedback tab — same reason
                  # as the display-ad proof above, and the longer prefix again
                  # so the builder at /tools/commercial-builder keeps its own.
                  #
                  # This is the second half of a pair: the first is the login
                  # exemption in modules/commercial_builder/__init__.py. That
                  # module is a BLUEPRINT on this app rather than a mounted
                  # one, so wsgi.py's PUBLIC_PREFIXES — which does both halves
                  # for ads_builder and scans — never sees it, and each half
                  # has to be written out separately. A page exempted from the
                  # login and not from the chrome is a client looking at our
                  # nav; the other way round is a login form in front of
                  # somebody with no account.
                  "/tools/commercial-builder/review/") + PUBLIC_EMBED_PREFIXES

    @app.after_request
    def _embed_policy(resp):
        """Who may frame a hub page, and what they get when they do.

        Runs on every hub response, not only HTML: a stylesheet or a JSON fetch
        made by the framed page has to carry the same allowlist, or the page
        renders and its data does not.

        Two things happen here and nowhere else, because `bare_prefixes` in
        wsgi.py only covers dispatcher-mounted modules and Client 360 is a hub
        route -- the exact gap that makes this the page most worth embedding
        and the one nothing was protecting.
        """
        try:
            from . import suite_embed as embed
            if not embed.is_embedded(request.environ):
                return resp
            path = request.path or "/"
            # A public page of ours, framed on the marketing site. This is not
            # the Suite question at all: `is_embedded()` is true for ANY
            # framer, so without this the marketing site's iframe is answered
            # with the Suite refusal -- and the answer to "who may frame the
            # calculators" is hub/embed.py's allowlist, the same one every
            # /land/<tool>/embed already carries. Checked before `embeddable`
            # so the refusal below can never reach a prospect.
            if path.startswith(PUBLIC_EMBED_PREFIXES):
                from . import embed as _site_embed
                return _site_embed.framable(resp)
            if not embed.embeddable(path):
                # Refuse in words. A blank frame gets reported as a broken
                # integration; a named path gets fixed by whoever configured
                # the menu link.
                out = make_response(embed.refuse(path), 403)
                out.mimetype = "text/plain"
                return embed.framable(out)
            return embed.framable(resp)
        except Exception:  # noqa: BLE001 — never 500 a page over its chrome
            return resp

    @app.after_request
    def _inject_sidebar_response(resp):
        try:
            if resp.status_code != 200:
                return resp
            if not (resp.mimetype or "").startswith("text/html"):
                return resp
            path = request.path or "/"
            # Sign-in and the client-facing pages are deliberately chrome-free.
            if any(path.startswith(p) for p in CHROMELESS):
                return resp
            # A page inside a frame already sits in somebody else's navigation.
            # HubBar applies this test to every mounted module; the hub app's
            # own pages had no equivalent, so Client 360 inside Suite would
            # have arrived with a second full sidebar in it.
            #
            # Asked directly rather than read off a flag the other handler
            # sets: Flask runs after_request handlers in reverse registration
            # order, so this one runs FIRST and any flag would still be unset.
            # Both call hub.embed so they cannot drift apart.
            from . import suite_embed as _embed
            if _embed.is_embedded(request.environ):
                return resp
            if resp.direct_passthrough:
                return resp
            body = resp.get_data()
            if b"s1hub-sb" in body or b'class="sidebar"' in body:
                return resp                      # already has one
            if b"</body>" not in body:
                return resp                      # a fragment, not a page
            from .sidebar import render_sidebar, collapses_by_default
            # A creative tool is itself a full-width workbench and opens with
            # the nav as an icon rail. The Display Ad Builder is a
            # three-column bench -- controls, canvas, size rail -- and every
            # other creative tool wants the same room for the same reason,
            # which is why the list lives in hub/sidebar.py and all three
            # renderers of the nav read it. A stored preference still wins,
            # so this is a starting point rather than the page overruling
            # anybody.
            bar = render_sidebar(_hub_active(path),
                is_admin=viewer_is_admin(),
                collapsed_default=collapses_by_default(path))
            # The help/demo/autofill layer has to come with the sidebar. It
            # was injected by HubBar for dispatcher-mounted modules and by
            # base.html for hub pages, which left blueprint-registered pages
            # — Tickets, Calculators, Page Images, Google Access, Stale
            # Creative — with neither. No scripts means no bubbles and no
            # walkthrough button, silently.
            # Tag <body> so the walkthrough launcher knows which tool it's on.
            # Only when the page hasn't already declared one.
            if b"data-module=" not in body and b"<body" in body:
                seg = path.strip("/").split("/")
                slug = seg[1] if len(seg) > 1 and seg[0] == "tools" else (seg[0] if seg else "")
                mod = {"tickets": "tickets", "calculators": "calculators",
                       "page-images": "page_image_optimizer",
                       "google-access": "google_access",
                       "image-picker": "image_picker",
                       "sites-match": "sites_admin",
                       "domains": "sites_admin",
                       "google-match": "google_access",
                       "stale-creative": "qa", "qa": "qa"}.get(slug, "")
                if mod:
                    body = re.sub(rb"<body\b",
                                  b'<body data-module="' + mod.encode() + b'"',
                                  body, count=1)

            extra = b""
            if b"hub-help.js" not in body:
                extra = (b'<script defer src="/hub-help.js"></script>'
                         b'<script defer src="/hub-demo.js"></script>'
                    b'<script defer src="/hub-crumbs.js"></script>'
                         b'<script defer src="/hub-thinking.js"></script>'
                         b'<script defer src="/hub-autofill.js"></script>'
                         b'<script defer src="/hub-accordion.js"></script>')
            # The third code path. hub/templates/base.html links these for the
            # Hub's own pages and wsgi.py's HubBar injects them into the twenty
            # dispatcher-mounted modules -- and a blueprint registered on the
            # hub app passes through neither, which is the same three-way split
            # hub-thinking.js already names. Google Access, the Image Picker,
            # Page Image Optimizer, Tickets, the Calculators, Video Search and
            # the Commercial Builder all arrive here, so a stylesheet added to
            # only the first two reaches none of them: theme.css was missing
            # from every one of those pages, and adopting the shared record-page
            # look on them did nothing at all until this line existed.
            if b"assets/theme.css" not in body and b"</head>" in body:
                body = body.replace(
                    b"</head>",
                    b'<link rel="stylesheet" href="/assets/theme.css">'
                    b'<link rel="stylesheet" href="/assets/hub-detail.css"></head>', 1)
            if b"hub-help.css" not in body and b"</head>" in body:
                body = body.replace(
                    b"</head>",
                    b'<link rel="stylesheet" href="/hub-help.css"></head>', 1)
            # The last </body>, not the first — see the note in wsgi.py's
            # HubBar. These are blueprint pages we do not control either.
            cut = body.rfind(b"</body>")
            resp.set_data(body[:cut] + bar + extra + body[cut:])
        except Exception:  # noqa: BLE001 — never break a page over navigation
            pass
        return resp

    # ---------------- v7.9 blueprint tools ----------------
    # These ship as Flask blueprints rather than standalone apps, so they
    # register on the hub app directly. Each is wrapped: a tool that fails to
    # load must degrade to "that tool is missing", never to a dead Hub.
    for _label, _mod, _fn, _prefix in (
        ("Calculators", "modules.calculators", "register_calculators", "/tools/calculators"),
        ("Google Access", "modules.google_access", "register_google_access", "/tools/google-access"),
        ("Image Picker", "modules.image_picker", "register_image_picker", "/tools/image-picker"),
        ("Page Image Optimizer", "modules.page_image_optimizer", "register", "/tools/page-images"),
        ("Web Tickets", "modules.tickets", "register_tickets", "/tools/tickets"),
        # The Display Ad Builder is a Node service in the same container; this
        # registers the proxy that puts it behind the Hub login. Same wrapper
        # as the rest, so a renderer that will not start costs the Hub nothing.
        ("Display Ad Builder", "hub.ad_builder_proxy", "register", "/tools/display-ads"),
        # The client and proposal joins. Registered separately from the
        # proxy so a fault in one does not take the other down: the
        # builder is still usable without attach, and attach still
        # explains itself if the renderer is down.
        ("Display Ad Builder links", "hub.ad_builder_link", "register", "/tools/display-ads"),
    ):
        try:
            _m = __import__(_mod, fromlist=[_fn])
            _register = getattr(_m, _fn)
            import inspect
            if "url_prefix" in inspect.signature(_register).parameters:
                _register(app, url_prefix=_prefix)
            else:
                _register(app)
        except Exception as _tool_exc:  # noqa: BLE001
            try:
                errors.log_exception("hub", _tool_exc)
            except Exception:  # noqa: BLE001
                pass

    # ---------------- Commercial Builder ----------------
    # A blueprint, not a standalone Flask app, so it registers here rather
    # than mounting through DispatcherMiddleware in wsgi.py.
    try:
        from modules.commercial_builder import register_commercial_builder
        register_commercial_builder(app)
    except Exception as _cb_exc:  # noqa: BLE001
        try:
            errors.log_exception("hub", _cb_exc)
        except Exception:  # noqa: BLE001
            pass

    # ---------------- Video Backgrounds ----------------
    # Also a blueprint, and for the same reason: /tools/video-backgrounds is
    # not one of the prefixes wsgi.py mounts, so it belongs to the hub app.
    try:
        from modules.video_backgrounds import register_video_backgrounds
        register_video_backgrounds(app)
    except Exception as _vb_exc:  # noqa: BLE001
        try:
            errors.log_exception("hub", _vb_exc)
        except Exception:  # noqa: BLE001
            pass

    # ---------------- User accounts ----------------
    # Registered after init_db (models bind to the shared instance) and before
    # the help layer, so /diagnostics/users exists by the time the sidebar
    # renders. Seeds the founding super admins on first boot.
    try:
        from .users_routes import register_users
        register_users(app)
        app.config["HUB_USERS_REGISTERED"] = True
    except Exception as _users_exc:  # noqa: BLE001
        # This failing is why /signup returned 404 with nothing to go on:
        # Flask-SQLAlchemy was missing from requirements.txt, the import
        # raised, and the except swallowed it. Record the reason so
        # /login/health can say so instead of leaving you guessing.
        app.config["HUB_USERS_REGISTERED"] = False
        app.config["HUB_USERS_BOOT_ERROR"] = (
            f"{type(_users_exc).__name__}: {_users_exc}")
        try:
            errors.log_exception("hub", _users_exc)
        except Exception:  # noqa: BLE001
            pass

    # ---------------- Inbound Suite (GoHighLevel) webhooks ----------------
    try:
        from .ghl_hooks import register_ghl_hooks
        register_ghl_hooks(app)
    except Exception as _hook_exc:  # noqa: BLE001
        app.config["HUB_HOOKS_BOOT_ERROR"] = str(_hook_exc)
        try:
            errors.log_exception("hub", _hook_exc)
        except Exception:  # noqa: BLE001
            pass

    @app.route("/api/scheduler/run/<name>", methods=["POST"])
    def api_scheduler_run(name):
        """Run one job now, without waiting for its next slot."""
        gate = _require_api()
        if gate:
            return gate
        from . import scheduler as _s
        audit.log("scheduler", "manual_run", actor=current_user(), job=name)
        return jsonify(_s.run_now(name, app))

    # ---------------- background jobs ----------------
    # Started last, so every module it might call is registered first. Exactly
    # one worker actually runs jobs — see hub/scheduler.py for why that
    # matters with two gunicorn workers.
    try:
        from . import scheduler as _sched
        _sched.start(app)
    except Exception as _sched_exc:  # noqa: BLE001
        app.config["HUB_SCHEDULER_BOOT_ERROR"] = str(_sched_exc)
        try:
            errors.log_exception("hub", _sched_exc)
        except Exception:  # noqa: BLE001
            pass

    # ---------------- Partner resource pages ----------------
    # Static, self-contained pages served behind the login. Same defensive
    # registration as everything else here: a page that fails to load must
    # cost the dashboard a button, not the Hub a boot.
    try:
        from .partner import register as register_partner
        register_partner(app)
    except Exception as _pp_exc:  # noqa: BLE001
        try:
            errors.log_exception("hub", _pp_exc)
        except Exception:  # noqa: BLE001
            pass

    # ---------------- Stale Creative audit ----------------
    # Same defensive registration as the help layer below: an audit that fails
    # to load must not take the Hub with it.
    try:
        from .stale_creative import register_stale_creative
        register_stale_creative(app)

        from .image_audit import register_image_audit
        register_image_audit(app)
    except Exception as _sc_exc:  # noqa: BLE001
        try:
            errors.log_exception("hub", _sc_exc)
        except Exception:  # noqa: BLE001
            pass

    # ---------------- Prospect 360 ----------------
    # The record a scanned business gets before it is a client. Blueprint, so
    # the login gate sits on the blueprint itself -- every route here names a
    # real person and their phone number.
    try:
        from .prospect_routes import register_prospect
        register_prospect(app)
    except Exception as _pr_exc:  # noqa: BLE001
        try:
            errors.log_exception("hub", _pr_exc)
        except Exception:  # noqa: BLE001
            pass

    # ---------------- Website Audit (Tools > Sales) ----------------
    # A blueprint rather than a mounted module, because everything it reads --
    # the client registry, the discovered-URL overlay, the scan facts and the
    # lead store -- is the hub's own. The login gate is on the blueprint
    # itself: wsgi.py's AuthGuard only wraps dispatcher-mounted modules, and
    # hub/auth.py names what happens when a blueprint is left to guard itself
    # view by view.
    try:
        from .website_audit_routes import register_website_audit
        register_website_audit(app)
    except Exception as _wa_exc:  # noqa: BLE001
        try:
            errors.log_exception("hub", _wa_exc)
        except Exception:  # noqa: BLE001
            pass

    # Imported for its side effect: the Presence model has to exist before the
    # create_all() below or `hub_presence` is never created, and every read of
    # it fails into "we could not look" for ever.
    try:
        from . import presence as _presence  # noqa: F401
    except Exception:  # noqa: BLE001
        pass

    # Create any tables the newly registered blueprints declared. Runs AFTER
    # all of them, so a module registered later still gets its tables. Guarded:
    # a sleeping database must not take the Hub down at boot.
    try:
        from .extensions import create_all as _create_all
        _tbl_err = _create_all(app)
        if _tbl_err:
            app.config["HUB_DB_BOOT_ERROR"] = _tbl_err
    except Exception:  # noqa: BLE001
        pass

    # Refill the persistent disk from the database if this is a *new* disk.
    # JSON files on /var/data are outside the database backup and do not
    # survive the disk being recreated, so hub/jsonstore.py mirrors the ones
    # that are the only copy of something. On an ordinary boot this is two
    # cheap queries and a no-op; on the first boot after a disk is recreated
    # it is the whole recovery, and it has to happen here rather than lazily
    # because /diagnostics and the sidebar read those files before any user
    # does. Guarded like every other boot step — but recorded, not swallowed.
    try:
        from . import jsonstore
        app.config["HUB_JSONSTORE_RESTORE"] = jsonstore.maybe_restore()
    except Exception as _js_exc:  # noqa: BLE001
        app.config["HUB_JSONSTORE_RESTORE"] = {
            "ran": False, "reason": f"{type(_js_exc).__name__}: {_js_exc}"}
        try:
            errors.log_exception("jsonstore", _js_exc)
        except Exception:  # noqa: BLE001
            pass

    # ---------------- v7: help bubbles, tool walkthroughs, demo mode -------
    # Registered last, because it needs _hub_user (defined with the login
    # routes above). The fallbacks below are not decoration: an earlier build
    # wrapped this in a bare try/except, the registration failed on a NameError,
    # and every page in the Hub 500'd on an undefined `demo_banner`. A failure
    # here must degrade to "no help" — never to "no Hub".
    try:
        from .help_routes import register_help
        register_help(app, current_user_fn=_hub_user)
    except Exception as _help_exc:  # noqa: BLE001
        from markupsafe import Markup as _M
        app.jinja_env.globals.setdefault("demo_banner", lambda *a, **k: _M(""))
        app.jinja_env.globals.setdefault("help_dot", lambda *a, **k: _M(""))
        app.jinja_env.globals.setdefault("demo_launcher", lambda *a, **k: _M(""))
        app.jinja_env.globals.setdefault("help_text", lambda *a, **k: "")
        # True, not False: with the help layer down there is no tour to serve
        # either way, and this keeps a failure here degrading to exactly
        # today's markup rather than quietly changing it.
        app.jinja_env.globals.setdefault("has_tour", lambda *a, **k: True)
        try:
            errors.log_exception("hub", _help_exc)
        except Exception:  # noqa: BLE001
            pass

    return app
